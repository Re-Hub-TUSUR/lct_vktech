"""Анимации, переходы и скрытые элементы Студии в собранном .pptx (export/motion.py)."""

from io import BytesIO

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from app.pipeline.export.motion import build_timing, resolve, target_rects, _Shape
from app.pipeline.export.pptx_export import build_deck
from app.pipeline.theme import slide_theme
from app.schemas.slide import Slide, SlideAnimation, SlideShape, SlideTransition, slide_key


def _anim(target: str, trigger: str = "click", **extra) -> SlideAnimation:
    return SlideAnimation(target=target, trigger=trigger, **extra)


def test_timing_groups_clicks_with_and_after():
    steps = [
        (_anim("title", "after", type="up"), [2]),
        (_anim("subtitle", "with", type="fade", delay_ms=200), [3]),
        (_anim("bullets.0", "click", type="wipe"), [4]),
        (_anim("bullets.1", "after", type="zoom", duration_ms=400), [5]),
    ]
    timing = build_timing(steps)
    xml = etree.tostring(timing).decode()
    main = timing.find(f".//{qn('p:cTn')}[@nodeType='mainSeq']")
    groups = main.find(qn("p:childTnLst")).findall(qn("p:par"))
    assert len(groups) == 2, "первая группа стартует сама, вторая — по щелчку"
    assert 'evt="onBegin"' in xml
    assert xml.count('nodeType="clickEffect"') == 1
    assert xml.count('nodeType="withEffect"') == 1
    assert xml.count('nodeType="afterEffect"') == 2
    assert 'filter="wipe(left)"' in xml and "ppt_y" in xml and "ppt_w" in xml
    ids = [int(node.get("id")) for node in timing.iter(qn("p:cTn"))]
    assert len(ids) == len(set(ids)), "id узлов времени уникальны"


def test_timing_empty_without_shapes():
    assert build_timing([(_anim("title"), [])]) is None


def test_resolve_prefers_contained_shapes():
    shapes = [
        _Shape(element=None, spid=1, rect=(0, 0, 1, 1)),  # фон
        _Shape(element=None, spid=2, rect=(0.05, 0.05, 0.5, 0.1)),  # заголовок
        _Shape(element=None, spid=3, rect=(0.05, 0.3, 0.25, 0.3)),  # карточка
        _Shape(element=None, spid=4, rect=(0.07, 0.32, 0.2, 0.05)),  # текст карточки
    ]
    rects = {"title": (0.05, 0.05, 0.5, 0.1), "columns.0": (0.05, 0.3, 0.25, 0.3), "shape:c": (0.05, 0.3, 0.25, 0.3)}
    assert [s.spid for s in resolve("title", rects, shapes, set())] == [2]
    assert [s.spid for s in resolve("columns.0.title", rects, shapes, set())] == [3, 4]
    assert [s.spid for s in resolve("shape:c", rects, shapes, set())] == [3]
    assert [s.spid for s in resolve("columns.0", rects, shapes, {3})] == [4]


def test_target_rects_include_shapes():
    spec = Slide(kind="bullets", title="T", shapes=[SlideShape(id="a", kind="rect", x=0.1, y=0.2, width=0.3, height=0.1)])
    assert target_rects({"title": [0, 0, 1, 0.1]}, spec)["shape:a"] == (0.1, 0.2, 0.3, 0.1)


def test_motion_fields_do_not_change_slide_key():
    base = Slide(kind="bullets", title="T", bullets=["a"])
    moved = base.model_copy(update={"animations": [_anim("title")], "transition": SlideTransition(type="push")})
    assert slide_key(base) == slide_key(moved)
    hidden = base.model_copy(update={"hidden_elements": ["title"]})
    assert slide_key(base) != slide_key(hidden)


def test_deck_carries_animations_transition_and_hidden(parsed):
    path, result = parsed
    theme = slide_theme(result.tokens)
    slides = [
        Slide(
            kind="bullets",
            title="Сервис сокращает подготовку колоды",
            bullets=["Шаблон", "Аудит", "Экспорт"],
            layout_key=result.roles.get("content"),
            animations=[_anim("title", type="up"), _anim("bullets", "after", type="fade")],
            transition=SlideTransition(type="push"),
        ),
        Slide(
            kind="title",
            title="Цифровой дизайнер",
            subtitle="Скрытый подзаголовок",
            layout_key=result.roles.get("title") or result.roles.get("content"),
            hidden_elements=["subtitle"],
        ),
    ]
    deck = build_deck(path, slides, roles=result.roles, theme=theme, patterns=result.patterns, strict=False)
    prs = Presentation(BytesIO(deck.data))
    first, second = prs.slides[0]._element, prs.slides[1]._element
    assert first.find(qn("p:timing")) is not None
    assert first.find(qn("p:transition")) is not None
    children = [child.tag for child in first]
    assert children.index(qn("p:transition")) < children.index(qn("p:timing")), "порядок элементов p:sld"
    spids = {int(node.get("spid")) for node in first.iter(qn("p:spTgt"))}
    present = {shape.shape_id for shape in prs.slides[0].shapes}
    assert spids and spids <= present, "анимации ссылаются на фигуры слайда"
    texts = " ".join(shape.text_frame.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert "Скрытый подзаголовок" not in texts
    assert "subtitle" not in deck.targets[1]
