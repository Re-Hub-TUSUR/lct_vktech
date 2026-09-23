"""Парсер 0.5: форматы шаблонов, паттерны из макетов, служебные паттерны, диаграммы, токены дизайн-системы."""

import io
import json
import zipfile
from dataclasses import asdict

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from app.core.config import settings
from app.pipeline.export import pattern_fill
from app.pipeline.export.pptx_export import build_deck
from app.pipeline.layout.matcher import pattern_score
from app.pipeline.parsing import patterns as patterns_module
from app.pipeline.parsing.formats import PPTX_MAIN, normalize_template
from app.pipeline.parsing.layout_patterns import PH_BASE
from app.pipeline.parsing.template_parser import parse_template
from app.pipeline.theme import slide_theme
from app.schemas.slide import ChartData, ChartSeries, Slide
from app.services import template_cache

POTX_MAIN = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"


def test_layouts_count(parsed):
    _, result = parsed
    assert result.layouts_count == len(result.layouts)


def test_potx_is_normalized(parsed, tmp_path):
    path, _ = parsed
    potx = tmp_path / "template.potx"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(potx, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(PPTX_MAIN.encode(), POTX_MAIN.encode())
            out.writestr(info, data)
    with pytest.raises(Exception):
        Presentation(str(potx))  # python-pptx не открывает шаблон .potx как есть
    normalized = normalize_template(potx, tmp_path / "normalized.pptx")
    assert len(Presentation(str(normalized)).slides) > 0


def test_design_tokens(parsed):
    _, result = parsed
    tokens = result.tokens
    typography = tokens["typography"]
    assert typography and typography["title"]["font"] and typography["body"]["font"]
    grid = tokens["grid"]
    assert grid and all(value is None or 0 <= value < 0.4 for value in grid["margins"].values())
    assert tokens["chrome"] is not None and "slide_number" in tokens["chrome"]
    assert "fonts_check" in tokens and "table_style" in tokens and "chart_style" in tokens


def test_every_pattern_has_schema(parsed):
    _, result = parsed
    for pattern in result.patterns:
        assert pattern["schema"]["type"] == pattern["type"]
        if pattern["usable"]:
            assert pattern["schema"]["title"] == pattern["title"]["chars"]


def test_layout_patterns_reference_placeholders(parsed):
    _, result = parsed
    for pattern in (p for p in result.patterns if p.get("origin") == "layout"):
        assert pattern["usable"] and pattern["layout_key"]
        ids = [pattern["title"]["id"], *(s["id"] for s in pattern["slots"])]
        ids += [s["id"] for item in pattern["items"] for s in item["slots"]]
        assert all(slot_id >= PH_BASE for slot_id in ids)


def test_service_patterns(parsed):
    _, result = parsed
    for pattern in (p for p in result.patterns if p.get("service")):
        assert pattern["category"] in {"cover", "section", "closing"}
        assert not pattern["usable"], "служебный пример не должен уходить под обычный контент"


def test_chart_pattern_gets_new_data(parsed):
    path, result = parsed
    charts = [p for p in result.patterns if p["usable"] and p["type"] == "chart"]
    if not charts:
        pytest.skip("в шаблоне нет простых диаграмм-примеров")
    pattern = charts[0]
    slide = Slide(
        kind="chart",
        title="Время подготовки колоды",
        chart=ChartData(
            type=pattern["frames"]["chart"]["kind"],
            categories=["Вручную", "Ассистент", "Сервис"],
            series=[ChartSeries(name="Минуты", values=[480, 180, 5])],
        ),
        pattern_id=pattern["id"],
    )
    assert pattern_score(slide, pattern) is not None
    theme = slide_theme(result.tokens, result.layouts, result.roles)
    deck = build_deck(path, [slide], roles=result.roles, theme=theme, patterns=result.patterns)
    built = Presentation(io.BytesIO(deck.data))
    charts_on_slide = [shape for shape in built.slides[0].shapes if getattr(shape, "has_chart", False) and shape.has_chart]
    assert charts_on_slide, "диаграмма слайда-примера должна попасть в клон"
    assert list(charts_on_slide[0].chart.plots[0].categories) == ["Вручную", "Ассистент", "Сервис"]


def test_broken_slide_does_not_break_parse(parsed, monkeypatch):
    path, _ = parsed
    original = patterns_module.analyze_slide

    def flaky(slide, index, *args, **kwargs):
        if index == 1:
            raise RuntimeError("битый слайд")
        return original(slide, index, *args, **kwargs)

    monkeypatch.setattr(patterns_module, "analyze_slide", flaky)
    result = parse_template(path)
    assert result.patterns[1]["error"].startswith("RuntimeError")
    assert not result.patterns[1]["usable"]
    assert result.tokens["stats"]["broken_slides"] == 1


def test_template_cache_round_trip(parsed, tmp_path, monkeypatch):
    """Кэш разбора: паттерны и рендеры возвращаются как есть, чужой кэш не подходит."""
    source, result = parsed
    storage, restored = tmp_path / "storage", tmp_path / "restored"
    (storage / "slides").mkdir(parents=True)
    (storage / "slides" / "1.jpg").write_bytes(b"jpeg")
    monkeypatch.setattr(settings, "template_cache_dir", tmp_path / "cache")
    monkeypatch.setattr(settings, "vlm_base_url", None)  # в тесте разбор без VLM

    # Каталога кэша нет — сохранять некуда, читать нечего: разбор идёт обычным путём.
    template_cache.save(source, storage, result)
    assert template_cache.load(source, restored) is None

    (tmp_path / "cache").mkdir()
    template_cache.save(source, storage, result)
    cached = template_cache.load(source, restored)
    assert cached is not None
    assert (restored / "slides" / "1.jpg").read_bytes() == b"jpeg"
    assert asdict(cached) == json.loads(json.dumps(asdict(result), ensure_ascii=False))

    # Разметка другой моделью и другая версия парсера — кэш непригоден.
    monkeypatch.setattr(settings, "vlm_base_url", "http://vlm.test/v1")
    monkeypatch.setattr(settings, "vlm_model", "other-model")
    assert template_cache.load(source, restored) is None
    monkeypatch.setattr(settings, "vlm_base_url", None)
    entry = template_cache.entry_dir(source) / template_cache.ENTRY_FILE
    entry.write_text(json.dumps({**json.loads(entry.read_text("utf-8")), "parser_version": "0.0.1"}), "utf-8")
    assert template_cache.load(source, restored) is None


def test_font_scale_fits_text_into_slot():
    """Текст длиннее образца из шаблона ужимается по рамке слота, а не наезжает сам на себя."""
    fit = pattern_fill.Fit(720, 405)  # слайд 10 × 5.63 дюйма из шаблона VK Tech
    # Слот крупного числа: образец «7» помещается, «40%» в ту же рамку — уже нет.
    kpi = {"size": 166.0, "bbox": [0.019, 0.324, 0.26, 0.426]}
    assert fit.scale(kpi, "7") == 1.0, "как в шаблоне — кегль не трогаем"
    assert 0.4 < fit.scale(kpi, "40%") < 0.7, "три широких знака в рамке под один — сильное уменьшение"
    wide = {"size": 166.0, "bbox": [0.344, 0.324, 0.472, 0.426]}
    assert fit.scale(wide, "10") == 1.0
    assert 0.8 < fit.scale(wide, "95%") < 1.0, "рамка шире — хватает небольшого уменьшения"
    body = {"size": 14, "bbox": [0.1, 0.3, 0.4, 0.3]}
    assert fit.scale(body, "тезис " * 20) == 1.0, "текст помещается — кегль не трогаем"
    assert fit.scale(body, "слово " * 4000) == pattern_fill.MIN_FONT_SCALE, "ниже предела не уменьшаем"
    assert fit.scale(None, "текст") == 1.0 and fit.scale({"size": 0, "bbox": []}, "текст") == 1.0


def _txbody(size: str = ' sz="16600"'):
    """Минимальный txBody с одним прогоном — как у слота крупного числа в шаблоне."""
    return etree.fromstring(
        '<a:txBody xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:bodyPr/><a:p><a:r><a:rPr lang="ru"{size}/><a:t>7</a:t></a:r></a:p></a:txBody>'
    )


def test_scaled_font_size_lands_in_runs():
    """Кегль уменьшается прямо в прогонах: fontScale из normAutofit LibreOffice не применяет."""
    body = _txbody()
    assert pattern_fill.set_body(body, ["40%"], scale=0.5)
    assert [run.get("sz") for run in body.iter(qn("a:rPr"))] == ["8300"]
    assert body.find(f'.//{qn("a:normAutofit")}') is not None, "PowerPoint сможет пересчитать сам"
    assert body.find(f'.//{qn("a:normAutofit")}').get("fontScale") is None

    kept = _txbody()
    pattern_fill.set_body(kept, ["95"])
    assert [run.get("sz") for run in kept.iter(qn("a:rPr"))] == ["16600"], "помещается — кегль не трогаем"

    # Прогон без своего кегля (размер наследуется от плейсхолдера) получает размер слота.
    inherited = _txbody(size="")
    pattern_fill.set_body(inherited, ["40%"], scale=0.5, base_size=166.0)
    assert [run.get("sz") for run in inherited.iter(qn("a:rPr"))] == ["8300"]
