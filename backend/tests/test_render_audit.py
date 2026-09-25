"""Проверки аудита по рендеру (audit/geometry.py) на синтетических фактах и на настоящей сборке шаблона."""

from io import BytesIO

import pymupdf
import pytest
from PIL import Image

from app.pipeline.audit.catalog import AuditContext, limits
from app.pipeline.audit.geometry import derived, family
from app.pipeline.audit.runner import run_checks
from app.pipeline.export.facts import template_inventory
from app.pipeline.export.pptx_export import add_picture_fill, build_deck
from app.pipeline.export.render import libreoffice_available, pptx_to_pdf
from app.pipeline.theme import slide_theme
from app.schemas.slide import Slide
from app.services.renders import text_lines

THEME = {
    "background": "#FFFFFF",
    "text": "#000000",
    "accent": "#0077FF",
    "accents": ["#0077FF"],
    "cover": {"background": "#0077FF", "text": "#FFFFFF"},
    "section": {"background": "#0077FF", "text": "#FFFFFF"},
    "closing": {"background": "#0077FF", "text": "#FFFFFF"},
}
TOKENS = {
    "theme": {"fonts": {"major": "Montserrat", "minor": "Inter"}, "colors": {"accent1": "#0077FF", "dk1": "#000000"}},
    "grid": {"margins": {"left": 0.06, "right": 0.06, "top": 0.06, "bottom": 0.06}},
    "type_scale": [{"size_pt": 32}, {"size_pt": 20}, {"size_pt": 14}],
}
INVENTORY = {"colors": ["#0077FF", "#FFFFFF", "#222222"], "fonts": ["Montserrat", "Inter"], "sizes": [32.0, 20.0, 14.0]}


def line(text: str, x0: float, y0: float, x1: float, y1: float, size: float = 20, block: int = 0) -> dict:
    return {"text": text, "bbox": [x0, y0, x1, y1], "size": size, "block": block}


def page(lines=(), targets=None, **facts) -> dict:
    return {"facts": {"lines": list(lines), "fonts": [], "colors": [], "sizes": [], "pictures": [], "chrome": [], **facts},
            "targets": targets or {}}


def audit(slides: list[Slide], renders: dict, **overrides) -> tuple[dict, dict]:
    options = {"render_status": "ready", **overrides}
    ctx = AuditContext(
        theme=THEME, layout_keys={"0:1"}, limits=limits(), tokens=TOKENS, inventory=INVENTORY, renders=renders, **options
    )
    results, runs = run_checks(slides, ctx)
    by_check: dict[str, list] = {}
    for result in results:
        by_check.setdefault(result.finding.check_id, []).append(result.finding)
    return by_check, runs


def slide(sid: str = "s1", **fields) -> Slide:
    return Slide(**({"id": sid, "kind": "bullets", "title": "Сервис сокращает подготовку колоды", "layout_key": "0:1",
                     "bullets": ["Тезис"]} | fields))


TITLE = [0.06, 0.08, 0.88, 0.12]
BODY = [0.06, 0.25, 0.88, 0.55]
CLEAN_LINES = [
    line("Сервис сокращает подготовку колоды", 0.07, 0.10, 0.60, 0.16, 32),
    line("• Шаблон разбирается автоматически, модель пишет тексты слайдов", 0.07, 0.27, 0.85, 0.31),
    line("• Три варианта вёрстки", 0.07, 0.33, 0.40, 0.37),
    line("• Аудит до показа", 0.07, 0.39, 0.35, 0.43),
    line("• Экспорт в pptx, pdf и html", 0.07, 0.45, 0.45, 0.49),
    line("• Правки из чата", 0.07, 0.51, 0.30, 0.55),
    line("• Кэш вердиктов модели", 0.07, 0.57, 0.40, 0.61),
    line("• Колода за пять минут", 0.07, 0.63, 0.40, 0.67),
]


def test_clean_render_passes_and_states_are_done():
    found, runs = audit([slide()], {"s1": page(CLEAN_LINES, {"title": TITLE, "bullets": BODY}, sizes=[20.0], fonts=["Inter"])})
    render_checks = [check_id for check_id, run in runs.items() if run.check.method == "render"]
    assert not {check_id for check_id in found} & set(render_checks), found
    assert all(runs[check_id].state == "done" for check_id in render_checks)


def test_layout_defects_are_found():
    lines = [
        *CLEAN_LINES[:2],
        line("Наложенный текст поверх тезиса", 0.08, 0.275, 0.50, 0.305),  # легла на строку тезиса
        line("Строка ушла за правый край слайда", 0.70, 0.01, 1.08, 0.05),
        # Текст рамки bullets (низ — 0.80) продолжается ниже неё — переполнение.
        line("• тезис у нижней границы рамки", 0.07, 0.73, 0.40, 0.77),
        line("• тезис на самой границе", 0.07, 0.78, 0.40, 0.82),
        line("• хвост тезиса ниже рамки", 0.07, 0.83, 0.40, 0.87),
    ]
    targets = {"title": TITLE, "bullets": [0.06, 0.25, 0.88, 0.55], "chart": [0.5, 0.5, 0.6, 0.3]}
    found, _ = audit([slide()], {"s1": page(lines, targets)})
    assert "layout.overlap" in found
    assert "layout.text_clipped" in found
    assert found["layout.out_of_bounds"][0].target == "chart"
    assert found["layout.text_overflow"][0].target == "bullets"


def test_alignment_margins_and_fill():
    lines = [
        line("Заголовок слайда", 0.078, 0.10, 0.50, 0.16, 32),  # на 14 px правее текста
        line("• Тезис у самого края", 0.067, 0.27, 0.50, 0.31),
        line("Подпись у края", 0.005, 0.93, 0.20, 0.96, 14),
    ]
    targets = {"title": TITLE, "bullets": [0.06, 0.25, 0.88, 0.5], "subtitle": [0.0, 0.9, 0.3, 0.08]}
    found, _ = audit([slide()], {"s1": page(lines, targets)})
    assert "layout.alignment" in found and "px" in found["layout.alignment"][0].message
    assert "layout.margins" in found
    sparse, _ = audit([slide()], {"s1": page(CLEAN_LINES[:2], {"title": TITLE, "bullets": [0.06, 0.25, 0.5, 0.1]})})
    assert "много пустого места" in sparse["density.fill_ratio"][0].message


def test_template_conformance():
    facts = {
        "fonts": ["Comic Sans MS", "Inter SemiBold"],
        "colors": ["#0077FF", "#7FBBFF", "#FF00AA"],  # акцент, его оттенок к белому и чужой цвет
        "sizes": [20.0, 17.3],
        "pictures": [{"bbox": [0.5, 0.3, 0.4, 0.4], "stretch": 1.25, "image": "abc", "user": True}],
        "chrome": [{"kind": "logo", "bbox": [0.85, 0.88, 0.1, 0.08]}],
        "only_pictures": True,
    }
    lines = [line("Текст поверх логотипа", 0.80, 0.885, 0.97, 0.955)]
    found, _ = audit([slide()], {"s1": page(lines, {"title": TITLE, "bullets": [0.06, 0.25, 0.92, 0.72]}, **facts)})
    assert found["template.fonts"][0].evidence == "Comic Sans MS"
    assert any("гарнитуры" in f.message for f in found["template.fonts"]), "Montserrat, Inter и Comic Sans — три"
    assert found["template.palette"][0].evidence == "#FF00AA"
    assert found["template.font_sizes"][0].evidence == "17.3 pt"
    assert "растянута по горизонтали" in found["layout.image_aspect"][0].message
    assert found["template.logo_position"][0].evidence == "logo"
    assert "integrity.raster_slide" in found


def test_text_squeezed_below_type_scale():
    found, _ = audit([slide()], {"s1": page(CLEAN_LINES, {"title": TITLE, "bullets": BODY}, sizes=[6.0, 20.0])})
    assert "мельче шкалы шаблона" in found["template.font_sizes"][0].message


def test_render_states():
    slides = [slide("s1"), slide("s2", title="Второй слайд про сервис")]
    _, runs = audit(slides, {"s1": page(CLEAN_LINES, {"title": TITLE, "bullets": BODY})}, render_status="pending")
    assert runs["layout.overlap"].state == "pending", "у второго слайда ещё нет рендера"
    assert runs["layout.text_overflow"].state == "pending"
    found, runs = audit(slides, {}, render_status="failed", render_error="bad zip", render_error_kind="build")
    assert found["integrity.file_opens"][0].slide_id is None
    assert runs["layout.overlap"].state == "unavailable" and runs["layout.text_overflow"].state == "fallback"


def test_palette_mix_and_font_family():
    base = [(0, 119, 255), (255, 255, 255), (0, 0, 0)]
    assert derived((127, 187, 255), base, 10), "оттенок акцента к белому — из палитры"
    assert not derived((255, 0, 170), base, 10)
    assert family("Montserrat SemiBold") == family("Montserrat") == "montserrat"


@pytest.fixture
def deck_parts(parsed, tmp_path):
    path, result = parsed
    image = tmp_path / "wide.png"
    Image.new("RGB", (1600, 400), "#3366AA").save(image)
    return path, result, image


def test_build_facts_and_picture_crop(deck_parts):
    path, result, image = deck_parts
    theme = slide_theme(result.tokens)
    slides = [
        Slide(kind="title", title="Цифровой дизайнер", subtitle="Презентации за минуты"),
        Slide(kind="bullets", title="Сервис сокращает подготовку колоды", bullets=["Шаблон", "Аудит"]),
    ]
    for spec in slides:
        spec.layout_key = result.roles.get("content")
    deck = build_deck(path, slides, roles=result.roles, theme=theme, patterns=result.patterns, strict=False)
    assert len(deck.facts) == len(slides)
    assert all({"fonts", "colors", "sizes", "pictures", "chrome", "only_pictures"} <= set(f) for f in deck.facts)
    inventory = template_inventory(path)
    assert inventory["colors"] and inventory["sizes"]

    from pptx import Presentation

    prs = Presentation(BytesIO(deck.data))
    picture = add_picture_fill(prs.slides[1], image, 0, 0, 3_000_000, 3_000_000)
    assert picture.crop_left == pytest.approx(0.375, abs=1e-3), "широкая картинка в квадрате — срезаны бока"


@pytest.mark.skipif(not libreoffice_available(), reason="LibreOffice не установлен")
def test_real_render_lines(deck_parts):
    """Настоящий рендер: строки текста PDF извлекаются в долях страницы, заголовок — на слайде."""
    path, result, _ = deck_parts
    theme = slide_theme(result.tokens)
    slides = [Slide(kind="bullets", title="Сервис сокращает подготовку колоды", bullets=["Шаблон", "Аудит"],
                    layout_key=result.roles.get("content"))]
    deck = build_deck(path, slides, roles=result.roles, theme=theme, patterns=result.patterns, strict=False)
    with pymupdf.open(stream=pptx_to_pdf(deck.data), filetype="pdf") as doc:
        lines = text_lines(doc[0])
    assert any("подготовку" in item["text"] for item in lines)
    assert all(-0.01 <= item["bbox"][0] <= 1.01 for item in lines)
