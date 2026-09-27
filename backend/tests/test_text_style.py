"""Оформление текста из редактора (шрифт, масштаб, цвет, жирность, выравнивание) доходит до .pptx —
и на слайдах, собранных из слайдов-примеров шаблона, и на слайдах на макете."""

from io import BytesIO

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

from app.pipeline.export.pptx_export import build_deck
from app.pipeline.export.text_style import TextStyle, style_paragraph
from app.pipeline.theme import slide_theme
from tests.test_patterns import compose

STYLE = {"font": "Georgia", "color": "#C0392B", "bold": False, "align": "center", "scale": 1.5}


def test_style_paragraph_sets_run_and_paragraph_properties():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_textbox(0, 0, 1_000_000, 500_000).text_frame
    run = frame.paragraphs[0].add_run()
    run.text = "Текст"
    run.font.size = Pt(20)
    run.font.bold = True

    style_paragraph(frame.paragraphs[0]._p, TextStyle(**STYLE))

    props = run._r.find(qn("a:rPr"))
    assert props.get("sz") == "3000"
    assert props.get("b") == "0"
    assert props.find(qn("a:latin")).get("typeface") == "Georgia"
    assert props.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}").get("val") == "C0392B"
    assert frame.paragraphs[0]._p.find(qn("a:pPr")).get("algn") == "ctr"
    # Порядок элементов rPr по схеме: заливка раньше шрифта — иначе PowerPoint просит «восстановить» файл.
    tags = [child.tag for child in props]
    assert tags.index(qn("a:solidFill")) < tags.index(qn("a:latin"))
    buffer = BytesIO()
    prs.save(buffer)
    Presentation(BytesIO(buffer.getvalue()))


def _styled(slide):
    """Поля слайда, которые правит тест: заголовок и первый элемент содержимого."""
    fields = ["title"]
    if slide.kind in {"cards", "two_columns"} and slide.columns:
        fields.append("columns.0.title")
    elif slide.kind == "kpi" and slide.kpis:
        fields.append("kpis.0.value")
    elif slide.kind == "timeline" and slide.timeline:
        fields.append("timeline.0.title")
    elif slide.bullets:
        fields.append("bullets.0")
    slide.element_fonts = {field: STYLE["font"] for field in fields}
    slide.element_colors = {field: STYLE["color"] for field in fields}
    slide.element_aligns = {field: STYLE["align"] for field in fields}
    return fields


def test_editor_text_style_reaches_pptx(parsed):
    path, result = parsed
    theme = slide_theme(result.tokens, result.layouts, result.roles)
    slides = compose(parsed)[0].slides
    for slide in slides:
        _styled(slide)
    deck = build_deck(path, slides, roles=result.roles, theme=theme, patterns=result.patterns)
    prs = Presentation(BytesIO(deck.data))

    for spec, slide in zip(slides, prs.slides):
        xml = slide._element
        fonts = {latin.get("typeface") for latin in xml.iter(qn("a:latin"))}
        colors = {clr.get("val") for clr in xml.iter(qn("a:srgbClr"))}
        aligns = {ppr.get("algn") for ppr in xml.iter(qn("a:pPr"))}
        where = f"слайд {spec.kind} {spec.pattern_id or 'на макете'}"
        assert "Georgia" in fonts, f"шрифт заголовка не попал в .pptx ({where})"
        assert "C0392B" in colors, f"цвет заголовка не попал в .pptx ({where})"
        assert "ctr" in aligns, f"выравнивание не попало в .pptx ({where})"
