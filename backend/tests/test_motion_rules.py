"""Автоанимации по правилам (pipeline/visuals/motion_rules.py) и их перенос в .pptx."""

from io import BytesIO

from pptx import Presentation
from pptx.oxml.ns import qn

from app.pipeline.export.pptx_export import build_deck
from app.pipeline.theme import slide_theme
from app.pipeline.visuals.motion_rules import MAX_STEPS, animate_deck, animate_slide
from app.schemas.slide import ChartData, ChartSeries, Column, ImageSpec, Kpi, Slide, SlideAnimation, SlideTransition, TimelineItem


def _targets(slide: Slide) -> list[str]:
    return [animation.target for animation in slide.animations]


def test_bullets_appear_one_by_one_automatically():
    slide = animate_slide(Slide(kind="bullets", title="T", bullets=["a", "b", "c"]), "subtle")
    assert _targets(slide) == ["bullets.0", "bullets.1", "bullets.2"]
    assert {a.trigger for a in slide.animations} == {"after"}, "слайд оживает сам, без кликера"
    assert {a.type for a in slide.animations} == {"up"}
    assert slide.animations[0].delay_ms > 0 and slide.animations[1].delay_ms == 0
    assert slide.transition == SlideTransition(type="fade", duration_ms=500)
    assert "title" not in _targets(slide), "заголовок открывается вместе со слайдом"


def test_kinds_get_their_blocks():
    cards = Slide(kind="cards", title="T", columns=[Column(title="a"), Column(title="b"), Column(title="c")])
    kpi = Slide(kind="kpi", title="T", kpis=[Kpi(value="1"), Kpi(value="2")])
    timeline = Slide(kind="timeline", title="T", timeline=[TimelineItem(title="a"), TimelineItem(title="b")])
    chart = Slide(kind="chart", title="T", chart=ChartData(categories=["a"], series=[ChartSeries(name="s", values=[1])]))
    cover = Slide(kind="title", title="T", subtitle="S", image=ImageSpec(prompt="p", url="/x.png"))
    assert _targets(animate_slide(cards, "subtle")) == ["columns.0", "columns.1", "columns.2"]
    assert _targets(animate_slide(kpi, "lively")) == ["kpis.0", "kpis.1"]
    assert {a.type for a in animate_slide(timeline, "subtle").animations} == {"wipe"}
    assert _targets(animate_slide(chart, "subtle")) == ["chart"]
    animated_cover = animate_slide(cover, "lively")
    assert _targets(animated_cover) == ["subtitle", "image"]
    assert animated_cover.animations[1].trigger == "with"
    assert animated_cover.transition.type == "zoom"
    assert animate_slide(Slide(kind="section", title="T", subtitle="02"), "lively").transition.type == "push"


def test_long_lists_animate_as_one_block():
    slide = animate_slide(Slide(kind="bullets", title="T", bullets=[str(i) for i in range(MAX_STEPS + 1)]), "subtle")
    assert _targets(slide) == ["bullets"]


def test_hidden_elements_are_skipped():
    slide = animate_slide(Slide(kind="bullets", title="T", bullets=["a", "b"], hidden_elements=["bullets.1"]), "subtle")
    assert _targets(slide) == ["bullets.0"]


def test_manual_animations_survive_without_overwrite():
    manual = [SlideAnimation(target="title", type="zoom")]
    slide = Slide(kind="bullets", title="T", bullets=["a"], animations=manual, transition=SlideTransition(type="wipe"))
    assert animate_slide(slide, "lively") == slide
    replaced = animate_slide(slide, "lively", overwrite=True)
    assert _targets(replaced) == ["bullets.0"] and replaced.transition.type == "fade"


def test_none_mode():
    slide = Slide(kind="bullets", title="T", bullets=["a"])
    assert animate_deck([slide], "none") == [slide]
    animated = animate_slide(slide, "subtle")
    cleared = animate_slide(animated, "none", overwrite=True)
    assert cleared.animations == [] and cleared.transition is None


def test_rules_reach_pptx(parsed):
    path, result = parsed
    theme = slide_theme(result.tokens)
    slides = animate_deck(
        [
            Slide(kind="bullets", title="Сервис", bullets=["Шаблон", "Аудит", "Экспорт"], layout_key=result.roles.get("content")),
            Slide(
                kind="cards",
                title="Карточки",
                columns=[Column(title="A", bullets=["x"]), Column(title="B", bullets=["y"]), Column(title="C", bullets=["z"])],
                layout_key=result.roles.get("content"),
            ),
        ],
        "subtle",
    )
    deck = build_deck(path, slides, roles=result.roles, theme=theme, patterns=result.patterns, strict=False)
    prs = Presentation(BytesIO(deck.data))
    for index in range(len(slides)):
        element = prs.slides[index]._element
        assert element.find(qn("p:timing")) is not None, f"слайд {index + 1}: анимации в .pptx"
        assert element.find(qn("p:transition")) is not None, f"слайд {index + 1}: переход в .pptx"
        effects = [node for node in element.iter(qn("p:cTn")) if node.get("presetClass") == "entr"]
        assert effects, f"слайд {index + 1}: есть эффекты появления"
    first = prs.slides[0]._element
    paragraphs = {node.get("st") for node in first.iter(qn("p:pRg"))}
    if paragraphs:
        # Тезисы на макете — абзацы одного плейсхолдера: каждый появляется сам, список строится по абзацам.
        assert paragraphs == {"0", "1", "2"}
        assert first.find(f".//{qn('p:bldP')}").get("build") == "p"
