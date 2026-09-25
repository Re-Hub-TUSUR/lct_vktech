"""Сквозной прогон без БД: план → три варианта → визуализация → аудит → экспорт в три формата."""

from io import BytesIO

import pymupdf
import pytest
from pptx import Presentation

from app.pipeline.audit.catalog import AuditContext, limits
from app.pipeline.audit.runner import run_checks
from app.pipeline.export.html_export import build_html
from app.pipeline.export.pdf_export import build_simple_pdf
from app.pipeline.export.pptx_export import build_pptx
from app.pipeline.export.render import libreoffice_available, pptx_to_pdf
from app.pipeline.layout.composer import axis_config, compose_single_variant, compose_variants
from app.pipeline.parsing.renders import render_template
from app.pipeline.planning.outline import build_outline, group_labels
from app.pipeline.theme import slide_theme
from app.pipeline.visuals.visualizer import VisualsReport, refine_variant
from tests.plans import plan

BRIEF = (
    "Сервис собирает презентации в фирменном стиле за минуты. "
    "Команды тратят часы на ручную вёрстку слайдов. "
    "Дизайнеры перегружены типовыми задачами. "
    "AI-инструменты не соблюдают брендбук."
)

def compose(parsed, axis: str = "density"):
    _, result = parsed
    return compose_variants(
        plan(),
        axis_key=axis,
        slides_count=12,
        speaker_notes=True,
        layouts=result.layouts,
        roles=result.roles,
        groups=group_labels(),
    )


@pytest.mark.parametrize("axis", ["density", "visualization", "structure", "layouts"])
def test_three_distinct_variants(parsed, axis):
    drafts = compose(parsed, axis)
    assert [d.letter for d in drafts] == ["A", "B", "C"]
    signatures = {tuple((s.kind, s.title) for s in d.slides) for d in drafts}
    assert len(signatures) == 3, "варианты должны различаться"
    _, result = parsed
    keys = {layout["key"] for layout in result.layouts}
    for draft in drafts:
        assert draft.slides[0].kind == "title" and draft.slides[-1].kind == "closing"
        assert all(slide.layout_key in keys for slide in draft.slides)
    assert axis_config(axis)[0] == axis


def test_audit_runs_on_composed_variants(parsed):
    """Без рендера и без модели: проверки по данным выполняются, по рендеру недоступны,
    контекстуальные — по запасным правилам там, где правило есть."""
    _, result = parsed
    ctx = AuditContext(
        theme=slide_theme(result.tokens),
        layout_keys={layout["key"] for layout in result.layouts},
        sources_text=BRIEF,
        limits=limits(),
    )
    for draft in compose(parsed):
        results, runs = run_checks(draft.slides, ctx)
        assert "template.layout_from_template" not in {r.finding.check_id for r in results}
        assert runs["density.max_bullets"].state == "done"
        assert runs["layout.overlap"].state == "unavailable"
        assert runs["layout.text_overflow"].state == "fallback", "без рендера — оценка по ёмкости рамки"
        assert runs["content.title_is_insight"].state == "fallback"
        assert runs["content.typos"].state == "unavailable"


def test_exports(parsed):
    path, result = parsed
    theme = slide_theme(result.tokens)
    slides = compose(parsed)[1].slides

    pptx_bytes = build_pptx(path, slides, roles=result.roles, theme=theme)
    deck = Presentation(BytesIO(pptx_bytes))
    assert len(deck.slides) == len(slides)
    assert deck.slides[0].shapes.title.text == slides[0].title

    html = build_html(slides, theme, title="Тест")
    assert html.count('class="slide ') == len(slides)

    pdf = pymupdf.open(stream=build_simple_pdf(slides, theme), filetype="pdf")
    assert pdf.page_count == len(slides)


def test_deck_without_model_uses_brief_only(parsed):
    """Запасной путь без модели: колода из тезисов брифа, ничего не придумано."""
    _, result = parsed
    specs = build_outline(title="Цифровой дизайнер презентаций", purpose="product", brief=BRIEF)
    assert specs[0]["kind"] == "title" and specs[-1]["kind"] == "closing"
    assert specs[0]["subtitle"].rstrip(".") == "Сервис собирает презентации в фирменном стиле за минуты"
    body = specs[1:-1]
    assert body and all(spec["kind"] == "bullets" for spec in body)
    text = " ".join(" ".join([spec["topic"], *spec["bullets"]]) for spec in body)
    assert "Дизайнеры перегружены типовыми задачами" in text
    drafts = compose_variants(
        specs, axis_key="density", slides_count=12, speaker_notes=False,
        layouts=result.layouts, roles=result.roles, groups=group_labels(),
    )
    assert all(draft.slides[0].kind == "title" for draft in drafts)


def test_visuals_refine_composed_variants(parsed):
    """Стадия визуализаций: ключевая категория, названная автором текстов, выделена на диаграмме."""
    _, result = parsed
    drafts = compose(parsed, "visualization")
    names = {layout["key"]: layout["name"] for layout in result.layouts}
    report = VisualsReport()
    for draft in drafts:
        report.merge(refine_variant(draft.slides, result.patterns, position=draft.position, layout_names=names))
    charts = [slide.chart for draft in drafts for slide in draft.slides if slide.chart]
    assert charts and report.charts == len(charts)
    assert all(chart.categories[chart.highlight] == "С сервисом" for chart in charts if chart.type in {"bar", "horizontal_bar"})
    assert "Диаграмм" in report.detail()


def test_compose_single_variant(parsed):
    _, result = parsed
    specs = plan()
    draft1 = compose_single_variant(
        specs,
        axis_key="density",
        position=1,
        slides_count=12,
        speaker_notes=True,
        layouts=result.layouts,
        roles=result.roles,
        groups=group_labels(),
        seed=42,
    )
    draft2 = compose_single_variant(
        specs,
        axis_key="density",
        position=1,
        slides_count=12,
        speaker_notes=True,
        layouts=result.layouts,
        roles=result.roles,
        groups=group_labels(),
        seed=99,
    )
    assert draft1.letter == "B"
    assert draft2.letter == "B"
    assert len(draft1.slides) == len(draft2.slides)


def test_layout_skins_cover_all_roles(parsed):
    _, result = parsed
    theme = slide_theme(result.tokens, result.layouts, result.roles, image_url=lambda layout: layout["key"])
    for role, key in result.roles.items():
        if key:
            skin = theme["layouts"][key]
            assert skin["frames"]["body"], f"у макета роли {role} нет области контента"
            assert skin["title"]["color"].startswith("#")


needs_libreoffice = pytest.mark.skipif(not libreoffice_available(), reason="LibreOffice не установлен")


@needs_libreoffice
def test_pdf_keeps_template_design(parsed):
    """Регрессия «пустого фона»: PDF рендерится из .pptx и содержит оформление мастеров и макетов."""
    path, result = parsed
    theme = slide_theme(result.tokens)
    slides = compose(parsed)[0].slides
    pdf = pymupdf.open(stream=pptx_to_pdf(build_pptx(path, slides, roles=result.roles, theme=theme)), filetype="pdf")
    assert pdf.page_count == len(slides)
    cover = pdf[0]
    # Упрощённый рендер давал ровно один прямоугольник фона и ни одной картинки.
    assert len(cover.get_images()) + len(cover.get_drawings()) > 1


@needs_libreoffice
def test_template_renders(parsed, tmp_path):
    path, result = parsed
    renders = render_template(path, tmp_path / "slides", tmp_path / "layouts")
    assert len(renders.slides) == result.slides_count
    assert set(renders.layouts) == {layout["key"] for layout in result.layouts}
    assert (tmp_path / "layouts" / renders.layouts[result.roles["title"]]).stat().st_size > 5_000
    theme = slide_theme(result.tokens, result.layouts, result.roles)
    images = {key: (tmp_path / "layouts" / name).read_bytes() for key, name in renders.layouts.items()}
    html = build_html(compose(parsed)[0].slides, theme, title="Тест", images=images)
    assert "data:image/jpeg;base64," in html
