"""Картинки в .pptx по модели PowerPoint — как их показывает редактор (frontend/src/lib/imageCrop.ts):
картинка не растягивается, обрезка исходника плюс срез по центру под пропорции рамки."""

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from app.pipeline.export.pptx_export import _apply_crop


def _picture(tmp_path, size, frame):
    path = tmp_path / "pic.png"
    Image.new("RGB", size, "#3366ff").save(path)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(str(path), 0, 0, Emu(frame[0]), Emu(frame[1]))
    return pic, path


def _crop(pic):
    return tuple(round(v, 3) for v in (pic.crop_top, pic.crop_right, pic.crop_bottom, pic.crop_left))


def test_without_crop_picture_fills_frame_like_cover(tmp_path):
    # Широкая картинка 2:1 в квадратной рамке — бока срезаются поровну, без растягивания.
    pic, path = _picture(tmp_path, (200, 100), (1_000_000, 1_000_000))
    _apply_crop(pic, None, path)
    assert _crop(pic) == (0.0, 0.25, 0.0, 0.25)


def test_user_crop_is_shown_exactly_when_frame_matches(tmp_path):
    # Пользователь оставил левую половину (квадрат 100×100) — рамка квадратная, лишнего не срезается.
    pic, path = _picture(tmp_path, (200, 100), (1_000_000, 1_000_000))
    _apply_crop(pic, [0, 0.5, 0, 0], path)
    assert _crop(pic) == (0.0, 0.5, 0.0, 0.0)


def test_user_crop_then_center_cut_to_frame_aspect(tmp_path):
    # Оставили левую половину, но рамка широкая 2:1 — от квадрата срезаются верх и низ по центру.
    pic, path = _picture(tmp_path, (200, 100), (2_000_000, 1_000_000))
    _apply_crop(pic, [0, 0.5, 0, 0], path)
    assert _crop(pic) == (0.25, 0.5, 0.25, 0.0)


def test_invalid_crop_is_ignored(tmp_path):
    pic, path = _picture(tmp_path, (100, 100), (1_000_000, 1_000_000))
    _apply_crop(pic, [0.6, 0, 0.6, 0], path)
    assert _crop(pic) == (0.0, 0.0, 0.0, 0.0)
