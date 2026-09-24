"""Паттерны слайдов: разбор слайдов-примеров, подбор под контент, сборка клонов и заполнение слотов."""

import re
from io import BytesIO

import pytest
from lxml import etree
from pptx import Presentation

from app.pipeline.export.pattern_fill import fill_table, shape_elements
from app.pipeline.export.pptx_export import build_deck
from app.pipeline.layout.composer import compose_single_variant, compose_variants
from app.pipeline.layout.matcher import assign_patterns, fits, pattern_score
from app.pipeline.parsing.patterns import CATEGORIES, USABLE_TYPES, is_usable
from app.pipeline.planning.outline import build_outline, group_labels
from app.pipeline.theme import slide_theme
from app.schemas.slide import Column, Slide, SlideShape, TableData

BRIEF = "Сервис собирает презентации в фирменном стиле за минуты."
# Текст-заглушка слайдов-примеров, который не должен попадать в готовую колоду.
LEFTOVER_RX = re.compile(r"^(заголовок|текст|пункт|описание|lorem ipsum|вставить|иллюстрация|кейс|заметка)\b", re.I)
CONTENT_KINDS = {"bullets", "cards", "two_columns", "kpi", "timeline", "agenda", "table", "image_text"}


def compose(parsed, axis: str = "density"):
    _, result = parsed
    specs = build_outline(title="Цифровой дизайнер презентаций", purpose="product", brief=BRIEF)
    return compose_variants(
        specs,
        axis_key=axis,
        slides_count=12,
        speaker_notes=True,
        layouts=result.layouts,
        roles=result.roles,
        groups=group_labels(),
        patterns=result.patterns,
    )


def test_patterns_are_extracted(parsed):
    _, result = parsed
    patterns = result.patterns
    # Паттерны из макетов (origin=layout) — не слайды-примеры, их в колоде шаблона нет.
    assert len([p for p in patterns if p.get("origin") != "layout"]) == result.slides_count
    usable = [pattern for pattern in patterns if pattern["usable"]]
    assert len(usable) >= 10, "в шаблоне должно найтись хотя бы 10 паттернов под контент"
    assert patterns[0]["category"] == "cover"
    assert {pattern["category"] for pattern in patterns} <= set(CATEGORIES)
    for pattern in usable:
        assert is_usable(pattern) and pattern["type"] in USABLE_TYPES
        assert pattern["category"] == "content" and pattern["title"]
        for item in pattern["items"]:
            assert item["slots"], f"у элемента паттерна {pattern['id']} нет слотов"
            assert all(0 <= value <= 1.05 for value in item["bbox"])
    assert result.tokens["stats"]["patterns"] == len(usable)


def test_service_slides_are_not_content(parsed):
    """Обложки, финальные слайды, визитки, код и витрины не должны уходить под контент."""
    _, result = parsed
    for pattern in result.patterns:
        title = (pattern["title"] or {}).get("sample", "").lower()
        if "спасибо" in title or "визитк" in title:
            assert not pattern["usable"], pattern["id"]


def test_most_content_slides_get_patterns(parsed):
    for draft in compose(parsed):
        content = [slide for slide in draft.slides if slide.kind in CONTENT_KINDS]
        patterned = [slide for slide in content if slide.pattern_id]
        assert len(patterned) >= len(content) // 2, f"вариант {draft.letter}: паттерны у {len(patterned)} из {len(content)}"
        _, result = parsed
        by_id = {pattern["id"]: pattern for pattern in result.patterns}
        for slide in patterned:
            assert fits(slide, by_id[slide.pattern_id])
            assert slide.layout_key == by_id[slide.pattern_id]["layout_key"]


def test_variants_differ_in_patterns(parsed):
    drafts = compose(parsed)
    signatures = {tuple(slide.pattern_id for slide in draft.slides) for draft in drafts}
    assert len(signatures) > 1


def test_reroll_changes_patterns(parsed):
    _, result = parsed
    specs = build_outline(title="Цифровой дизайнер презентаций", purpose="product", brief=BRIEF)

    def reroll(seed: int):
        draft = compose_single_variant(
            specs,
            axis_key="density",
            position=0,
            slides_count=12,
            speaker_notes=False,
            layouts=result.layouts,
            roles=result.roles,
            groups=group_labels(),
            seed=seed,
            patterns=result.patterns,
        )
        return tuple(slide.pattern_id for slide in draft.slides)

    assert len({reroll(seed) for seed in range(8)}) > 1, "рерол должен давать разную вёрстку"


def test_pattern_score_rejects_overflow(parsed):
    _, result = parsed
    cards = [pattern for pattern in result.patterns if pattern["usable"] and pattern["type"] == "cards"]
    if not cards:
        pytest.skip("в шаблоне нет карточек")
    pattern = cards[0]
    too_many = Slide(kind="cards", title="Много карточек", columns=[Column(title="Пункт", bullets=["Текст"])] * 12)
    assert pattern_score(too_many, pattern) is None
    huge = Slide(kind="cards", title="Длинный текст", columns=[Column(title="Пункт", bullets=["слово " * 400])])
    assert pattern_score(huge, pattern) is None


def test_keep_reassigns_only_misfits(parsed):
    drafts = compose(parsed)
    _, result = parsed
    slides = [slide.model_copy(deep=True) for slide in drafts[0].slides]
    before = [slide.pattern_id for slide in slides]
    assign_patterns(slides, result.patterns, position=0, keep=True)
    assert [slide.pattern_id for slide in slides] == before


def test_template_shapes_follow_pattern(parsed):
    """Фигуры шаблона меняются вместе с паттерном; фигуры пользователя и его удаления сохраняются."""
    _, result = parsed
    by_id = {pattern["id"]: pattern for pattern in result.patterns}
    slide = next(
        (s for d in compose(parsed) for s in d.slides if s.pattern_id and by_id[s.pattern_id].get("shapes")), None
    )
    if slide is None:
        pytest.skip("в шаблоне нет паттернов с фигурами")
    user = SlideShape(kind="circle")

    # Паттерн прежний: удалённая пользователем фигура шаблона не возвращается.
    slide.shapes = [*slide.shapes[1:], user]
    before = [shape.id for shape in slide.shapes]
    assign_patterns([slide], result.patterns, keep=True)
    assert [shape.id for shape in slide.shapes] == before

    # У слайда чужой паттерн (как после автоисправления аудита): у подобранного взамен — его фигуры.
    foreign = next(p for p in result.patterns if p.get("shapes") and pattern_score(slide, p) is None)
    slide.pattern_id = foreign["id"]
    slide.shapes = [SlideShape.model_validate(shape) for shape in foreign["shapes"]] + [user]
    assign_patterns([slide], result.patterns, keep=True)
    assert slide.pattern_id != foreign["id"]
    expected = [shape["template_shape_id"] for shape in (by_id.get(slide.pattern_id or "") or {}).get("shapes") or []]
    assert [shape.template_shape_id for shape in slide.shapes if shape.template_shape_id is not None] == expected
    assert user in slide.shapes


def test_untouched_template_shapes_keep_their_look(parsed):
    """Сборка переносит в клон только правки пользователя: нетронутые фигуры шаблона не перекрашиваются
    и не сдвигаются (цвет парсер знает без прозрачности и оттенков темы)."""
    path, result = parsed
    by_id = {pattern["id"]: pattern for pattern in result.patterns}
    slide = next(
        (s for s in compose(parsed)[0].slides if s.pattern_id and any(shape.color for shape in s.shapes)), None
    )
    if slide is None:
        pytest.skip("в шаблоне нет паттернов с цветными фигурами")
    theme = slide_theme(result.tokens, result.layouts, result.roles)
    source = Presentation(str(path)).slides[by_id[slide.pattern_id]["slide_index"]]
    colored = {shape.template_shape_id for shape in slide.shapes if shape.color}
    before = {sp.shape_id: etree.tostring(sp._element.spPr) for sp in source.shapes if sp.shape_id in colored}

    def build(spec: Slide) -> Presentation:
        deck = build_deck(path, [spec], roles=result.roles, theme=theme, patterns=result.patterns)
        return Presentation(BytesIO(deck.data))

    built = build(slide).slides[0]
    after = {sp.shape_id: etree.tostring(sp._element.spPr) for sp in built.shapes if sp.shape_id in before}
    if not after:
        pytest.skip("цветные фигуры паттерна лежат внутри групп или уходят вместе с лишними элементами")
    assert after == {shape_id: before[shape_id] for shape_id in after}, "нетронутые фигуры шаблона изменились"

    moved = slide.model_copy(deep=True)
    target = next(shape for shape in moved.shapes if shape.template_shape_id in after)
    target.x = round(target.x + 0.05, 3)
    deck = build(moved)
    shifted = next(sp for sp in deck.slides[0].shapes if sp.shape_id == target.template_shape_id)
    assert abs(shifted.left - int(target.x * deck.slide_width)) <= 1, "правка пользователя должна попасть в .pptx"


def test_deck_is_built_from_patterns(parsed):
    path, result = parsed
    theme = slide_theme(result.tokens, result.layouts, result.roles)
    slides = compose(parsed)[0].slides
    deck = build_deck(path, slides, roles=result.roles, theme=theme, patterns=result.patterns)
    prs = Presentation(BytesIO(deck.data))
    assert len(prs.slides) == len(slides), "слайды-примеры шаблона должны быть удалены из колоды"
    assert len(deck.targets) == len(slides)

    for spec, slide, targets in zip(slides, prs.slides, deck.targets):
        assert "title" in targets
        if not spec.pattern_id:
            continue
        texts = [shape.text_frame.text.strip() for shape in slide.shapes if shape.has_text_frame]
        texts += [
            inner.text_frame.text.strip()
            for shape in slide.shapes
            if shape.shape_type == 6
            for inner in shape.shapes
            if inner.has_text_frame
        ]
        assert spec.title in texts, f"заголовок не попал в слайд {spec.pattern_id}"
        leftovers = [text for text in texts if text and LEFTOVER_RX.match(text) and text != spec.title]
        assert not leftovers, f"в слайде {spec.pattern_id} остался текст примера: {leftovers}"


def test_table_pattern_follows_data(parsed):
    path, result = parsed
    tables = [pattern for pattern in result.patterns if pattern["usable"] and pattern["type"] == "table"]
    if not tables:
        pytest.skip("в шаблоне нет таблиц")
    pattern = tables[0]
    prs = Presentation(str(path))
    slide = prs.slides[pattern["slide_index"]]
    frame = shape_elements(slide)[pattern["frames"]["table"]["id"]]
    data = TableData(columns=["Критерий", "Было", "Стало"], rows=[["Скорость", "8 ч", "5 мин"], ["Правки", "3", "1"]])
    fill_table(frame, data)
    table = next(shape for shape in slide.shapes if shape.shape_id == pattern["frames"]["table"]["id"]).table
    assert len(table.rows) == 3 and len(table.columns) == 3
    assert table.cell(0, 0).text == "Критерий" and table.cell(2, 2).text == "1"
