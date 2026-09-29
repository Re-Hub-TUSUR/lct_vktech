"""Иллюстрации text-to-image (ТЗ, п. 3.1): клиент модели, стиль шаблона, задания, сюжеты, вставка в слайды."""

import base64
import copy
import io
from io import BytesIO

import httpx
import pytest
from PIL import Image
from pptx import Presentation

from app.core.config import settings
from app.core.storage import generated_image_path
from app.pipeline.ai import t2i
from app.pipeline.export.facts import digest
from app.pipeline.export.pptx_export import build_deck
from app.pipeline.layout.composer import compose_single_variant, compose_variants
from app.pipeline.planning.outline import group_labels
from app.pipeline.theme import slide_theme
from app.pipeline.visuals import illustrations
from app.pipeline.visuals.illustrations import (
    apply_images,
    brand_colors,
    compose_prompt,
    fallback_style,
    plan_images,
    remember_images,
    run_images,
    style_is_current,
    template_style,
)
from app.schemas.slide import ImageSpec, Slide
from tests.plans import plan

STYLE = {
    "source": "vlm",
    "medium": "3d render",
    "style": "Glossy 3D render of blue glass.",
    "background": "plain light background close to #FAFCFF",
    "subjects": "abstract objects",
    "colors": ["#0077FF"],
}


def png(width: int = 64, height: int = 48, color: str = "blue") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buffer, "PNG")
    return buffer.getvalue()


def image_plan() -> list[dict]:
    """План колоды со слайдом «текст и изображение» — таким, каким его отдаёт автор текстов."""
    specs = plan()
    specs.insert(
        3,
        {
            "ref": "case",
            "kind": "image_text",
            "plan_kind": "image_text",
            "group": "solution",
            "priority": 1,
            "topic": "Кейс",
            "insight": "Команда продаж собрала колоду за пять минут",
            "bullets": ["Бриф на полстраницы", "Шаблон компании", "Колода за пять минут"],
            "image_prompt": "A glass hourglass with a few grains of sand",
        },
    )
    return specs


def enabled_t2i(monkeypatch, generate=None) -> list[dict]:
    calls: list[dict] = []
    cfg = t2i.T2IConfig(
        base_url="https://t2i.test/v1", api_key="k", model="flux.2-klein-4b", api="chat", timeout=5, concurrency=2
    )
    monkeypatch.setattr(t2i, "t2i_config", lambda: cfg)

    def fake(http, prompt, *, ratio, seed, cfg=None):
        calls.append({"prompt": prompt, "ratio": ratio, "seed": seed})
        if generate is not None:
            return generate(prompt)
        return png()

    monkeypatch.setattr(t2i, "generate_image", fake)
    return calls


# --- клиент модели ----------------------------------------------------------------------------------


def test_chat_api_image_is_decoded():
    data = png()

    def handler(request: httpx.Request) -> httpx.Response:
        body = request.read().decode()
        assert '"modalities": ["image"]' in body or '"modalities":["image"]' in body
        assert "aspect_ratio" in body
        url = "data:image/png;base64," + base64.b64encode(data).decode()
        return httpx.Response(200, json={"choices": [{"message": {"images": [{"image_url": {"url": url}}]}}]})

    cfg = t2i.T2IConfig(base_url="https://t2i.test/v1", api_key=None, model="m", api="chat", timeout=5, concurrency=1)
    with httpx.Client(base_url=cfg.base_url, transport=httpx.MockTransport(handler)) as http:
        assert t2i.generate_image(http, "a cube", ratio="4:3", seed=1, cfg=cfg) == data


def test_images_api_retries_rate_limit(monkeypatch):
    data = png()
    answers = [
        httpx.Response(429, json={"error": "slow down"}),
        httpx.Response(200, json={"data": [{"b64_json": base64.b64encode(data).decode()}]}),
    ]
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request.url.path)
        return answers.pop(0)

    monkeypatch.setattr(t2i.time, "sleep", lambda _: None)
    cfg = t2i.T2IConfig(base_url="https://t2i.test/v1", api_key=None, model="m", api="images", timeout=5, concurrency=1)
    with httpx.Client(base_url=cfg.base_url, transport=httpx.MockTransport(handler)) as http:
        assert t2i.generate_image(http, "a cube", ratio="16:9", seed=1, cfg=cfg) == data
    assert seen == ["/v1/images/generations", "/v1/images/generations"]


def test_client_error_is_not_retried(monkeypatch):
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(1)
        return httpx.Response(400, json={"error": "bad prompt"})

    monkeypatch.setattr(t2i.time, "sleep", lambda _: None)
    cfg = t2i.T2IConfig(base_url="https://t2i.test/v1", api_key=None, model="m", api="chat", timeout=5, concurrency=1)
    with httpx.Client(base_url=cfg.base_url, transport=httpx.MockTransport(handler)) as http:
        with pytest.raises(t2i.T2IError):
            t2i.generate_image(http, "a cube", ratio="1:1", seed=1, cfg=cfg)
    assert len(seen) == 1


def test_llm_key_goes_only_to_the_same_host(monkeypatch):
    monkeypatch.setattr(settings, "llm_base_url", "https://routerai.ru/api/v1")
    monkeypatch.setattr(settings, "llm_api_key", "secret")
    monkeypatch.setattr(settings, "t2i_api_key", None)
    monkeypatch.setattr(settings, "t2i_model", "black-forest-labs/flux.2-klein-4b")
    monkeypatch.setattr(settings, "t2i_base_url", "https://routerai.ru/api/v1")
    assert t2i.t2i_config().api_key == "secret"
    monkeypatch.setattr(settings, "t2i_base_url", "http://t2i:8000/v1")
    assert t2i.t2i_config().api_key is None


def test_nearest_ratio():
    assert t2i.nearest_ratio(16 / 9) == "16:9"
    assert t2i.nearest_ratio(1.02) == "1:1"
    assert t2i.nearest_ratio(0.5) == "9:16"
    assert t2i.nearest_ratio(1.4) == "4:3"
    assert t2i.nearest_ratio(0) == "4:3"


# --- стиль шаблона ----------------------------------------------------------------------------------


def test_fallback_style_follows_palette_and_background():
    tokens = {
        "palette": [{"hex": "#FFFFFF"}, {"hex": "#000000"}, {"hex": "#808080"}, {"hex": "#0077FF"}, {"hex": "#FF3885"}],
        "backgrounds": {"content": "#0A0A12"},
    }
    assert brand_colors(tokens) == ["#0077FF", "#FF3885"]
    style = fallback_style(tokens)
    assert style["source"] == "fallback"
    assert style["colors"] == ["#0077FF", "#FF3885"]
    assert "dark" in style["background"]
    assert "light" in fallback_style({"backgrounds": {"content": "#FAFCFF"}})["background"]


def test_template_style_is_described_by_vlm(monkeypatch, tmp_path):
    (tmp_path / "1.jpg").write_bytes(png(320, 180))
    (tmp_path / "7.jpg").write_bytes(png(320, 180, "white"))
    patterns = [
        {"id": "s1", "category": "cover", "render": "1.jpg", "images": []},
        {"id": "s7", "category": "content", "render": "7.jpg", "images": [{"id": 4, "bbox": [0.5, 0.2, 0.4, 0.6]}]},
        {"id": "s9", "category": "library", "render": "9.jpg", "images": [{"id": 4, "bbox": [0, 0, 1, 1]}]},
    ]
    seen = {}

    def fake_ask(http, *, system, user, schema, usage, label, max_tokens=None, cfg=None, images=None, **kwargs):
        seen["images"] = len(images or [])
        seen["user"] = user
        return {
            "medium": "3d render",
            "style": "Glossy translucent blue glass, soft studio light.",
            "background": "very light bluish white",
            "subjects": "cubes, spheres",
            "colors": ["#0077ff", "blue", "#FFFFFF"],
        }

    cfg = illustrations.ModelConfig(
        base_url="https://vlm.test/v1", api_key=None, model="qwen", timeout=5, max_tokens=500, temperature=0.2,
        disable_reasoning=True, extra_body=None, concurrency=1, label="VLM",
    )
    monkeypatch.setattr(illustrations, "_vision_model", lambda: cfg)
    monkeypatch.setattr(illustrations, "ask", fake_ask)
    tokens = {"palette": [{"hex": "#0077FF"}], "backgrounds": {"content": "#FAFCFF"}}
    style = template_style("VK Tech", tokens, patterns, tmp_path)
    assert style["source"] == "vlm"
    assert style["colors"] == ["#0077FF"]  # белый и нераспознанный цвет — не акценты
    assert "glass" in style["style"]
    # Фон картинок — фон слайдов с текстом, а не то, что VLM увидела на обложке.
    assert style["background"] == "plain light background close to #FAFCFF"
    assert style["background_note"] == "very light bluish white"
    # Обложка и слайд с крупной картинкой; витрина иконок и слайд без рендера — нет.
    assert seen["images"] == 2
    assert "VK Tech" in seen["user"]
    assert style_is_current(style)
    assert not style_is_current({**style, "skill_version": "0.0.1"})


def test_template_style_without_vlm_is_fallback(monkeypatch, tmp_path):
    cfg = illustrations.ModelConfig(
        base_url=None, api_key=None, model=None, timeout=5, max_tokens=500, temperature=0.2,
        disable_reasoning=True, extra_body=None, concurrency=1, label="VLM",
    )
    monkeypatch.setattr(illustrations, "_vision_model", lambda: cfg)
    style = template_style("T", {"palette": [{"hex": "#0077FF"}]}, [], tmp_path)
    assert style["source"] == "fallback"
    assert style_is_current(style)


# --- задания и вставка ------------------------------------------------------------------------------


def drafts(axis: str = "density", patterns=None, specs=None):
    return compose_variants(
        specs or image_plan(),
        axis_key=axis,
        slides_count=12,
        speaker_notes=True,
        layouts=[],
        roles={},
        groups=group_labels(),
        patterns=patterns,
        images_available=True,
    )


def test_one_image_per_plan_slide_for_all_variants():
    variants = drafts()
    assert all(any(slide.kind == "image_text" for slide in draft.slides) for draft in variants)
    image_plan_ = plan_images(variants, None, aspect=16 / 9, max_images=8, seed=100)
    assert [job.ref for job in image_plan_.jobs] == ["case"]
    job = image_plan_.jobs[0]
    assert job.seed == 100
    assert job.idea.startswith("A glass hourglass")
    assert job.ratio == "1:1"  # слайд на макете: картинка справа от текста


def test_image_limit_is_respected():
    variants = drafts("layouts")
    full = plan_images(variants, None, aspect=16 / 9, max_images=8)
    assert len(full.jobs) > 1  # вариант C верстает гибкие тезисы с иллюстрацией
    limited = plan_images(variants, None, aspect=16 / 9, max_images=1)
    assert len(limited.jobs) == 1
    assert limited.over_limit == len(full.jobs) - 1


def test_images_are_drawn_saved_and_placed(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "storage_dir", tmp_path)
    monkeypatch.setattr(illustrations, "llm_enabled", lambda: False)
    calls = enabled_t2i(monkeypatch)
    variants = drafts()
    image_plan_ = plan_images(variants, None, aspect=16 / 9)
    result = run_images(image_plan_, presentation_id="p1", title="Колода", purpose="продукт", style=STYLE)
    assert result.generated == 1 and result.failed == 0
    assert result.director == "fallback"
    prompt = calls[0]["prompt"]
    assert prompt.startswith("A glass hourglass")  # без модели сюжет — идея автора текстов
    assert "Glossy 3D render of blue glass." in prompt and "#0077FF" in prompt and "no text" in prompt

    dropped = apply_images(variants, image_plan_, patterns=None, roles={}, layouts=[], aspect=16 / 9)
    assert dropped == 0
    urls = {slide.image.url for draft in variants for slide in draft.slides if slide.kind == "image_text"}
    assert len(urls) == 1
    url = urls.pop()
    name = url.rsplit("/", 1)[1]
    assert url == f"/api/presentations/p1/images/{name}"
    path = generated_image_path("p1", name)
    assert path is not None and path.is_file()
    with Image.open(path) as saved:
        assert saved.format == "JPEG"

    specs = image_plan()
    remember_images(specs, variants)
    assert next(spec for spec in specs if spec["ref"] == "case")["image"]["url"] == url
    # Рерол: вариант собирается из плана с картинкой — рисовать заново нечего.
    reroll = compose_single_variant(
        specs, axis_key="density", position=1, slides_count=12, speaker_notes=True, layouts=[], roles={},
        groups=group_labels(), seed=5, images_available=True,
    )
    case = next(slide for slide in reroll.slides if slide.ref == "case")
    assert case.image.url == url and case.image.status == "generated"
    assert not plan_images([reroll], None, aspect=16 / 9).jobs


def test_director_subjects_are_used(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "storage_dir", tmp_path)
    monkeypatch.setattr(settings, "llm_base_url", "https://llm.test/v1")
    monkeypatch.setattr(settings, "llm_model", "qwen")
    monkeypatch.setattr(illustrations, "llm_enabled", lambda: True)
    seen = {}

    def fake_ask(http, *, system, user, schema, usage, label, **kwargs):
        seen["system"], seen["user"] = system, user
        return {"images": [{"n": 1, "subject": "A glass stopwatch next to a neat stack of cards."}, {"n": 7, "subject": "x"}]}

    monkeypatch.setattr(illustrations, "ask", fake_ask)
    calls = enabled_t2i(monkeypatch)
    image_plan_ = plan_images(drafts(), None, aspect=16 / 9)
    result = run_images(image_plan_, presentation_id="p1", title="Колода", purpose="продукт", style=STYLE)
    assert result.director == "llm"
    assert image_plan_.jobs[0].subject == "A glass stopwatch next to a neat stack of cards"
    assert calls[0]["prompt"].startswith("A glass stopwatch")
    assert "Glossy 3D render of blue glass." in seen["system"]
    assert "Команда продаж собрала колоду за пять минут" in seen["user"]
    assert "A glass hourglass" in seen["user"]  # идея автора текстов видна арт-директору


def test_failed_image_turns_slide_into_bullets(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "storage_dir", tmp_path)
    monkeypatch.setattr(illustrations, "llm_enabled", lambda: False)

    def broken(prompt):
        raise t2i.T2IError("нет связи")

    enabled_t2i(monkeypatch, broken)
    variants = drafts()
    image_plan_ = plan_images(variants, None, aspect=16 / 9)
    result = run_images(image_plan_, presentation_id="p1", title="Колода", purpose="продукт", style=STYLE)
    assert result.generated == 0 and result.failed == 1
    kept = copy.deepcopy(variants[0].slides)
    assert apply_images([kept], image_plan_, patterns=None, roles={}, layouts=[], aspect=16 / 9, drop_missing=False) == 0
    assert any(slide.kind == "image_text" and slide.image and not slide.image.url for slide in kept)

    dropped = apply_images(variants, image_plan_, patterns=None, roles={}, layouts=[], aspect=16 / 9)
    assert dropped == 3  # по слайду в каждом варианте
    case = next(slide for slide in variants[0].slides if slide.ref == "case")
    assert case.kind == "bullets" and case.image is None and case.bullets


def test_prompt_template_fills_style():
    prompt = compose_prompt("A glass cube.", STYLE)
    assert prompt.startswith("A glass cube. Glossy 3D render of blue glass.")
    assert "Background: plain light background close to #FAFCFF" in prompt
    assert "{" not in prompt


def test_generated_image_path_rejects_foreign_names():
    assert generated_image_path("p1", "../../etc/passwd") is None
    assert generated_image_path("p1", "a" * 32 + ".png") is None
    assert generated_image_path("p1", "a" * 32 + ".jpg") is not None


# --- шаблоны: место картинки ------------------------------------------------------------------------


def test_picture_pattern_needs_clear_image_slot():
    from app.pipeline.layout.matcher import image_slot_clear, pattern_score

    body = {"id": 2, "role": "body", "bbox": [0.05, 0.25, 0.4, 0.6], "chars": 300}
    title = {"id": 1, "bbox": [0.05, 0.05, 0.9, 0.12], "chars": 80}
    side = {
        "id": "p1", "type": "image_text", "usable": True, "title": title, "slots": [body],
        "images": [{"id": 5, "bbox": [0.55, 0.2, 0.4, 0.7]}],
    }
    under_text = {**side, "id": "p2", "images": [{"id": 5, "bbox": [0.03, 0.2, 0.45, 0.7]}]}
    no_picture = {**side, "id": "p3", "images": []}
    assert image_slot_clear(side)
    assert not image_slot_clear(under_text)
    assert not image_slot_clear(no_picture)
    slide = Slide(kind="image_text", title="Кейс", bullets=["Бриф на полстраницы", "Колода за пять минут"])
    assert pattern_score(slide, side) is not None
    assert pattern_score(slide, under_text) is None
    assert pattern_score(slide, no_picture) is None


def test_fresh_id_skips_ids_of_removed_template_shapes(tmp_path):
    from pptx.util import Emu

    from app.pipeline.export.pptx_export import _fresh_id

    path = tmp_path / "p.png"
    path.write_bytes(png())
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture = slide.shapes.add_picture(str(path), Emu(0), Emu(0), Emu(100000), Emu(100000))
    _fresh_id(picture, ceiling=40)
    assert picture.shape_id == 41


def test_image_slides_use_only_picture_patterns(parsed):
    _, result = parsed
    by_id = {pattern["id"]: pattern for pattern in result.patterns}
    for draft in drafts(patterns=result.patterns):
        for slide in draft.slides:
            if slide.kind == "image_text" and slide.pattern_id:
                assert by_id[slide.pattern_id]["type"] == "image_text"


def test_generated_picture_replaces_example_picture(parsed, monkeypatch, tmp_path):
    """Картинка встаёт на место фото слайда-примера: та же фигура (рамка, маска, эффекты), новое изображение."""
    path, result = parsed
    from app.pipeline.layout.matcher import image_slot_clear

    candidates = [
        p for p in result.patterns
        if p.get("type") == "image_text" and p.get("usable") and p.get("origin") != "layout" and p.get("images")
    ]
    pattern = next((p for p in candidates if image_slot_clear(p)), candidates[0] if candidates else None)
    if pattern is None:
        pytest.skip("в шаблоне нет слайда-примера с местом под картинку")
    monkeypatch.setattr(settings, "storage_dir", tmp_path)
    presentation_id = "0e5ed156-0000-4000-8000-000000000001"
    name = "b" * 32 + ".jpg"
    target = tmp_path / "generated" / presentation_id / name
    target.parent.mkdir(parents=True)
    buffer = BytesIO()
    Image.new("RGB", (300, 200), "orange").save(buffer, "JPEG")
    target.write_bytes(buffer.getvalue())

    spec = Slide(
        kind="image_text",
        title="Кейс: колода за пять минут",
        bullets=["Бриф на полстраницы", "Шаблон компании"],
        pattern_id=pattern["id"],
        layout_key=pattern.get("layout_key"),
        image=ImageSpec(
            prompt="A glass hourglass", url=f"/api/presentations/{presentation_id}/images/{name}", status="generated"
        ),
    )
    theme = slide_theme(result.tokens, result.layouts, result.roles)
    deck = build_deck(path, [spec], roles=result.roles, theme=theme, patterns=result.patterns)
    slide = Presentation(BytesIO(deck.data)).slides[0]

    def pictures(shapes):
        for shape in shapes:
            if shape.shape_type == 6:
                yield from pictures(shape.shapes)
                continue
            try:
                yield shape, digest(shape.image.blob)
            except (AttributeError, ValueError, KeyError):
                continue

    ours = [shape for shape, key in pictures(slide.shapes) if key == digest(target.read_bytes())]
    assert len(ours) == 1, "картинка должна попасть в слайд ровно один раз"
    source = Presentation(str(path)).slides[pattern["slide_index"]]
    example_ids = {shape.shape_id for shape, _ in pictures(source.shapes)}
    if pattern["images"][0]["id"] in example_ids and pattern["images"][0]["id"] in {s.shape_id for s, _ in pictures(slide.shapes)}:
        assert ours[0].shape_id == pattern["images"][0]["id"], "картинка встаёт в фигуру фото слайда-примера"
    assert any(picture["user"] for picture in deck.facts[0]["pictures"]), "аудит видит картинку как содержимое слайда"


def test_unfinished_images_are_counted_as_late():
    from app.pipeline.visuals.illustrations import ImageJob, ImagePlan, unfinished_result

    job = dict(ratio="1:1", seed=1, title="T", summary="", bullets=[], idea="")
    plan = ImagePlan(jobs=[ImageJob(ref="a", url="/api/presentations/p/images/x.jpg", **job), ImageJob(ref="b", **job)])
    result = unfinished_result(plan, STYLE, seconds=60)
    assert (result.requested, result.generated, result.late) == (2, 1, 1)
    assert "не дождались" in result.detail()
