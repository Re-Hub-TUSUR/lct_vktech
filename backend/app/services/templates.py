import logging
import uuid
from pathlib import Path
from typing import Any

from app.core.config import settings
from app.models import Template
from app.pipeline.ai.vlm import enrich_patterns, vlm_enabled
from app.pipeline.export.render import libreoffice_available
from app.pipeline.parsing import design_tokens
from app.pipeline.parsing.renders import TemplateRenders, layout_image_name, render_template, slide_image_name
from app.pipeline.parsing.patterns import pattern_summary
from app.pipeline.parsing.template_parser import PARSER_VERSION, ParsedTemplate, parse_template
from app.pipeline.theme import slide_theme
from app.services import template_cache

logger = logging.getLogger(__name__)


def template_dir(template_id: uuid.UUID) -> Path:
    return Path(settings.storage_dir) / "templates" / str(template_id)


def layout_images_dir(template_id: uuid.UUID) -> Path:
    return template_dir(template_id) / "layouts"


def slide_images_dir(template_id: uuid.UUID) -> Path:
    return template_dir(template_id) / "slides"


def icons_dir(template_id: uuid.UUID) -> Path:
    return template_dir(template_id) / "icons"


def _icons(file_path: str, directory: Path, patterns: list[dict[str, Any]]) -> list[dict[str, Any]]:
    from pptx import Presentation

    return design_tokens.extract_icons(Presentation(file_path), directory, patterns)


def analyze_template(template_id: uuid.UUID, file_path: str, builtin: bool = False) -> ParsedTemplate:
    """Декомпозиция шаблона, рендеры слайдов и макетов, разметка паттернов VLM (вызывать в потоке).

    builtin=True — шаблон контент-пакета: его готовый разбор лежит в кэше репозитория, тогда
    ни LibreOffice, ни VLM не нужны, а свежий разбор пополняет кэш. Всё, что загрузил
    пользователь, разбирается по-настоящему — кэш на загруженные шаблоны не распространяется."""
    if builtin and (cached := template_cache.load(file_path, template_dir(template_id))) is not None:
        return cached

    parsed = parse_template(file_path)
    renders = TemplateRenders()
    if libreoffice_available():
        try:
            renders = render_template(file_path, slide_images_dir(template_id), layout_images_dir(template_id))
        except Exception:
            # Без рендеров превью останется на цветах темы — это не повод валить разбор шаблона.
            logger.exception("Не удалось отрендерить шаблон %s", template_id)
    else:
        logger.warning("LibreOffice не найден — шаблон %s разобран без рендеров", template_id)

    for layout in parsed.layouts:
        layout["image"] = renders.layouts.get(layout["key"])
    for pattern in parsed.patterns:
        if pattern.get("origin") == "layout":
            # Паттерн из макета: превью — рендер самого макета.
            pattern["render"] = None
            pattern["render_layout"] = renders.layouts.get(pattern["layout_key"])
        else:
            pattern["render"] = renders.slides.get(pattern["slide_index"])
    # Библиотека пиктограмм шаблона — для будущих слайдов с иконками.
    icons = design_tokens.safe("icons", _icons, file_path, icons_dir(template_id), parsed.patterns)
    parsed.tokens["icons"] = icons or []

    if vlm_enabled() and renders.slides:
        parsed.tokens["vlm"] = enrich_patterns(parsed.patterns, slide_images_dir(template_id))
        # Модель могла поменять типы и назначения — пересчитываем сводку паттернов.
        parsed.tokens["stats"]["patterns"] = sum(1 for pattern in parsed.patterns if pattern["usable"])
        parsed.tokens["stats"]["service_patterns"] = sum(1 for pattern in parsed.patterns if pattern.get("service"))
        parsed.tokens["patterns"] = pattern_summary(parsed.patterns)

    if builtin:
        template_cache.save(file_path, template_dir(template_id), parsed)
    return parsed


def vlm_outdated(template: Template) -> bool:
    """VLM включили или сменили модель, а паттерны шаблона размечены без неё (или другой моделью)."""
    if not vlm_enabled():
        return False
    return ((template.tokens or {}).get("vlm") or {}).get("model") != settings.vlm_model


def needs_analysis(template: Template) -> bool:
    return (
        template.status != "ready"
        or not template.layouts
        or template.parser_version != PARSER_VERSION
        or vlm_outdated(template)
    )


def apply_parsed(template: Template, parsed: ParsedTemplate) -> None:
    template.slide_width = parsed.slide_width
    template.slide_height = parsed.slide_height
    template.slides_count = parsed.slides_count
    template.tokens = parsed.tokens
    template.layouts = parsed.layouts
    template.layout_roles = parsed.roles
    template.patterns = parsed.patterns
    template.parser_version = PARSER_VERSION
    template.status = "ready"
    template.error = None


DESIGN_TOKENS = ("typography", "grid", "chrome", "table_style", "chart_style")


def pattern_render_url(template: Template, pattern: dict[str, Any], version: int | None = None) -> str | None:
    """Картинка паттерна: рендер слайда-примера или (для паттерна из макета) рендер макета."""
    if version is None:
        version = int(template.updated_at.timestamp()) if template.updated_at else 0
    if pattern.get("origin") == "layout":
        name = pattern.get("render_layout") or layout_image_name(pattern.get("layout_key") or "0:0")
        return f"/api/templates/{template.id}/layouts/{name}?v={version}"
    render_file = pattern.get("render") or slide_image_name(max(pattern.get("slide_index", 0), 0))
    return f"/api/templates/{template.id}/slides/{render_file}?v={version}"


def template_theme(template: Template | None) -> dict[str, Any]:
    if template is None:
        return slide_theme(None)
    version = int(template.updated_at.timestamp()) if template.updated_at else 0

    def image_url(layout: dict[str, Any]) -> str:
        return f"/api/templates/{template.id}/layouts/{layout_image_name(layout['key'])}?v={version}"

    pattern_layouts = {
        pattern["layout_key"] for pattern in template.patterns or [] if pattern.get("usable") and pattern.get("layout_key")
    }
    theme = slide_theme(
        template.tokens, template.layouts, template.layout_roles, image_url=image_url, extra_keys=pattern_layouts
    )
    patterns = []
    for p in template.patterns or []:
        p_copy = dict(p)
        if p.get("origin") != "layout":
            p_copy["render"] = p.get("render") or slide_image_name(max(p.get("slide_index", 0), 0))
        p_copy["render_url"] = pattern_render_url(template, p, version)
        patterns.append(p_copy)
    theme["patterns"] = patterns
    theme["template_id"] = str(template.id)
    # Токены дизайн-системы для сборки .pptx: стиль таблиц, цвета диаграмм, типографика, обвязка.
    tokens = template.tokens or {}
    theme["icons"] = tokens.get("icons") or []
    theme["design"] = {key: tokens.get(key) for key in DESIGN_TOKENS}
    theme["palette"] = [color["hex"] for color in (tokens.get("palette") or [])[:8] if color.get("hex")]
    return theme


def layout_images(template: Template | None) -> dict[str, bytes]:
    """Рендеры фонов макетов для встраивания в HTML-экспорт."""
    if template is None:
        return {}
    directory = layout_images_dir(template.id)
    images = {}
    for layout in template.layouts or []:
        if layout.get("image") and (path := directory / layout["image"]).exists():
            images[layout["key"]] = path.read_bytes()
    return images


def slide_size_inches(template: Template | None) -> tuple[float, float]:
    size = (template.tokens or {}).get("slide_size", {}) if template is not None else {}
    return size.get("width_in") or 13.333, size.get("height_in") or 7.5
