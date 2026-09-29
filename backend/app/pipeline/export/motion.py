"""Анимации, переходы и скрытые элементы слайда в собранном .pptx.

Редактор адресует элементы так же, как аудит: title, subtitle, bullets.1, columns.0, image, chart…
и shape:<id> для фигур слайда. Сборка возвращает рамку каждого адреса (Targets, доли слайда); по ней
здесь находятся фигуры собранного слайда верхнего уровня, которые этот элемент нарисовали:

- для фигуры слайда — одна фигура, сильнее всего совпадающая с её рамкой;
- для текстового элемента или блока — все фигуры, почти целиком лежащие внутри его рамки
  (текстовое поле, подложка карточки, маркер, иконка), кроме фона и крупного декора.

Скрытые элементы удаляются из слайда (как их не показывает редактор), анимации пишутся в
p:timing нативными эффектами PowerPoint, переход — в p:transition.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from lxml import etree
from pptx.oxml.ns import qn

from app.schemas.slide import Slide, SlideAnimation, SlideShape

logger = logging.getLogger(__name__)

Rect = tuple[float, float, float, float]
# Цель эффекта: фигура целиком (spid) или абзац текстовой фигуры (spid, номер абзаца).
Target = int | tuple[int, int]

# Фигура, занимающая больше этой доли слайда, — фон или подложка макета, а не элемент.
BACKGROUND_SHARE = 0.55
# Доля площади фигуры, которая должна лежать в рамке элемента, чтобы фигура считалась его частью.
INSIDE_SHARE = 0.82
# Допуск рамки: сборка округляет доли до 4 знаков, кегли ужимают текст внутри рамки.
SLACK = 0.012
# Минимальное совпадение (IoU) фигуры слайда с её рамкой.
SHAPE_IOU = 0.55

# Эффект редактора → (presetID, presetSubtype) PowerPoint. Поведение задаётся явно, preset — подпись в UI.
PRESETS = {
    "fade": (10, 0),
    "up": (2, 4),  # «Вылет» снизу
    "zoom": (53, 16),  # «Масштабирование»
    "left": (2, 8),  # «Вылет» слева
    "wipe": (22, 8),  # «Появление» (шторка) слева
}
# Смещение в долях слайда, как в редакторе: 40 px и 120 px на слайде 960×540.
RISE_Y = 40 / 540
SLIDE_X = 120 / 960
TRANSITIONS = {"fade": "fade", "push": "push", "zoom": "zoom", "wipe": "wipe"}
# Пункт списка без своей фигуры: на макете все тезисы — абзацы одного плейсхолдера.
LIST_ITEM = re.compile(r"^(bullets)\.(\d+)$")


@dataclass
class _Shape:
    element: etree._Element
    spid: int
    rect: Rect


def _has_text(shape: _Shape) -> bool:
    return shape.element is not None and any((node.text or "").strip() for node in shape.element.iter(qn("a:t")))


def _paragraph_count(shape: _Shape) -> int:
    body = shape.element.find(qn("p:txBody")) if shape.element is not None else None
    return len(body.findall(qn("a:p"))) if body is not None else 0


def _area(rect: Rect) -> float:
    return max(rect[2], 0.0) * max(rect[3], 0.0)


def _intersection(a: Rect, b: Rect) -> float:
    left, top = max(a[0], b[0]), max(a[1], b[1])
    right, bottom = min(a[0] + a[2], b[0] + b[2]), min(a[1] + a[3], b[1] + b[3])
    return max(0.0, right - left) * max(0.0, bottom - top)


def _grow(rect: Rect, slack: float = SLACK) -> Rect:
    return (rect[0] - slack, rect[1] - slack, rect[2] + 2 * slack, rect[3] + 2 * slack)


def _top_shapes(slide, width: int, height: int) -> list[_Shape]:
    """Фигуры верхнего уровня с рамками в долях слайда (плейсхолдеры — с геометрией макета)."""
    shapes: list[_Shape] = []
    for shape in slide.shapes:
        try:
            left, top, cx, cy = shape.left, shape.top, shape.width, shape.height
        except Exception:  # noqa: BLE001 — у фигуры без геометрии нет и рамки
            continue
        if None in (left, top, cx, cy):
            continue
        rect = (left / width, top / height, cx / width, cy / height)
        shapes.append(_Shape(element=shape._element, spid=int(shape.shape_id), rect=rect))
    return shapes


def target_rects(targets: dict[str, list[float]], spec: Slide) -> dict[str, Rect]:
    """Рамки всех адресов слайда: из сборки и фигуры слайда (shape:<id>). Фигуры — ровно spec.shapes:
    пустой список значит, что фигур шаблона на слайде нет (так их собирает и pptx_export._shapes)."""
    rects: dict[str, Rect] = {key: tuple(value) for key, value in targets.items() if len(value) == 4}  # type: ignore[misc]
    shapes: list[SlideShape] = list(spec.shapes or [])
    for shape in shapes:
        rects[f"shape:{shape.id}"] = (shape.x, shape.y, shape.width, max(shape.height, 0.004))
    return rects


def _parent_target(targets: dict[str, Rect], address: str) -> Rect | None:
    """Рамка адреса или ближайшего родителя: columns.0.title → columns.0 → columns."""
    parts = address.split(".")
    for length in range(len(parts), 0, -1):
        rect = targets.get(".".join(parts[:length]))
        if rect is not None:
            return rect
    return None


def resolve(address: str, rects: dict[str, Rect], shapes: list[_Shape], taken: set[int]) -> list[_Shape]:
    """Фигуры собранного слайда, которые рисуют элемент address (ещё не занятые другими адресами)."""
    free = [shape for shape in shapes if shape.spid not in taken]
    if address.startswith("shape:"):
        rect = rects.get(address)
        if rect is None:
            return []
        best, best_iou = None, 0.0
        for shape in free:
            overlap = _intersection(shape.rect, rect)
            union = _area(shape.rect) + _area(rect) - overlap
            iou = overlap / union if union > 0 else 0.0
            if iou > best_iou:
                best, best_iou = shape, iou
        return [best] if best is not None and best_iou >= SHAPE_IOU else []

    rect = _parent_target(rects, address)
    if rect is None:
        return []
    area = _grow(rect)
    found = []
    for shape in free:
        size = _area(shape.rect)
        # Крупная фигура — фон или подложка, но не текстовый блок: плейсхолдер тезисов бывает больше полслайда.
        if size <= 0 or (size > BACKGROUND_SHARE and not _has_text(shape)):
            continue
        if _intersection(shape.rect, area) / size >= INSIDE_SHARE:
            found.append(shape)
    return found


def resolve_paragraph(address: str, rects: dict[str, Rect], shapes: list[_Shape]) -> tuple[int, int] | None:
    """Абзац текстовой фигуры для пункта списка, у которого нет своей рамки (bullets.2 → 3-й абзац блока
    тезисов). Так PowerPoint анимирует список «по абзацам»."""
    match = LIST_ITEM.match(address)
    if match is None or address in rects:
        return None
    rect = rects.get(match.group(1))
    if rect is None:
        return None
    area, index = _grow(rect), int(match.group(2))
    for shape in shapes:
        size = _area(shape.rect)
        if size <= 0 or not _has_text(shape) or index >= _paragraph_count(shape):
            continue
        if _intersection(shape.rect, area) / size >= INSIDE_SHARE:
            return shape.spid, index
    return None


# --- p:timing ------------------------------------------------------------------------------------


class _Ids:
    def __init__(self) -> None:
        self.value = 0

    def __call__(self) -> str:
        self.value += 1
        return str(self.value)


def _el(tag: str, attrib: dict[str, str] | None = None, *children) -> etree._Element:
    element = etree.Element(qn(tag), attrib or {})
    for child in children:
        if child is not None:
            element.append(child)
    return element


def _cond(delay: str = "0", **attrib: str) -> etree._Element:
    return _el("p:cond", {"delay": delay, **attrib})


def _target(target: Target) -> etree._Element:
    spid, paragraph = (target, None) if isinstance(target, int) else target
    shape = _el("p:spTgt", {"spid": str(spid)})
    if paragraph is not None:
        shape.append(_el("p:txEl", None, _el("p:pRg", {"st": str(paragraph), "end": str(paragraph)})))
    return _el("p:tgtEl", None, shape)


def _behavior(ids: _Ids, spid: Target, duration: int, attr: str | None = None, **extra: str) -> etree._Element:
    ctn = _el("p:cTn", {"id": ids(), "dur": str(duration), "fill": "hold"})
    children = [ctn, _target(spid)]
    if attr:
        children.append(_el("p:attrNameLst", None, _text_el("p:attrName", attr)))
    return _el("p:cBhvr", extra or None, *children)


def _text_el(tag: str, text: str) -> etree._Element:
    element = _el(tag)
    element.text = text
    return element


def _set_visible(ids: _Ids, spid: Target) -> etree._Element:
    ctn = _el("p:cTn", {"id": ids(), "dur": "1", "fill": "hold"}, _el("p:stCondLst", None, _cond("0")))
    behavior = _el(
        "p:cBhvr",
        None,
        ctn,
        _target(spid),
        _el("p:attrNameLst", None, _text_el("p:attrName", "style.visibility")),
    )
    return _el("p:set", None, behavior, _el("p:to", None, _el("p:strVal", {"val": "visible"})))


def _fade(ids: _Ids, spid: Target, duration: int) -> etree._Element:
    return _el("p:animEffect", {"transition": "in", "filter": "fade"}, _behavior(ids, spid, duration))


def _move(ids: _Ids, spid: Target, duration: int, attr: str, start: str, end: str) -> etree._Element:
    tavs = _el(
        "p:tavLst",
        None,
        _el("p:tav", {"tm": "0"}, _el("p:val", None, _el("p:strVal", {"val": start}))),
        _el("p:tav", {"tm": "100000"}, _el("p:val", None, _el("p:strVal", {"val": end}))),
    )
    return _el(
        "p:anim",
        {"calcmode": "lin", "valueType": "num"},
        _behavior(ids, spid, duration, attr, additive="base"),
        tavs,
    )


def _effect_behaviors(ids: _Ids, animation: SlideAnimation, spid: Target) -> list[etree._Element]:
    duration = animation.duration_ms
    behaviors = [_set_visible(ids, spid)]
    kind = animation.type
    if kind == "wipe":
        behaviors.append(
            _el("p:animEffect", {"transition": "in", "filter": "wipe(left)"}, _behavior(ids, spid, duration))
        )
        return behaviors
    behaviors.append(_fade(ids, spid, duration))
    if kind == "up":
        behaviors.append(_move(ids, spid, duration, "ppt_y", f"#ppt_y+{RISE_Y:.4f}", "#ppt_y"))
    elif kind == "left":
        behaviors.append(_move(ids, spid, duration, "ppt_x", f"#ppt_x-{SLIDE_X:.4f}", "#ppt_x"))
    elif kind == "zoom":
        behaviors.append(_move(ids, spid, duration, "ppt_w", "#ppt_w*0.6", "#ppt_w"))
        behaviors.append(_move(ids, spid, duration, "ppt_h", "#ppt_h*0.6", "#ppt_h"))
    return behaviors


NODE_TYPES = {"click": "clickEffect", "with": "withEffect", "after": "afterEffect"}


def _effect_par(ids: _Ids, animation: SlideAnimation, spids: list[Target], node_type: str, group: int) -> etree._Element:
    preset, subtype = PRESETS.get(animation.type, PRESETS["fade"])
    behaviors: list[etree._Element] = []
    for spid in spids:
        behaviors.extend(_effect_behaviors(ids, animation, spid))
    ctn = _el(
        "p:cTn",
        {
            "id": ids(),
            "presetID": str(preset),
            "presetClass": "entr",
            "presetSubtype": str(subtype),
            "fill": "hold",
            "grpId": str(group),
            "nodeType": node_type,
        },
        _el("p:stCondLst", None, _cond(str(animation.delay_ms))),
        _el("p:childTnLst", None, *behaviors),
    )
    return _el("p:par", None, ctn)


def build_timing(steps: list[tuple[SlideAnimation, list[Target]]]) -> etree._Element | None:
    """p:timing по списку (анимация, её фигуры) в порядке показа.

    Щелчок открывает группу; «вместе с предыдущей» идёт в тот же момент, «после предыдущей» — когда
    закончится самый долгий эффект текущего момента. Если первая анимация не по щелчку, группа
    стартует сама при появлении слайда.
    """
    steps = [(animation, spids) for animation, spids in steps if spids]
    if not steps:
        return None
    ids = _Ids()
    root_ctn_id = ids()
    main_ctn_id = ids()

    # Группы щелчков → моменты (параллельные блоки) → эффекты.
    groups: list[list[list[tuple[SlideAnimation, list[Target], str]]]] = []
    for animation, spids in steps:
        trigger = animation.trigger
        if trigger == "click" or not groups:
            groups.append([[(animation, spids, "clickEffect" if trigger == "click" else NODE_TYPES[trigger])]])
        elif trigger == "with":
            groups[-1][-1].append((animation, spids, "withEffect"))
        else:
            groups[-1].append([(animation, spids, "afterEffect")])

    click_pars = []
    for group_index, moments in enumerate(groups):
        first_trigger = moments[0][0][0].trigger
        if group_index == 0 and first_trigger != "click":
            start = _el(
                "p:stCondLst",
                None,
                _cond("indefinite"),
                _el("p:cond", {"evt": "onBegin", "delay": "0"}, _el("p:tn", {"val": main_ctn_id})),
            )
        else:
            start = _el("p:stCondLst", None, _cond("indefinite"))
        offset = 0
        moment_pars = []
        for moment in moments:
            effects = [_effect_par(ids, animation, spids, node_type, 0) for animation, spids, node_type in moment]
            moment_pars.append(
                _el(
                    "p:par",
                    None,
                    _el(
                        "p:cTn",
                        {"id": ids(), "fill": "hold"},
                        _el("p:stCondLst", None, _cond(str(offset))),
                        _el("p:childTnLst", None, *effects),
                    ),
                )
            )
            offset += max(animation.delay_ms + animation.duration_ms for animation, _, _ in moment)
        click_pars.append(
            _el("p:par", None, _el("p:cTn", {"id": ids(), "fill": "hold"}, start, _el("p:childTnLst", None, *moment_pars)))
        )

    main = _el(
        "p:seq",
        {"concurrent": "1", "nextAc": "seek"},
        _el("p:cTn", {"id": main_ctn_id, "dur": "indefinite", "nodeType": "mainSeq"}, _el("p:childTnLst", None, *click_pars)),
        _el("p:prevCondLst", None, _el("p:cond", {"evt": "onPrev", "delay": "0"}, _el("p:tgtEl", None, _el("p:sldTgt")))),
        _el("p:nextCondLst", None, _el("p:cond", {"evt": "onNext", "delay": "0"}, _el("p:tgtEl", None, _el("p:sldTgt")))),
    )
    root = _el(
        "p:par",
        None,
        _el(
            "p:cTn",
            {"id": root_ctn_id, "dur": "indefinite", "restart": "never", "nodeType": "tmRoot"},
            _el("p:childTnLst", None, main),
        ),
    )
    # Абзацы в анимации — список строится «по абзацам» (p:bldP build="p"), иначе PowerPoint покажет текст сразу.
    paragraph_shapes = sorted({target[0] for _, targets in steps for target in targets if not isinstance(target, int)})
    builds = (
        _el("p:bldLst", None, *(_el("p:bldP", {"spid": str(spid), "grpId": "0", "build": "p"}) for spid in paragraph_shapes))
        if paragraph_shapes
        else None
    )
    return _el("p:timing", None, _el("p:tnLst", None, root), builds)


def build_transition(spec: Slide) -> etree._Element | None:
    transition = spec.transition
    if transition is None or transition.type == "none" or transition.type not in TRANSITIONS:
        return None
    duration = transition.duration_ms
    speed = "fast" if duration <= 600 else "med" if duration <= 900 else "slow"
    element = _el("p:transition", {"spd": speed})
    effect = TRANSITIONS[transition.type]
    if effect == "push":
        element.append(_el("p:push", {"dir": "l"}))
    elif effect == "wipe":
        element.append(_el("p:wipe", {"dir": "r"}))
    elif effect == "zoom":
        element.append(_el("p:zoom"))
    else:
        element.append(_el("p:fade"))
    return element


def _place(slide_element: etree._Element, element: etree._Element, tag: str) -> None:
    """Вставляет p:transition / p:timing на своё место: после cSld и clrMapOvr, перед extLst."""
    for old in slide_element.findall(qn(tag)):
        slide_element.remove(old)
    order = ["p:cSld", "p:clrMapOvr", "p:transition", "p:timing"]
    anchor = None
    for name in order[: order.index(tag)]:
        found = slide_element.find(qn(name))
        if found is not None:
            anchor = found
    if anchor is not None:
        anchor.addnext(element)
    else:
        slide_element.insert(0, element)


def apply_motion(
    slide,
    spec: Slide,
    targets: dict[str, list[float]],
    *,
    width: int,
    height: int,
) -> dict[str, list[float]]:
    """Скрывает элементы, пишет анимации и переход. Возвращает рамки без скрытых элементов (для аудита)."""
    if not (spec.hidden_elements or spec.animations or spec.transition):
        return targets
    rects = target_rects(targets, spec)
    shapes = _top_shapes(slide, width, height)

    hidden = set(spec.hidden_elements or [])
    if hidden:
        removed: set[int] = set()
        for address in sorted(hidden, key=lambda item: not item.startswith("shape:")):
            for shape in resolve(address, rects, shapes, removed):
                parent = shape.element.getparent()
                if parent is not None:
                    parent.remove(shape.element)
                removed.add(shape.spid)
        shapes = [shape for shape in shapes if shape.spid not in removed]
        targets = {key: value for key, value in targets.items() if key not in hidden}

    steps: list[tuple[SlideAnimation, list[Target]]] = []
    taken: set[int] = set()
    # Фигуры слайда разбираются первыми: иначе карточка «заберёт» фигуру, у которой своя анимация.
    ordered = sorted(spec.animations or [], key=lambda item: not item.target.startswith("shape:"))
    claimed: dict[int, list[Target]] = {}
    for animation in ordered:
        if animation.target in hidden:
            continue
        paragraph = resolve_paragraph(animation.target, rects, shapes)
        if paragraph is not None:
            claimed[id(animation)] = [paragraph]
            continue
        found = resolve(animation.target, rects, shapes, taken)
        taken.update(shape.spid for shape in found)
        claimed[id(animation)] = [shape.spid for shape in found]
    for animation in spec.animations or []:
        spids = claimed.get(id(animation)) or []
        if spids:
            steps.append((animation, spids))

    try:
        timing = build_timing(steps)
        if timing is not None:
            _place(slide._element, timing, "p:timing")
        transition = build_transition(spec)
        if transition is not None:
            _place(slide._element, transition, "p:transition")
    except Exception:  # noqa: BLE001 — анимация не должна ломать сборку колоды
        logger.exception("Анимации слайда %s не записаны", spec.id)
    return targets
