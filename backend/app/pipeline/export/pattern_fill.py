"""Вёрстка по паттернам: клон слайда-примера шаблона + заполнение слотов.

Слайд-пример копируется со всеми фигурами, картинками и фоном, затем текст-заглушка
в слотах заменяется содержимым с сохранением форматирования шаблона (шрифт, кегль,
цвет, маркеры — из абзацев самого слота). Лишние карточки удаляются целиком, а текст
примера, которому не нашлось содержимого, убирается — в готовом слайде остаются только
сгенерированные тексты и оформление шаблона.

Диаграмма слайда-примера копируется вместе со своими частями (данные, стиль, цвета)
и получает новые данные через replace_data — оформление графика остаётся шаблонным.
Паттерн из макета (origin=layout) — новый слайд на макете, текст ставится в плейсхолдеры.
"""

import copy
import re
from dataclasses import dataclass, field
from typing import Any

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, XmlPart
from pptx.oxml.ns import qn

from app.pipeline.export.chart_style import highlight_points
from app.pipeline.parsing.layout_patterns import slot_id
from app.pipeline.parsing.patterns import _walk
from app.pipeline.visuals.numbers import numeric_columns
from app.schemas.slide import INVERSE_KINDS, ChartData, Slide, TableData

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
COPYABLE_RELTYPES = {
    RT.IMAGE,
    RT.MEDIA,
    RT.VIDEO,
    RT.AUDIO,
    RT.HYPERLINK,
    "http://schemas.microsoft.com/office/2007/relationships/hdphoto",
    "http://schemas.microsoft.com/office/2007/relationships/media",
}
LINK_TAGS = {qn("a:hlinkClick"), qn("a:hlinkHover")}
# Порядок дочерних элементов a:pPr по схеме: маркер ставим перед табуляцией и стилем текста.
PPR_TAIL = {qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst")}
BULLET_TAGS = {qn("a:buNone"), qn("a:buAutoNum"), qn("a:buChar"), qn("a:buBlip")}
TREE_HEADER = {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}
AUTOFIT_TAGS = {qn("a:noAutofit"), qn("a:normAutofit"), qn("a:spAutoFit")}
NUMBERING_RX = re.compile(r"^0?\d{1,2}[.)]?$")
# Адрес элемента для подсветки аудита: совпадает с target в проверках (bullets.2, columns.0, kpis.1…).
ITEM_TARGETS = {"two_columns": "columns", "cards": "columns", "kpi": "kpis", "timeline": "timeline", "agenda": "bullets"}


@dataclass
class ItemContent:
    title: str = ""
    lines: list[str] = field(default_factory=list)
    value: str = ""
    date: str = ""


@dataclass
class PatternContent:
    title: str
    subtitle: str | None = None
    body: list[str] = field(default_factory=list)
    items: list[ItemContent] = field(default_factory=list)


def content_for(slide: Slide) -> PatternContent:
    """Содержимое слайда в терминах слотов паттерна (общее для подбора и заполнения)."""
    content = PatternContent(title=slide.title, subtitle=slide.subtitle)
    kind = slide.kind
    if kind in {"two_columns", "cards"}:
        content.items = [ItemContent(title=column.title, lines=list(column.bullets)) for column in slide.columns]
    elif kind == "kpi":
        content.items = [ItemContent(value=kpi.value, lines=[kpi.label] if kpi.label else []) for kpi in slide.kpis]
    elif kind == "timeline":
        content.items = [
            ItemContent(title=item.title, date=item.date, lines=[item.text] if item.text else []) for item in slide.timeline
        ]
    elif kind == "agenda":
        content.items = [ItemContent(title=bullet) for bullet in slide.bullets]
        content.body = list(slide.bullets)
    elif kind == "quote" and slide.quote:
        content.body = [f"«{slide.quote.text}»"] + ([f"— {slide.quote.author}"] if slide.quote.author else [])
    else:
        content.body = list(slide.bullets)
    return content


def units(pattern: dict[str, Any], kind: str) -> list[dict[str, Any]]:
    """Повторяющиеся элементы паттерна. Для KPI к карточкам добавляются одиночные показатели с подписью."""
    result = [dict(item, single=False) for item in pattern["items"]]
    if kind != "kpi":
        return result
    slots = pattern["slots"]
    for index, slot in enumerate(slots):
        if slot["role"] != "value":
            continue
        unit_slots = [slot]
        label = slots[index + 1] if index + 1 < len(slots) and slots[index + 1]["role"] == "value_label" else None
        if label is not None:
            unit_slots.append(dict(label, role="text"))
        result.append(
            {
                "bbox": slot["bbox"],
                "group": None,
                "shapes": [s["id"] for s in unit_slots],
                "slots": unit_slots,
                "image": None,
                "single": True,
            }
        )
    return result


# --- клонирование ------------------------------------------------------------------


def _numbered(partname: str) -> str:
    """/ppt/charts/chart12.xml → /ppt/charts/chart%d.xml — шаблон имени для следующей свободной части."""
    return re.sub(r"\d+(\.\w+)$", r"%d\1", partname)


def _remap_rids(element, mapping: dict[str, str]) -> None:
    for node in element.iter():
        for attr, value in list(node.attrib.items()):
            if attr.startswith(R_NS) and value in mapping:
                node.set(attr, mapping[value])


def copy_part(part, package):
    """Копия части пакета (диаграмма, её книга Excel, стиль и цвета) со своими связями: у каждого
    слайда должна быть своя диаграмма, иначе новые данные попадут и в соседний слайд."""
    partname = package.next_partname(_numbered(str(part.partname)))
    # Тот же класс части (EmbeddedXlsxPart, ChartPart…): replace_data пишет новую книгу Excel через его API.
    loader = type(part) if isinstance(part, (Part, XmlPart)) else Part
    new = loader.load(partname, part.content_type, package, part.blob)
    mapping: dict[str, str] = {}
    for rel in part.rels.values():
        if rel.is_external:
            mapping[rel.rId] = new.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.IMAGE:
            mapping[rel.rId] = new.relate_to(rel.target_part, rel.reltype)
        else:
            mapping[rel.rId] = new.relate_to(copy_part(rel.target_part, package), rel.reltype)
    if isinstance(new, XmlPart) and mapping:
        _remap_rids(new._element, mapping)
    return new


def clone_slide(prs, source):
    """Копия слайда-примера в конце презентации: фигуры, картинки, диаграммы, ссылки и фон."""
    slide = prs.slides.add_slide(source.slide_layout)
    tree = slide.shapes._spTree
    for element in list(tree):
        if element.tag not in TREE_HEADER:
            tree.remove(element)

    rid_map: dict[str, str] = {}
    for rel in source.part.rels.values():
        if rel.reltype in COPYABLE_RELTYPES:
            target = rel.target_ref if rel.is_external else rel.target_part
            rid_map[rel.rId] = slide.part.relate_to(target, rel.reltype, is_external=rel.is_external)
        elif rel.reltype == RT.CHART and not rel.is_external:
            try:
                rid_map[rel.rId] = slide.part.relate_to(copy_part(rel.target_part, slide.part.package), RT.CHART)
            except Exception:  # noqa: BLE001 — не скопировалась диаграмма: фигура просто не попадёт в клон
                continue

    def remap(element) -> bool:
        for node in list(element.iter()):
            for attr, value in list(node.attrib.items()):
                if not attr.startswith(R_NS):
                    continue
                if value in rid_map:
                    node.set(attr, rid_map[value])
                elif node.tag in LINK_TAGS:
                    # Переход на другой слайд шаблона — в новой колоде его нет.
                    node.getparent().remove(node)
                    break
                else:
                    return False  # диаграммы, OLE, SmartArt: без своих частей не копируются
        return True

    for element in source.shapes._spTree.iterchildren():
        if element.tag in TREE_HEADER:
            continue
        clone = copy.deepcopy(element)
        if remap(clone):
            tree.append(clone)

    c_sld = slide._element.find(qn("p:cSld"))
    background = source._element.find(f"{qn('p:cSld')}/{qn('p:bg')}")
    if background is not None:
        clone = copy.deepcopy(background)
        if remap(clone):
            c_sld.insert(0, clone)
    color_map = source._element.find(qn("p:clrMapOvr"))
    if color_map is not None:
        current = slide._element.find(qn("p:clrMapOvr"))
        if current is not None:
            slide._element.remove(current)
        c_sld.addnext(copy.deepcopy(color_map))
    return slide


# --- работа с текстом ---------------------------------------------------------------


def shape_elements(slide) -> dict[int, Any]:
    """id фигуры → её элемент (sp, pic, grpSp, graphicFrame, cxnSp)."""
    result = {}
    for c_nv_pr in slide._element.iter(qn("p:cNvPr")):
        result[int(c_nv_pr.get("id"))] = c_nv_pr.getparent().getparent()
    return result


def placeholder_elements(slide) -> dict[int, Any]:
    """Для слайда на макете: id слота (PH_BASE + idx плейсхолдера) → элемент плейсхолдера."""
    return {slot_id(ph.placeholder_format.idx): ph._element for ph in slide.placeholders}


def prune_placeholders(slide, keep: set[int] | None = None) -> None:
    """Пустые плейсхолдеры в PowerPoint показывают подсказки «Введите текст» — убираем их."""
    keep = keep or set()
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx in keep:
            continue
        text = ph.text_frame.text.strip() if ph.has_text_frame else ""
        if not text and ph._element.find(f".//{qn('p:blipFill')}") is None:
            remove(ph._element)


def _paragraph_text(paragraph) -> str:
    return "".join(node.text or "" for node in paragraph.iter(qn("a:t"))).strip()


def _filled_paragraph(template, text: str):
    paragraph = copy.deepcopy(template)
    runs = paragraph.findall(qn("a:r"))
    for child in list(paragraph):
        if child.tag in {qn("a:r"), qn("a:br"), qn("a:fld")} and (not runs or child is not runs[0]):
            paragraph.remove(child)
    if runs:
        runs[0].find(qn("a:t")).text = text
        return paragraph
    run = etree.Element(qn("a:r"))
    end = paragraph.find(qn("a:endParaRPr"))
    if end is not None:
        props = copy.deepcopy(end)
        props.tag = qn("a:rPr")
        run.append(props)
    etree.SubElement(run, qn("a:t")).text = text
    if end is not None:
        end.addprevious(run)
    else:
        paragraph.append(run)
    return paragraph



# Кегль слота подгоняется под рамку: модель та же, что у парсера (ёмкость ~ площадь / кегль²),
# поэтому уменьшение кегля в k раз даёт в k² больше знаков.
EMU_PER_PT = 12700
CHAR_WIDTH = 0.58  # средняя ширина знака в долях кегля для обычного текста
CHAR_WIDTH_SHORT = 0.8  # у крупных чисел («40%») знаки заметно шире средних букв
TEXT_INSET_PT = 7.2  # стандартные поля текстовой рамки OOXML с двух сторон
SHORT_TEXT = 6  # столько знаков — это ещё «число», а не фраза
LINE_HEIGHT = 1.2  # интерлиньяж при переносе
MIN_FONT_SCALE = 0.35  # ниже текст нечитаем — пусть лучше аудит сообщит о переполнении


@dataclass(frozen=True)
class Fit:
    """Размер слайда в пунктах: по нему считается, во сколько раз ужать текст в слоте.
    steps — ступени типографической шкалы шаблона: ужатый кегль встаёт на ближайшую меньшую."""

    width_pt: float
    height_pt: float
    steps: tuple[float, ...] = ()

    def scale(self, slot: dict[str, Any] | None, text: str) -> float:
        """Во сколько раз уменьшить кегль слота, чтобы текст поместился в его рамку.

        Два ограничения. По ширине: самое длинное слово перенести некуда, оно должно влезть
        в строку (для «40%» в рамке под одну цифру это главное условие). По высоте: весь текст
        с переносами должен уложиться в рамку — уменьшение кегля в k раз и укорачивает строки,
        и увеличивает их число, поэтому для многострочного текста масштаб берётся под корнем."""
        text = " ".join((text or "").split())
        if not slot or not text:
            return 1.0
        size = slot.get("size") or 18
        bbox = slot.get("bbox") or []
        if len(bbox) != 4 or size <= 0:
            return 1.0
        # Текст живёт не во всей фигуре: по краям рамки стандартные поля.
        box_width = bbox[2] * self.width_pt - 2 * TEXT_INSET_PT
        box_height = bbox[3] * self.height_pt
        char_width = (CHAR_WIDTH_SHORT if len(text) <= SHORT_TEXT else CHAR_WIDTH) * size
        if box_width <= 0 or box_height <= 0 or char_width <= 0:
            return 1.0

        # Крупные числа шаблон часто ставит с очень плотным интерлиньяжем (16% кегля):
        # перенос такой строки — это наложение цифр друг на друга, поэтому короткий текст
        # обязан поместиться в одну строку целиком, а не только самым длинным словом.
        longest = len(text) if len(text) <= SHORT_TEXT else max(len(word) for word in text.split())
        width_scale = box_width / (longest * char_width)
        per_line = box_width / char_width
        if len(text) <= per_line:
            height_scale = box_height / size  # одна строка: её высота примерно равна кеглю
        else:
            height_scale = (box_width * box_height / (len(text) * char_width * size * LINE_HEIGHT)) ** 0.5
        return max(MIN_FONT_SCALE, min(1.0, width_scale, height_scale))


def slide_fit(slide, steps: list[float] | None = None) -> Fit | None:
    """Размер слайда из презентации, которой он принадлежит."""
    try:
        presentation = slide.part.package.presentation_part.presentation
        return Fit(
            presentation.slide_width / EMU_PER_PT,
            presentation.slide_height / EMU_PER_PT,
            tuple(sorted(steps or ())),
        )
    except Exception:  # noqa: BLE001 — без размеров слайда просто не подгоняем кегль
        return None


def _autofit(body) -> None:
    """Разрешает PowerPoint пересчитывать кегль при правках текста (сам по себе ничего не меняет:
    уменьшение кегля под рамку делает _scale_runs, потому что fontScale LibreOffice игнорирует)."""
    body_pr = body.find(qn("a:bodyPr"))
    if body_pr is None:
        return
    for child in list(body_pr):
        if child.tag in AUTOFIT_TAGS:
            body_pr.remove(child)
    autofit = etree.Element(qn("a:normAutofit"))
    warp = body_pr.find(qn("a:prstTxWarp"))
    if warp is not None:
        warp.addnext(autofit)
    else:
        body_pr.insert(0, autofit)


MIN_FONT_SZ = 600  # 6 pt в сотых долях пункта — мельче текст бессмысленно уменьшать
STEP_DROP = 0.8  # до ступени шкалы кегль уменьшается не больше чем ещё на пятую часть


def snap_size(size: int, steps: tuple[float, ...]) -> int:
    """Ужатый кегль (в сотых пункта) — на ближайшую меньшую ступень шкалы шаблона, если она рядом."""
    lower = [int(step * 100) for step in steps if step * 100 <= size]
    if lower and lower[-1] >= size * STEP_DROP:
        return lower[-1]
    return size


def _scale_runs(body, scale: float, base_size: float | None, steps: tuple[float, ...] = ()) -> None:
    """Уменьшает кегль текста в рамке. LibreOffice не применяет fontScale из normAutofit,
    поэтому размер проставляется каждому прогону явно — так работает в любом редакторе."""
    for props in body.iter(qn("a:rPr"), qn("a:endParaRPr"), qn("a:defRPr")):
        size = props.get("sz")
        if size and size.isdigit():
            props.set("sz", str(max(snap_size(int(int(size) * scale), steps), MIN_FONT_SZ)))
        elif base_size:
            props.set("sz", str(max(snap_size(int(base_size * 100 * scale), steps), MIN_FONT_SZ)))


def _plain_head(paragraph) -> None:
    """Заголовок элемента внутри плейсхолдера списка: без маркера и отступа, жирным."""
    props = paragraph.find(qn("a:pPr"))
    if props is None:
        props = etree.Element(qn("a:pPr"))
        paragraph.insert(0, props)
    props.set("marL", "0")
    props.set("indent", "0")
    for child in list(props):
        if child.tag in BULLET_TAGS:
            props.remove(child)
    tail = next((child for child in props if child.tag in PPR_TAIL), None)
    bullet = etree.Element(qn("a:buNone"))
    if tail is not None:
        tail.addprevious(bullet)
    else:
        props.append(bullet)
    for run in paragraph.findall(qn("a:r")):
        run_props = run.find(qn("a:rPr"))
        if run_props is None:
            run_props = etree.Element(qn("a:rPr"))
            run.insert(0, run_props)
        run_props.set("b", "1")


def set_body(
    body,
    lines: list[str],
    *,
    head: str | None = None,
    autofit: bool = True,
    plain_head: bool = False,
    scale: float = 1.0,
    base_size: float | None = None,
    steps: tuple[float, ...] = (),
    align: str | None = None,
) -> bool:
    """Заменяет абзацы txBody. head — первый абзац в стиле первого абзаца шаблона (заголовок карточки)."""
    if body is None:
        return False
    paragraphs = body.findall(qn("a:p"))
    if not paragraphs:
        return False
    # Стиль — из абзацев самого слота: первый для первой строки (или заголовка карточки),
    # второй для следующих — у них бывают свои отступы и интервалы.
    templates = [p for p in paragraphs if _paragraph_text(p)] or paragraphs
    for paragraph in paragraphs:
        body.remove(paragraph)
    if head is not None:
        head_paragraph = _filled_paragraph(templates[0], head)
        if plain_head:
            _plain_head(head_paragraph)
        body.append(head_paragraph)
        templates = templates[1:] or templates
    for index, line in enumerate(lines if lines or head is not None else [""]):
        body.append(_filled_paragraph(templates[min(index, len(templates) - 1, 1)], line))
    if scale < 1:
        _scale_runs(body, scale, base_size, steps)
    if align:
        for paragraph in body.findall(qn("a:p")):
            props = paragraph.find(qn("a:pPr"))
            if props is None:
                props = etree.Element(qn("a:pPr"))
                paragraph.insert(0, props)
            props.set("algn", align)
    if autofit:
        _autofit(body)
    return True


def set_text(
    element,
    lines: list[str],
    *,
    head: str | None = None,
    plain_head: bool = False,
    slot: dict[str, Any] | None = None,
    fit: Fit | None = None,
    scale: float | None = None,
) -> bool:
    """Ставит текст в фигуру. slot и fit — чтобы длинный текст ужался по рамке слота;
    scale — готовый масштаб, когда кегль согласован по всей серии элементов."""
    text = " ".join([head or "", *lines]).strip()
    if scale is None:
        scale = fit.scale(slot, text) if fit else 1.0
    return set_body(
        element.find(qn("p:txBody")),
        lines,
        head=head,
        plain_head=plain_head,
        scale=scale,
        base_size=(slot or {}).get("size"),
        steps=fit.steps if fit else (),
    )


def remove(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def drop(elements: dict[int, Any], slot: dict[str, Any]) -> None:
    """Удаляет фигуру слота вместе с оформлением, которое без неё теряет смысл: маркер строки,
    иконку, плашку-подложку (companions из парсера)."""
    for shape_id in [*slot.get("companions", []), slot["id"]]:
        if (element := elements.get(shape_id)) is not None:
            remove(element)


# --- заполнение паттерна ------------------------------------------------------------


def _fill_item(
    elements: dict[int, Any],
    unit: dict[str, Any],
    content: ItemContent,
    number: int,
    plain_head: bool = False,
    fit: Fit | None = None,
    value_scale: float | None = None,
) -> None:
    slots: dict[str, list[tuple[dict[str, Any], Any]]] = {}
    for slot in unit["slots"]:
        if (element := elements.get(slot["id"])) is not None:
            slots.setdefault(slot["role"], []).append((slot, element))
    title, lines, date = content.title, list(content.lines), content.date

    for index, (slot, element) in enumerate(slots.get("value", [])):
        if index:
            drop(elements, slot)
            continue
        set_text(element, [content.value or date or f"{number:02d}"], slot=slot, fit=fit, scale=value_scale)
        date = "" if not content.value else date
    for slot, element in slots.get("date", []):
        set_text(element, [date or f"{number:02d}"], slot=slot, fit=fit)
        date = ""
    for slot, element in slots.get("number", []):
        # В кружок номера помещается «01» или «Q4», но не «Шаг 1» — такая дата уходит в заголовок.
        fits = date and len(date) <= min(slot["chars"], 3)
        set_text(element, [date if fits else f"{number:02d}"], slot=slot, fit=fit)
        date = "" if fits else date
    if date:
        # Даты некуда поставить отдельно — она становится частью заголовка шага.
        title = f"{date} · {title}" if title else date

    heads = slots.get("title_text", []) + slots.get("title", [])
    texts = sorted(slots.get("text", []), key=lambda pair: (-(pair[0]["size"] or 0), pair[0]["bbox"][1]))
    line_slots = sorted(slots.get("lines", []), key=lambda pair: pair[0]["bbox"][1])
    if not title:
        # Без заголовка (подпись KPI) слоты заголовка работают как текстовые.
        heads, texts = [], heads + texts
    elif not heads and texts:
        # Нет отдельного слота заголовка: берём самый крупный текстовый.
        heads, texts = [texts[0]], texts[1:]
    text_capable = bool(texts or line_slots)

    for index, (slot, element) in enumerate(heads):
        if index:
            drop(elements, slot)
        elif slot["role"] == "title_text" or not text_capable:
            set_text(element, lines, head=title, plain_head=plain_head, slot=slot, fit=fit)
            lines = []
        else:
            set_text(element, [title], slot=slot, fit=fit)
    for index, (slot, element) in enumerate(texts):
        if index == 0 and lines:
            set_text(element, lines, slot=slot, fit=fit)
            lines = []
        else:
            drop(elements, slot)
    for index, (slot, element) in enumerate(line_slots):
        if not lines:
            drop(elements, slot)
        elif index == len(line_slots) - 1:
            set_text(element, lines, slot=slot, fit=fit)
            lines = []
        else:
            set_text(element, [lines.pop(0)], slot=slot, fit=fit)


def fill_table(frame, table: TableData) -> None:
    """Таблица шаблона под данные слайда: строки и столбцы клонируются из шаблонных, стиль сохраняется."""
    tbl = frame.find(f".//{qn('a:tbl')}")
    if tbl is None:
        return
    rows_data = [list(table.columns)] + [list(row) for row in table.rows]
    n_cols = max((len(row) for row in rows_data), default=1) or 1
    grid = tbl.find(qn("a:tblGrid"))
    if grid is None:
        return
    columns = grid.findall(qn("a:gridCol"))
    if not columns:
        return
    total_width = sum(int(column.get("w", "0")) for column in columns)
    for column in columns[n_cols:]:
        grid.remove(column)
    while len(grid.findall(qn("a:gridCol"))) < n_cols:
        grid.append(copy.deepcopy(columns[-1]))
    for column in grid.findall(qn("a:gridCol")):
        column.set("w", str(total_width // n_cols))

    templates = tbl.findall(qn("a:tr"))
    if not templates:
        return
    numeric = numeric_columns(table)
    frame_height = sum(int(row.get("h", "0")) for row in templates)
    row_height = max(frame_height // max(len(rows_data), 1), 1)
    # Шапка — первая строка шаблона, данные — одна «типовая» строка из середины: крайние строки
    # примера часто выделены (итог, акцент), а чередование полос задаёт стиль таблицы.
    body_rows = templates[1:] or templates
    body_template = body_rows[len(body_rows) // 2]
    for row in templates:
        tbl.remove(row)
    for index, values in enumerate(rows_data):
        template = templates[0] if index == 0 else body_template
        row = copy.deepcopy(template)
        row.set("h", str(min(int(template.get("h", "0")) or row_height, row_height)))
        cells = row.findall(qn("a:tc"))
        if not cells:
            continue
        for cell in cells[n_cols:]:
            row.remove(cell)
        while len(row.findall(qn("a:tc"))) < n_cols:
            row.append(copy.deepcopy(cells[-1]))
        for column, (cell, text) in enumerate(zip(row.findall(qn("a:tc")), values + [""] * n_cols)):
            for attr in ("gridSpan", "rowSpan", "hMerge", "vMerge"):
                cell.attrib.pop(attr, None)
            # Столбец чисел — по правому краю вместе с шапкой: так сравнивают разряды.
            set_body(cell.find(qn("a:txBody")), [text], autofit=False, align="r" if column in numeric else None)
        tbl.append(row)

    extent = frame.find(f"{qn('p:xfrm')}/{qn('a:ext')}")
    if extent is not None:
        extent.set("cy", str(sum(int(row.get("h", "0")) for row in tbl.findall(qn("a:tr")))))


def _union(boxes: list[list[float]]) -> list[float]:
    x0 = min(b[0] for b in boxes)
    y0 = min(b[1] for b in boxes)
    x1 = max(b[0] + b[2] for b in boxes)
    y1 = max(b[1] + b[3] for b in boxes)
    return [round(x0, 4), round(y0, 4), round(x1 - x0, 4), round(y1 - y0, 4)]


def fill_chart(
    slide, frame_id: int, chart_spec: ChartData, accents: list[str] | None = None, background: str = "#FFFFFF"
) -> bool:
    """Данные слайда в диаграмму слайда-примера: цвета, шрифты и оси остаются шаблонными."""
    shape = next(
        (s for s in _walk(slide.shapes) if s.shape_id == frame_id and getattr(s, "has_chart", False) and s.has_chart),
        None,
    )
    if shape is None or not chart_spec.series:
        return False
    chart = shape.chart
    before = sum(len(list(plot.series)) for plot in chart.plots)
    data = CategoryChartData()
    data.categories = chart_spec.categories
    for series in chart_spec.series:
        data.add_series(series.name or "Серия", series.values)
    chart.replace_data(data)
    if chart.has_title:
        if chart_spec.caption:
            chart.chart_title.text_frame.text = chart_spec.caption
        else:
            chart.has_title = False
    # Лишние ряды python-pptx клонирует из последнего — даём им свои цвета, чтобы не слились.
    if accents and chart.plots:
        for index, series in enumerate(chart.plots[0].series):
            if index >= before:
                color = accents[index % len(accents)].lstrip("#")
                try:
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = RGBColor.from_string(color)
                except Exception:  # noqa: BLE001 — у линий нет заливки, красим линию
                    series.format.line.color.rgb = RGBColor.from_string(color)
    if chart.plots:
        highlight_points(chart.plots[0], chart_spec, background, fallback=(accents or [None])[0])
    return True


def series_scale(unit_list: list[dict[str, Any]], content: PatternContent, fit: Fit | None) -> float | None:
    """Крупные числа серии («8 ч», «40%», «×96») показываются одним кеглем: берём самый мелкий
    из нужных, иначе соседние KPI отличаются размером в полтора раза."""
    if fit is None:
        return None
    scales = [
        fit.scale(slot, item.value)
        for unit, item in zip(unit_list, content.items)
        if item.value
        for slot in unit["slots"]
        if slot["role"] == "value"
    ]
    return min(scales) if scales else None


def fill_pattern(
    slide,
    pattern: dict[str, Any],
    spec: Slide,
    *,
    accents: list[str] | None = None,
    background: str = "#FFFFFF",
    steps: list[float] | None = None,
) -> dict[str, list[float]]:
    """Заполняет клон слайда-примера (или новый слайд на макете для origin=layout).
    Возвращает рамки элементов для подсветки аудита (адрес → [x, y, w, h])."""
    from_layout = pattern.get("origin") == "layout"
    # Размеры слайда нужны, чтобы ужать кегль там, где текст длиннее образца из шаблона.
    fit = slide_fit(slide, steps)
    elements = placeholder_elements(slide) if from_layout else shape_elements(slide)
    plain_head = bool(pattern.get("plain_heads"))
    content = content_for(spec)
    targets: dict[str, list[float]] = {}
    handled: set[int] = set()

    title = pattern.get("title")
    if title and (element := elements.get(title["id"])) is not None:
        set_text(element, [content.title], slot=title, fit=fit)
        targets["title"] = title["bbox"]
        handled.add(title["id"])

    prefix = ITEM_TARGETS.get(spec.kind)
    kept: list[list[float]] = []
    unit_list = list(units(pattern, spec.kind))
    value_scale = series_scale(unit_list, content, fit)
    for index, unit in enumerate(unit_list):
        handled.update(slot["id"] for slot in unit["slots"])
        if index < len(content.items):
            _fill_item(
                elements, unit, content.items[index], index + 1, plain_head=plain_head, fit=fit, value_scale=value_scale
            )
            kept.append(unit["bbox"])
            if prefix:
                targets[f"{prefix}.{index}"] = unit["bbox"]
        elif unit.get("group") and (group := elements.get(unit["group"])) is not None:
            remove(group)
        else:
            for shape_id in unit["shapes"]:
                if (element := elements.get(shape_id)) is not None:
                    remove(element)
    if prefix and kept:
        targets[prefix] = _union(kept)

    table = pattern["frames"].get("table")
    if table and spec.table and (frame := elements.get(table["id"])) is not None:
        fill_table(frame, spec.table)
        targets["table"] = table["bbox"]
    chart = pattern["frames"].get("chart")
    if chart and spec.chart and fill_chart(slide, chart["id"], spec.chart, accents, background):
        targets["chart"] = chart["bbox"]

    # Обложка, разделитель, финал: подзаголовок — в свой слот или в текстовый, номер раздела — в «01».
    service = spec.kind in INVERSE_KINDS
    has_subtitle_slot = any(slot["role"] == "subtitle" for slot in pattern["slots"])
    number_slot = None
    if spec.kind == "section" and content.subtitle and re.fullmatch(r"\d{1,2}", content.subtitle.strip()):
        number_slot = next(
            (slot for slot in pattern["slots"] if slot["role"] == "decor_text" and NUMBERING_RX.match(slot["sample"])),
            None,
        )
    filled: set[str] = set()
    use_body = bool(content.body) and not kept
    for slot in pattern["slots"]:
        element = elements.get(slot["id"])
        if element is None or slot["id"] in handled:
            continue
        role = slot["role"]
        if slot is number_slot:
            set_text(element, [content.subtitle.strip()], slot=slot, fit=fit)
            filled.add("subtitle")
        elif (
            service
            and role in {"subtitle", "body"}
            and (role == "subtitle" or not has_subtitle_slot)
            and content.subtitle
            and number_slot is None
            and "subtitle" not in filled
        ):
            set_text(element, [content.subtitle], slot=slot, fit=fit)
            targets["subtitle"] = slot["bbox"]
            filled.add("subtitle")
        elif role == "body" and use_body and "body" not in filled:
            set_text(element, content.body, slot=slot, fit=fit)
            targets["quote" if spec.kind == "quote" else "bullets"] = slot["bbox"]
            filled.add("body")
        elif role == "subtitle" and content.subtitle and "subtitle" not in filled:
            set_text(element, [content.subtitle], slot=slot, fit=fit)
            targets["subtitle"] = slot["bbox"]
            filled.add("subtitle")
        elif role == "decor_text" and NUMBERING_RX.match(slot["sample"]):
            continue  # номера пунктов — часть оформления
        else:
            # Текст примера без содержимого (заглушки, подписи, места под картинки) — вместе с плашкой.
            drop(elements, slot)

    if pattern["images"]:
        targets["image"] = pattern["images"][0]["bbox"]
    if from_layout:
        # Плейсхолдер картинки оставляем — в него вставит изображение сборка (или уберёт его сама).
        keep = {image["id"] - slot_id(0) for image in pattern["images"] if image.get("placeholder")}
        prune_placeholders(slide, keep=keep)
    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes
    return targets
