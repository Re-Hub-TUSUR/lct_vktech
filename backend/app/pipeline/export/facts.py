"""Факты собранного слайда для аудита: что на нём реально стоит после сборки .pptx.

Собираются при рендере (services/renders.py) вместе с разметкой строк текста из PDF и хранятся
у рендера варианта по слайдам. Проверки аудита по рендеру (audit/geometry.py) читают их, а не
пересобирают колоду. Здесь же — «инвентарь» файла шаблона: все цвета, гарнитуры и кегли, которые
в нём встречаются. Сравнение с ним отличает то, что внесла генерация, от собственного оформления
шаблона, которое клоны слайдов-примеров переносят как есть.
"""

import hashlib
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

PART_PREFIXES = ("ppt/slides/", "ppt/slideLayouts/", "ppt/slideMasters/", "ppt/theme/", "ppt/charts/")
SIZE_TAGS = ("a:rPr", "a:endParaRPr", "a:defRPr")
CHROME_TYPES = {PP_PLACEHOLDER.SLIDE_NUMBER: "slide_number", PP_PLACEHOLDER.FOOTER: "footer", PP_PLACEHOLDER.DATE: "date"}
LOGO_MAX_AREA = 0.05  # крупная картинка мастера — фон или иллюстрация, а не логотип
_PARSER = etree.XMLParser(resolve_entities=False, huge_tree=True)


def digest(blob: bytes) -> str:
    return hashlib.sha1(blob).hexdigest()[:12]


def _theme_face(face: str, theme_fonts: dict[str, Any]) -> str:
    if face.startswith("+mj"):
        return theme_fonts.get("major") or ""
    if face.startswith("+mn"):
        return theme_fonts.get("minor") or ""
    return face


@lru_cache(maxsize=16)
def _inventory(path: str, mtime: float) -> dict[str, list]:
    colors: set[str] = set()
    fonts: set[str] = set()
    sizes: set[float] = set()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml") or not name.startswith(PART_PREFIXES):
                continue
            root = etree.fromstring(archive.read(name), _PARSER)
            for element in root.iter(qn("a:srgbClr")):
                if value := element.get("val"):
                    colors.add("#" + value.upper())
            for element in root.iter(qn("a:latin")):
                face = element.get("typeface") or ""
                if face and not face.startswith("+"):
                    fonts.add(face)
            for tag in SIZE_TAGS:
                for element in root.iter(qn(tag)):
                    size = element.get("sz")
                    if size and size.isdigit():
                        sizes.add(int(size) / 100)
    return {"colors": sorted(colors), "fonts": sorted(fonts), "sizes": sorted(sizes)}


def template_inventory(path: str | Path | None) -> dict[str, list]:
    """Все цвета (srgb), гарнитуры и кегли файла шаблона. Кэшируется по пути и времени изменения."""
    if not path or not Path(path).exists():
        return {"colors": [], "fonts": [], "sizes": []}
    return _inventory(str(path), Path(path).stat().st_mtime)


def _walk(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def _box(shape, width: int, height: int) -> list[float] | None:
    if None in (shape.left, shape.top, shape.width, shape.height) or not width or not height:
        return None
    return [
        round(shape.left / width, 4),
        round(shape.top / height, 4),
        round(shape.width / width, 4),
        round(shape.height / height, 4),
    ]


def _picture(shape, width: int, height: int, user_digests: set[str]) -> dict[str, Any] | None:
    box = _box(shape, width, height)
    try:
        blob = shape.image.blob
        pixels_w, pixels_h = shape.image.size
    except Exception:  # noqa: BLE001 — картинка-ссылка или битый медиафайл: пропорции не проверить
        return None
    if not box or not pixels_w or not pixels_h or not shape.width or not shape.height:
        return None
    visible_w = pixels_w * max(1 - shape.crop_left - shape.crop_right, 1e-3)
    visible_h = pixels_h * max(1 - shape.crop_top - shape.crop_bottom, 1e-3)
    # Во сколько раз рамка вытянула картинку по горизонтали относительно её собственных пропорций.
    stretch = (shape.width / shape.height) / (visible_w / visible_h)
    key = digest(blob)
    return {"bbox": box, "stretch": round(stretch, 3), "image": key, "user": key in user_digests}


def _masters_visible(layout) -> bool:
    return layout._element.get("showMasterSp") != "0"


def chrome_boxes(slide, width: int, height: int, logo_digests: set[str]) -> list[dict[str, Any]]:
    """Где на слайде логотипы и колонтитулы: картинки макета и мастера, плейсхолдеры номера и колонтитула."""
    boxes: list[dict[str, Any]] = []
    layout = slide.slide_layout
    owners = [layout, layout.slide_master] if _masters_visible(layout) else [layout]
    for owner in owners:
        for shape in owner.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or not (box := _box(shape, width, height)):
                continue
            if box[2] * box[3] <= LOGO_MAX_AREA:
                boxes.append({"kind": "logo", "bbox": box})
    for shape in _walk(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and logo_digests:
            try:
                is_logo = digest(shape.image.blob) in logo_digests
            except Exception:  # noqa: BLE001
                is_logo = False
            if is_logo and (box := _box(shape, width, height)):
                boxes.append({"kind": "logo", "bbox": box})
    for placeholder in slide.placeholders:
        kind = CHROME_TYPES.get(placeholder.placeholder_format.type)
        if kind and (box := _box(placeholder, width, height)):
            boxes.append({"kind": kind, "bbox": box})
    return boxes


def slide_facts(
    slide,
    *,
    width: int,
    height: int,
    theme_fonts: dict[str, Any],
    user_digests: set[str],
    logo_digests: set[str],
) -> dict[str, Any]:
    """Шрифты, цвета и кегли, заданные в самом слайде; картинки с пропорциями; логотипы и колонтитулы."""
    root = slide._element
    # Шрифт и кегль считаются только у прогонов с текстом: пустые абзацы шаблона на слайде не видны.
    props = [
        run.find(qn("a:rPr"))
        for run in (*root.iter(qn("a:r")), *root.iter(qn("a:fld")))
        if (run.findtext(qn("a:t")) or "").strip()
    ]
    props = [prop for prop in props if prop is not None]
    fonts = {
        _theme_face(latin.get("typeface") or "", theme_fonts)
        for prop in props
        if (latin := prop.find(qn("a:latin"))) is not None
    }
    colors = {"#" + element.get("val").upper() for element in root.iter(qn("a:srgbClr")) if element.get("val")}
    sizes = {int(prop.get("sz")) / 100 for prop in props if (prop.get("sz") or "").isdigit()}
    shapes = list(slide.shapes)
    # Картинка в плейсхолдере (место под фото макета или слайда-примера) — тоже картинка слайда.
    pictures = [
        picture
        for shape in _walk(shapes)
        if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.PLACEHOLDER)
        and (picture := _picture(shape, width, height, user_digests)) is not None
    ]
    return {
        "fonts": sorted(face for face in fonts if face),
        "colors": sorted(colors),
        "sizes": sorted(sizes),
        "pictures": pictures,
        "chrome": chrome_boxes(slide, width, height, logo_digests),
        # integrity.raster_slide: на слайде есть объекты, и все они — картинки.
        "only_pictures": bool(shapes) and all(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes),
    }
