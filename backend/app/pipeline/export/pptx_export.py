"""Экспорт в .pptx нативными объектами на основе исходного шаблона.

Берём сам файл шаблона (мастера, макеты, тема, шрифты) и собираем новые слайды.
Слайд с паттерном — клон слайда-примера шаблона с заполненными слотами
(pattern_fill.py): карточки, KPI, таймлайны и таблицы выглядят как в шаблоне.
Остальные — на макете шаблона: заголовки и тезисы в плейсхолдеры или текстовые
поля, данные в нативные диаграммы, таблицы и фигуры. В конце слайды-примеры
удаляются. Растровых слайдов нет.

Токены дизайн-системы шаблона (theme["design"], см. parsing/design_tokens.py): новые
таблицы берут стиль таблиц шаблона, диаграммы — цвета его рядов, текст на макетах —
кегли из типографики; номер слайда и колонтитул переносятся, если их ставит шаблон.
"""

import copy
import logging
import math
import re
from dataclasses import dataclass, field, replace
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image
import pymupdf
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.core.storage import storage_path
from app.pipeline.export.chart_style import chart_caption, highlight_points, number_format
from app.pipeline.export.facts import digest, slide_facts
from app.pipeline.export.motion import apply_motion
from app.pipeline.export.text_style import style_for, style_paragraph
from app.pipeline.export.pattern_fill import clone_slide, fill_pattern, remove, set_body
from app.pipeline.geometry import content_frame
from app.pipeline.parsing.layout_patterns import PH_BASE
from app.pipeline.parsing.template_parser import extract_theme, layout_background
from app.pipeline.theme import mix, text_on
from app.pipeline.visuals.numbers import numeric_columns
from app.schemas.slide import Slide

logger = logging.getLogger(__name__)

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE}
REFERENCE_WIDTH = 12192000  # 13,333" — ширина, под которую подобраны кегли
# Кегль ближе этой доли к ступени типографической шкалы шаблона встаёт на ступень.
SCALE_SNAP = 0.15
CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "horizontal_bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "donut": XL_CHART_TYPE.DOUGHNUT,
}


# Рамка элемента в долях слайда: [x, y, w, h]; адреса — как target в проверках аудита.
Targets = dict[str, list[float]]
# Где на слайде стоит содержимое каждого типа (для подсветки аудита на макетном пути).
CONTENT_TARGET = {
    "bullets": "bullets",
    "agenda": "bullets",
    "two_columns": "columns",
    "cards": "columns",
    "chart": "chart",
    "table": "table",
    "kpi": "kpis",
    "timeline": "timeline",
    "quote": "quote",
    "image_text": "bullets",
}


class ExportError(RuntimeError):
    pass


@dataclass
class DeckBuild:
    data: bytes
    # Рамки элементов каждого слайда — для подсветки проблем аудита на рендере.
    targets: list[Targets] = field(default_factory=list)
    # Что реально стоит на слайдах (export/facts.py) — для проверок аудита по рендеру.
    facts: list[dict] = field(default_factory=list)


@dataclass
class _Ctx:
    width: int
    height: int
    theme: dict
    # Цвета текущего слайда: зависят от фона его макета (в тёмных шаблонах текст светлый).
    fg: str = "#000000"
    muted: str = "#737373"
    surface: str = "#E6F1FF"
    background: str = "#FFFFFF"
    font_family: str | None = None
    # Гарнитуры шаблона по умолчанию — те же, что у превью (theme.heading_font / body_font): без них новый
    # текст наследовал бы шрифт темы, а он в шаблонах часто формальный Arial вместо фирменной гарнитуры.
    heading_font: str | None = None
    body_font: str | None = None

    @property
    def scale(self) -> float:
        return self.width / REFERENCE_WIDTH

    @property
    def steps(self) -> list[float]:
        """Ступени типографической шкалы шаблона (частые кегли его слайдов), от 8 pt."""
        return [item["size_pt"] for item in self.design.get("type_scale") or [] if (item.get("size_pt") or 0) >= 8]

    def pt(self, size: float) -> Pt:
        """Кегль в масштабе слайда; близкий к ступени шкалы шаблона встаёт на неё (template.font_sizes)."""
        value = size * self.scale
        nearest = min(self.steps, key=lambda step: abs(step - value), default=None)
        if nearest is not None and abs(nearest - value) <= value * SCALE_SNAP:
            value = nearest
        return Pt(round(value, 1))

    @property
    def design(self) -> dict:
        return self.theme.get("design") or {}

    def absolute(self, size_pt: float | None, fallback: float) -> float:
        """Кегль шаблона (в пунктах его слайда) в масштабе _run/pt; без токена — запасной кегль."""
        return size_pt / self.scale if size_pt else fallback

    def role_size(self, role: str, fallback: float) -> float:
        """Кегль роли текста из типографики шаблона (title, body, body_2, caption…)."""
        typography = self.design.get("typography") or {}
        return self.absolute((typography.get(role) or {}).get("size_pt"), fallback)

    def rgb(self, color: str) -> RGBColor:
        return RGBColor.from_string(color.lstrip("#"))


def build_pptx(
    template_path: str | Path, slides: list[Slide], *, roles: dict, theme: dict, patterns: list[dict] | None = None
) -> bytes:
    return build_deck(template_path, slides, roles=roles, theme=theme, patterns=patterns).data


def build_deck(
    template_path: str | Path,
    slides: list[Slide],
    *,
    roles: dict,
    theme: dict,
    patterns: list[dict] | None = None,
    strict: bool = True,
) -> DeckBuild:
    """Собирает колоду. strict=False — для рендера: слайд-картинка не прерывает сборку, а попадает
    в факты (integrity.raster_slide), чтобы аудит показал проблему на превью."""
    prs = Presentation(str(template_path))
    # Слайды-примеры нужны как источники паттернов до конца сборки, потом удаляются.
    sources = list(prs.slides)
    by_id = {pattern["id"]: pattern for pattern in patterns or []}
    layouts = {
        f"{mi}:{li}": layout
        for mi, master in enumerate(prs.slide_masters)
        for li, layout in enumerate(master.slide_layouts)
    }
    fallback_key = roles.get("content") if roles.get("content") in layouts else next(iter(layouts))
    theme_colors = extract_theme(prs)["colors"]
    ctx = _Ctx(
        width=int(prs.slide_width),
        height=int(prs.slide_height),
        theme=theme,
        heading_font=theme.get("heading_font"),
        body_font=theme.get("body_font"),
    )

    chrome = ctx.design.get("chrome") or {}
    # Ряды диаграмм — цветами диаграмм шаблона, если они есть, иначе акцентами темы.
    chart_colors = (ctx.design.get("chart_style") or {}).get("series_colors") or theme["accents"]
    targets: list[Targets] = []
    for spec in slides:
        key = spec.layout_key if spec.layout_key in layouts else fallback_key
        layout = layouts[key]
        background = getattr(spec, "background_color", None) or layout_background(layout, theme_colors) or theme["background"]
        fg = getattr(spec, "text_color", None) or text_on(background, theme["text"])
        accent = getattr(spec, "accent_color", None) or theme.get("accent", "#0077FF")
        slide_ctx = replace(
            ctx,
            fg=fg,
            muted=mix(fg, background, 0.45),
            surface=mix(accent, background, 0.88),
            background=background,
            font_family=spec.font_family,
        )
        pattern = by_id.get(spec.pattern_id or "")
        built = None
        if pattern is not None:
            count_before = len(prs.slides)
            try:
                built = _pattern_slide(prs, pattern, spec, sources, layouts, slide_ctx, chart_colors, chrome)
            except Exception:  # noqa: BLE001 — сборка по паттерну не удалась: слайд собирается на макете
                logger.exception("Паттерн %s не собрался для слайда %s", pattern.get("id"), spec.id)
                _drop_slides_after(prs, count_before)
        if built is None:
            slide = prs.slides.add_slide(layout)
            built = _render(slide, spec, slide_ctx)
            _add_chrome(slide, layout, chrome)
        # Скрытые элементы, анимации и переход — поверх собранного слайда (export/motion.py).
        targets.append(
            apply_motion(
                prs.slides[len(prs.slides) - 1],
                spec,
                built,
                width=ctx.width,
                height=ctx.height,
            )
        )
    clear_slides(prs, len(sources))

    theme_fonts = extract_theme(prs).get("fonts") or {}
    logo_digests = {logo["image"] for logo in chrome.get("logos") or [] if logo.get("image")}
    facts = []
    for spec, slide in zip(slides, prs.slides):
        path = _user_image(spec)
        facts.append(
            slide_facts(
                slide,
                width=ctx.width,
                height=ctx.height,
                theme_fonts=theme_fonts,
                user_digests={digest(path.read_bytes())} if path else set(),
                logo_digests=logo_digests,
            )
        )

    buffer = BytesIO()
    prs.save(buffer)
    data = buffer.getvalue()
    _verify(data, expected=len(slides), strict=strict)
    return DeckBuild(data=data, targets=targets, facts=facts)


def _ensure_raster(path: Path) -> Path:
    """Если файл в формате SVG, растрирует его в PNG через PyMuPDF для python-pptx."""
    if path.suffix.lower() != ".svg":
        return path
    png_path = path.parent / f"{path.name}.png"
    try:
        if not png_path.exists() or png_path.stat().st_mtime < path.stat().st_mtime:
            with pymupdf.open(stream=path.read_bytes(), filetype="svg") as doc:
                if len(doc) > 0:
                    pix = doc[0].get_pixmap(dpi=150)
                    pix.save(str(png_path))
    except Exception as exc:
        logger.warning("Не удалось растрировать SVG %s: %s", path, exc)
        return path
    return png_path if png_path.exists() else path


def add_picture_fill(slide, path: Path, left: int, top: int, width: int, height: int):
    """Вставляет картинку с заполнением рамки без искажения пропорций (center crop)."""
    picture = slide.shapes.add_picture(str(path), left, top, width, height)
    _apply_crop(picture, None, path)
    return picture


def _resolve_image_url(url: str | None) -> Path | None:
    """Унифицированное разрешение пути к картинке (контент-пакет, user_images или прямой путь)."""
    if not url:
        return None
    url = url.strip()
    if not url:
        return None

    # 1. Контент-пакет презентации: /presentations/{id}/files/{file_id}[/content]
    match = re.search(r"/presentations/([0-9a-fA-F-]+)/files/([0-9a-fA-F-]+)", url)
    if match:
        pres_id, file_id = match.group(1), match.group(2)
        content_dir = storage_path("content", pres_id)
        if content_dir.is_dir():
            matches = [f for f in content_dir.glob(f"{file_id}.*") if f.is_file() and not f.name.endswith(".svg.png")]
            if not matches:
                matches = [f for f in content_dir.glob(f"{file_id}") if f.is_file()]
            if not matches:
                matches = [f for f in content_dir.glob(f"{file_id}.*") if f.is_file()]
            if matches:
                return _ensure_raster(matches[0])

    # 2. Пользовательские изображения: storage_path("user_images", filename)
    filename = Path(url).name
    if filename:
        user_img = storage_path("user_images", filename)
        if user_img.is_file():
            return _ensure_raster(user_img)

    # 3. Прямой путь на файловой системе
    direct_path = Path(url)
    try:
        if direct_path.is_file():
            return _ensure_raster(direct_path)
    except OSError:
        pass

    return None


def _valid_crop(crop: list[float] | None) -> tuple[float, float, float, float]:
    if not crop or len(crop) < 4:
        return 0.0, 0.0, 0.0, 0.0
    top, right, bottom, left = (max(0.0, min(0.95, float(v or 0))) for v in crop[:4])
    if top + bottom >= 0.95 or left + right >= 0.95:
        return 0.0, 0.0, 0.0, 0.0
    return top, right, bottom, left


def _apply_crop(pic, crop: list[float] | None, path: Path) -> None:
    """Картинка в рамке по модели PowerPoint — так же, как её показывает редактор (frontend/src/lib/imageCrop.ts):
    обрезка исходника crop = [top, right, bottom, left] (доли его сторон), затем срез по центру под пропорции
    рамки. Картинка никогда не растягивается; в .pptx обрезку можно снять."""
    try:
        with Image.open(path) as img:
            pixels_w, pixels_h = img.size
        width, height = pic.width, pic.height
        if not (pixels_w and pixels_h and width and height):
            return
        top, right, bottom, left = _valid_crop(crop)
        x0, x1, y0, y1 = left, 1 - right, top, 1 - bottom
        frame = width / height
        region = ((x1 - x0) * pixels_w) / ((y1 - y0) * pixels_h)
        if region > frame:  # видимая часть шире рамки — срезаем бока
            part = (y1 - y0) * pixels_h * frame / pixels_w
            center = (x0 + x1) / 2
            x0, x1 = center - part / 2, center + part / 2
        elif region < frame:  # выше рамки — срезаем верх и низ
            part = (x1 - x0) * pixels_w / frame / pixels_h
            center = (y0 + y1) / 2
            y0, y1 = center - part / 2, center + part / 2
        pic.crop_left, pic.crop_right = x0, 1 - x1
        pic.crop_top, pic.crop_bottom = y0, 1 - y1
    except Exception:  # noqa: BLE001 — картинка без кропа лучше, чем несобранный слайд
        logger.exception("Не удалось кадрировать картинку %s: %s", path, crop)


def _round_picture(pic, radius: float | None) -> None:
    """Скругление углов картинки (как в редакторе): фигура картинки — roundRect, adj — доля половины
    короткой стороны в стотысячных. radius 1 — полукруглые края, у квадрата — круг."""
    if not radius or radius <= 0:
        return
    try:
        sp_pr = pic._element.spPr
        geometry = sp_pr.find(qn("a:prstGeom"))
        if geometry is None:
            geometry = etree.SubElement(sp_pr, qn("a:prstGeom"))
            # По схеме prstGeom идёт сразу после xfrm.
            xfrm = sp_pr.find(qn("a:xfrm"))
            if xfrm is not None:
                xfrm.addnext(geometry)
        geometry.set("prst", "roundRect")
        for child in list(geometry):
            geometry.remove(child)
        guides = etree.SubElement(geometry, qn("a:avLst"))
        guide = etree.SubElement(guides, qn("a:gd"))
        guide.set("name", "adj")
        guide.set("fmla", f"val {int(min(radius, 1.0) * 50000)}")
    except Exception:  # noqa: BLE001 — картинка с прямыми углами лучше, чем несобранный слайд
        logger.exception("Не удалось скруглить картинку: %s", radius)


def _slide_image(spec: Slide) -> Path | None:
    """Унифицированное разрешение пути к картинке слайда (контент-пакет, user_images или прямой путь)."""
    if not spec.image or not spec.image.url:
        return None
    return _resolve_image_url(spec.image.url)


_user_image = _slide_image


def _place_image(slide, pattern: dict, spec: Slide, ctx: _Ctx) -> None:
    """Картинка слайда: в плейсхолдер картинки макета или на место картинки слайда-примера."""
    path = _slide_image(spec) if (spec.kind in {"image_text", "title"} or (spec.image and spec.image.url)) else None
    if path is None:
        return

    elem_boxes = getattr(spec, "element_boxes", {}) or {}
    elem_crops = getattr(spec, "element_crops", {}) or {}

    if "image" in elem_boxes:
        ix, iy, iw, ih = elem_boxes["image"]
        pic = slide.shapes.add_picture(str(path), int(ix * ctx.width), int(iy * ctx.height), int(iw * ctx.width), int(ih * ctx.height))
        _apply_crop(pic, elem_crops.get("image"), path)
        _round_picture(pic, (spec.element_radii or {}).get("image"))
        placeholders = [image for image in pattern.get("images") or [] if image.get("placeholder")]
        if placeholders:
            idx = placeholders[0]["id"] - PH_BASE
            holder = next((ph for ph in slide.placeholders if ph.placeholder_format.idx == idx), None)
            if holder is not None:
                remove(holder._element)
        return

    placeholders = [image for image in pattern.get("images") or [] if image.get("placeholder")]
    if placeholders:
        idx = placeholders[0]["id"] - PH_BASE
        holder = next((ph for ph in slide.placeholders if ph.placeholder_format.idx == idx), None)
        if holder is not None:
            if hasattr(holder, "insert_picture"):
                picture = holder.insert_picture(str(path))
                _apply_crop(picture, elem_crops.get("image"), path)
                _round_picture(picture, (spec.element_radii or {}).get("image"))
            else:
                remove(holder._element)
        return
    if spec.kind == "title":
        if pattern.get("images"):
            x, y, w, h = pattern["images"][0]["bbox"]
            pic = slide.shapes.add_picture(str(path), int(x * ctx.width), int(y * ctx.height), int(w * ctx.width), int(h * ctx.height))
            _apply_crop(pic, elem_crops.get("image"), path)
            _round_picture(pic, (spec.element_radii or {}).get("image"))
        else:
            logo_w, logo_h = int(ctx.width * 0.16), int(ctx.height * 0.08)
            pic = slide.shapes.add_picture(str(path), int(ctx.width * 0.06), int(ctx.height * 0.06), logo_w, logo_h)
            _apply_crop(pic, elem_crops.get("image"), path)
            _round_picture(pic, (spec.element_radii or {}).get("image"))
    else:
        x, y, w, h = pattern["images"][0]["bbox"] if pattern.get("images") else [0.55, 0.25, 0.38, 0.55]
        pic = slide.shapes.add_picture(str(path), int(x * ctx.width), int(y * ctx.height), int(w * ctx.width), int(h * ctx.height))
        _apply_crop(pic, elem_crops.get("image"), path)
        _round_picture(pic, (spec.element_radii or {}).get("image"))


def _pattern_slide(prs, pattern: dict, spec: Slide, sources, layouts, ctx: _Ctx, chart_colors, chrome) -> Targets:
    """Слайд по паттерну: клон слайда-примера или новый слайд на макете паттерна (origin=layout)."""
    if pattern.get("origin") == "layout":
        layout = layouts.get(pattern.get("layout_key") or "")
        if layout is None:
            raise ExportError(f"Макет паттерна {pattern['id']} не найден")
        slide = prs.slides.add_slide(layout)
        slide_targets = fill_pattern(
            slide, pattern, spec, accents=chart_colors, background=ctx.background, steps=ctx.steps
        )
        _add_chrome(slide, layout, chrome)
    else:
        index = pattern.get("slide_index", -1)
        if not 0 <= index < len(sources):
            raise ExportError(f"Слайд-пример паттерна {pattern['id']} не найден")
        slide = clone_slide(prs, sources[index])
        slide_targets = fill_pattern(
            slide, pattern, spec, accents=chart_colors, background=ctx.background, steps=ctx.steps
        )
    _place_image(slide, pattern, spec, ctx)
    _shapes(slide, spec.shapes, ctx, pattern_shapes=pattern.get("shapes"))
    return slide_targets


def _drop_slides_after(prs, count: int) -> None:
    """Убирает недособранные слайды (добавленные после count), чтобы не было дублей."""
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids)[count:]:
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def _add_chrome(slide, layout, chrome: dict) -> None:
    """Номер слайда и колонтитул с макета: python-pptx их не копирует, а шаблон их ставит."""
    wanted = set()
    if chrome.get("slide_number"):
        wanted.add(PP_PLACEHOLDER.SLIDE_NUMBER)
    if chrome.get("footer_text"):
        wanted.add(PP_PLACEHOLDER.FOOTER)
    if not wanted:
        return
    present = {ph.placeholder_format.type for ph in slide.placeholders}
    tree = slide.shapes._spTree
    next_id = max((int(el.get("id", "0")) for el in slide._element.iter(qn("p:cNvPr"))), default=1) + 1
    for ph in layout.placeholders:
        kind = ph.placeholder_format.type
        if kind not in wanted or kind in present:
            continue
        element = copy.deepcopy(ph._element)
        c_nv_pr = element.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}")
        if c_nv_pr is not None:
            c_nv_pr.set("id", str(next_id))
            next_id += 1
        if kind == PP_PLACEHOLDER.FOOTER:
            set_body(element.find(qn("p:txBody")), [chrome["footer_text"]], autofit=False)
        tree.append(element)
        present.add(kind)


def clear_slides(prs, count: int | None = None) -> None:
    """Удаляет слайды шаблона (первые count, по умолчанию все), оставляя мастера, макеты и тему."""
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids)[:count]:
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    # Разделы (p14:sectionLst) ссылаются на удалённые слайды — PowerPoint предложит «восстановить» файл.
    for section_list in prs.part._element.xpath(".//*[local-name()='sectionLst']"):
        ext = section_list.getparent()
        ext.getparent().remove(ext)


def _verify(data: bytes, *, expected: int, strict: bool = True) -> None:
    """integrity.file_opens и integrity.raster_slide: файл открывается, слайды не картинки.
    strict=False не проверяет слайды-картинки: при рендере их отмечают факты сборки."""
    try:
        prs = Presentation(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExportError(f"Собранный .pptx не открывается: {exc}") from exc
    if len(prs.slides) != expected:
        raise ExportError(f"В файле {len(prs.slides)} слайдов вместо {expected}")
    if not strict:
        return
    for number, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        if shapes and all(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes):
            raise ExportError(f"Слайд {number} состоит только из картинки")


# --- отрисовка слайда ------------------------------------------------------------


def _fraction(rect: tuple[int, int, int, int], ctx: _Ctx) -> list[float]:
    x, y, w, h = rect
    return [round(x / ctx.width, 4), round(y / ctx.height, 4), round(w / ctx.width, 4), round(h / ctx.height, 4)]


def _render(slide, spec: Slide, ctx: _Ctx) -> Targets:
    placeholders = list(slide.placeholders)
    title_ph = next((p for p in placeholders if p.placeholder_format.type in TITLE_TYPES), None)
    bodies = sorted(
        (p for p in placeholders if p.placeholder_format.type in BODY_TYPES),
        key=lambda p: p.placeholder_format.idx,
    )
    used: set[int] = set()

    targets: Targets = {}
    elem_fonts = getattr(spec, "element_fonts", {}) or {}
    elem_scales = getattr(spec, "element_scales", {}) or {}
    elem_boxes = getattr(spec, "element_boxes", {}) or {}
    elem_crops = getattr(spec, "element_crops", {}) or {}
    elem_colors = getattr(spec, "element_colors", {}) or {}
    elem_bolds = getattr(spec, "element_bolds", {}) or {}

    title_color = elem_colors.get("title") or getattr(spec, "text_color", None)
    # Жирность заголовка по умолчанию — из стиля заголовка макета, как в превью (SlideView: skin.title.bold).
    skins, theme_roles = ctx.theme.get("layouts") or {}, ctx.theme.get("roles") or {}
    role_key = theme_roles.get(spec.kind) or theme_roles.get("content") or ""
    skin_title = (skins.get(spec.layout_key or "") or skins.get(role_key) or {}).get("title") or {}
    default_bold = spec.is_bold if spec.is_bold is not None else skin_title.get("bold") is not False
    title_bold = elem_bolds.get("title") if "title" in elem_bolds else default_bold
    title_font = elem_fonts.get("title") or spec.font_family or ctx.heading_font
    title_scale = elem_scales.get("title", 1.0)
    if "title" in elem_boxes:
        tx, ty, tw, th = elem_boxes["title"]
        rect = (int(tx * ctx.width), int(ty * ctx.height), int(tw * ctx.width), int(th * ctx.height))
        box = slide.shapes.add_textbox(*rect)
        _write(
            # Кегль заголовка шаблона (как показывает редактор), а не условные 30 pt.
            box.text_frame, [spec.title], ctx, size=ctx.role_size("title", 30) * title_scale, bold=bool(title_bold), color=title_color, font_name=title_font, heading=True
        )
        targets["title"] = _fraction(rect, ctx)
        _styled_frame(box.text_frame, spec, "title", scale=False)
        if title_ph is not None:
            remove(title_ph._element)
    elif title_ph is not None:
        title_ph.text_frame.text = spec.title
        # Длинный заголовок ужимается в рамку макета, а не наезжает на декор (как автоподбор в PowerPoint).
        title_ph.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        used.add(title_ph.shape_id)
        for p in title_ph.text_frame.paragraphs:
            for r in p.runs:
                if title_font:
                    r.font.name = title_font
                if title_color:
                    r.font.color.rgb = ctx.rgb(title_color)
                if title_bold is not None:
                    r.font.bold = title_bold
                if title_scale != 1.0:
                    # Масштаб из редактора — от кегля заголовка в типографике шаблона.
                    r.font.size = ctx.pt(ctx.role_size("title", 30) * title_scale)
        _styled_frame(title_ph.text_frame, spec, "title", scale=False)
        if None not in (title_ph.left, title_ph.top, title_ph.width, title_ph.height):
            targets["title"] = _fraction((title_ph.left, title_ph.top, title_ph.width, title_ph.height), ctx)
    else:
        rect = (int(ctx.width * 0.06), int(ctx.height * 0.07), int(ctx.width * 0.88), int(ctx.height * 0.14))
        box = slide.shapes.add_textbox(*rect)
        _write(
            box.text_frame, [spec.title], ctx, size=ctx.role_size("title", 30) * title_scale, bold=bool(title_bold), color=title_color, font_name=title_font, heading=True
        )
        _styled_frame(box.text_frame, spec, "title", scale=False)
        targets["title"] = _fraction(rect, ctx)

    if "content" in elem_boxes:
        cx, cy, cw, ch = elem_boxes["content"]
        area = (int(cx * ctx.width), int(cy * ctx.height), int(cw * ctx.width), int(ch * ctx.height))
    else:
        area = _content_area(title_ph, ctx)
    kind = spec.kind
    if kind in CONTENT_TARGET:
        targets[CONTENT_TARGET[kind]] = _fraction(area, ctx)

    sub_font = elem_fonts.get("subtitle") or spec.font_family
    sub_scale = elem_scales.get("subtitle", 1.0)
    sub_color = elem_colors.get("subtitle") or getattr(spec, "text_color", None)
    sub_bold = elem_bolds.get("subtitle") if "subtitle" in elem_bolds else getattr(spec, "is_bold", None)
    if kind in {"title", "section", "closing"}:
        if spec.subtitle:
            if "subtitle" in elem_boxes:
                sx, sy, sw, sh = elem_boxes["subtitle"]
                rect = (int(sx * ctx.width), int(sy * ctx.height), int(sw * ctx.width), int(sh * ctx.height))
                box = slide.shapes.add_textbox(*rect)
                _write(
                    box.text_frame, [spec.subtitle], ctx, size=ctx.role_size("subtitle", 20) * sub_scale, bold=bool(sub_bold), color=sub_color, font_name=sub_font
                )
                targets["subtitle"] = _fraction(rect, ctx)
                _styled_frame(box.text_frame, spec, "subtitle", scale=False)
                if bodies:
                    remove(bodies[0]._element)
            elif bodies:
                bodies[0].text_frame.text = spec.subtitle
                used.add(bodies[0].shape_id)
                for p in bodies[0].text_frame.paragraphs:
                    for r in p.runs:
                        if sub_font:
                            r.font.name = sub_font
                        if sub_color:
                            r.font.color.rgb = ctx.rgb(sub_color)
                        if sub_bold is not None:
                            r.font.bold = sub_bold
                _styled_frame(bodies[0].text_frame, spec, "subtitle")
                body = bodies[0]
                if None not in (body.left, body.top, body.width, body.height):
                    targets["subtitle"] = _fraction((body.left, body.top, body.width, body.height), ctx)
            else:
                x, y, w, _ = area
                rect = (x, y, w, int(ctx.height * 0.12))
                box = slide.shapes.add_textbox(*rect)
                _write(
                    box.text_frame, [spec.subtitle], ctx, size=ctx.role_size("subtitle", 20) * sub_scale, bold=bool(sub_bold), color=sub_color, font_name=sub_font
                )
                _styled_frame(box.text_frame, spec, "subtitle", scale=False)
                # Рамка подзаголовка — для подсветки аудита и адресации в Студии (скрыть, анимировать).
                targets["subtitle"] = _fraction(rect, ctx)
        if kind == "title" and spec.image:
            path = _slide_image(spec)
            if path is not None:
                if "image" in elem_boxes:
                    ix, iy, iw, ih = elem_boxes["image"]
                    pic = slide.shapes.add_picture(str(path), int(ix * ctx.width), int(iy * ctx.height), int(iw * ctx.width), int(ih * ctx.height))
                    _apply_crop(pic, elem_crops.get("image"), path)
                    _round_picture(pic, (spec.element_radii or {}).get("image"))
                else:
                    img_ph = next(
                        (
                            p
                            for p in placeholders
                            if hasattr(p, "insert_picture")
                            and getattr(p.placeholder_format, "type", None) in {PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.BITMAP}
                        ),
                        None,
                    )
                    if img_ph is not None:
                        picture = img_ph.insert_picture(str(path))
                        _apply_crop(picture, elem_crops.get("image"), path)
                        _round_picture(picture, (spec.element_radii or {}).get("image"))
                        used.add(img_ph.shape_id)
                    else:
                        logo_w, logo_h = int(ctx.width * 0.16), int(ctx.height * 0.08)
                        pic = slide.shapes.add_picture(
                            str(path),
                            int(ctx.width * 0.06),
                            int(ctx.height * 0.06),
                            logo_w,
                            logo_h,
                        )
                        _apply_crop(pic, elem_crops.get("image"), path)
                        _round_picture(pic, (spec.element_radii or {}).get("image"))
    elif kind in {"bullets", "agenda"}:
        if kind == "bullets" and spec.image and spec.image.url:
            _image_text(slide, spec, area, ctx)
        else:
            items = spec.bullets
            if kind == "agenda":
                items = [f"{i:02d}   {item}" for i, item in enumerate(items, start=1)]
            _bullets_box(
                slide,
                area,
                items,
                ctx,
                bullets=kind == "bullets",
                size=ctx.role_size("body", 20),
                elem_boxes=elem_boxes,
                spec=spec,
            )
    elif kind == "two_columns":
        _two_columns(slide, spec, area, ctx)
    elif kind == "cards":
        _cards(slide, spec, area, ctx)
    elif kind == "image_text":
        _image_text(slide, spec, area, ctx)
    elif kind == "chart" and spec.chart:
        _chart(slide, spec, area, ctx)
    elif kind == "table" and spec.table:
        _table(slide, spec, area, ctx)
    elif kind == "kpi":
        _kpis(slide, spec, area, ctx)
    elif kind == "timeline":
        _timeline(slide, spec, area, ctx)
    elif kind == "quote" and spec.quote:
        x, y, w, h = area
        box = slide.shapes.add_textbox(x, y, w, h)
        _write(box.text_frame, [f"«{spec.quote.text}»"], ctx, size=30, italic=True, heading=True)
        _styled(box.text_frame.paragraphs[0], spec, "quote.text")
        if spec.quote.author:
            paragraph = box.text_frame.add_paragraph()
            _run(paragraph, f"— {spec.quote.author}", ctx, size=16, color=ctx.muted)
            _styled(paragraph, spec, "quote.author")

    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes

    if spec.shapes:
        _shapes(slide, spec.shapes, ctx)

    # Пустые плейсхолдеры макета показывают подсказки «Введите текст» — убираем их.
    for placeholder in list(slide.placeholders):
        if placeholder.shape_id not in used:
            placeholder._element.getparent().remove(placeholder._element)
    return targets


def _content_area(title_ph, ctx: _Ctx) -> tuple[int, int, int, int]:
    title = None
    if title_ph is not None and None not in (title_ph.left, title_ph.top, title_ph.width, title_ph.height):
        title = (int(title_ph.left), int(title_ph.top), int(title_ph.width), int(title_ph.height))
    return content_frame(title, ctx.width, ctx.height)


def _run(
    paragraph,
    text: str,
    ctx: _Ctx,
    *,
    size: float,
    bold=False,
    italic=False,
    color: str | None = None,
    styled: bool = False,
    font_name: str | None = None,
    heading: bool = False,
):
    """styled=True — цвет текста задаёт стиль объекта (таблицы), а не фон слайда.
    heading=True — заголовочный текст (заголовок, карточки, показатели): гарнитура заголовков шаблона."""
    run = paragraph.add_run()
    run.text = text
    run.font.size = ctx.pt(size)
    run.font.bold = bold
    run.font.italic = italic
    fn = font_name or ctx.font_family or (ctx.heading_font if heading else ctx.body_font)
    if fn:
        run.font.name = fn
    if not styled:
        run.font.color.rgb = ctx.rgb(color or ctx.fg)
    return run


def _styled(paragraph, spec: Slide, field: str, *, scale: bool = True) -> None:
    """Оформление поля из редактора поверх стиля сборки (text_style). Шрифт и цвет слайда целиком сборка уже
    учла через ctx, поэтому здесь — только правки самого поля. scale=False — кегль уже посчитан с масштабом."""
    style_paragraph(paragraph._p, style_for(spec, field, slide_defaults=False), scale=scale)


def _styled_frame(text_frame, spec: Slide, field: str, *, scale: bool = True) -> None:
    for paragraph in text_frame.paragraphs:
        _styled(paragraph, spec, field, scale=scale)


def _write(
    text_frame, lines: list[str], ctx: _Ctx, *, size: float, bold=False, italic=False, color=None, font_name=None, heading=False
):
    text_frame.word_wrap = True
    for i, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        _run(paragraph, line, ctx, size=size, bold=bold, italic=italic, color=color, font_name=font_name, heading=heading)


def _set_bullet(paragraph, ctx: _Ctx) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    indent = int(Pt(18 * ctx.scale))
    p_pr.set("marL", str(indent))
    p_pr.set("indent", str(-indent))
    # Порядок дочерних элементов a:pPr задан схемой: spcBef → buClr → buChar.
    spacing = etree.SubElement(p_pr, qn("a:spcBef"))
    etree.SubElement(spacing, qn("a:spcPts")).set("val", str(int(900 * ctx.scale)))
    bullet_color = etree.SubElement(p_pr, qn("a:buClr"))
    etree.SubElement(bullet_color, qn("a:srgbClr")).set("val", ctx.theme["accent"].lstrip("#"))
    etree.SubElement(p_pr, qn("a:buChar")).set("char", "•")


def _bullets_box(
    slide,
    area,
    items: list[str],
    ctx: _Ctx,
    *,
    bullets: bool = True,
    size: float = 20,
    elem_boxes: dict | None = None,
    spec: Slide | None = None,
    fields: list[str] | None = None,
) -> None:
    """Список. fields[i] — адрес поля пункта для оформления из редактора (по умолчанию bullets.i);
    свою рамку (element_boxes) могут иметь только пункты основного списка слайда."""
    if not items:
        return
    elem_boxes = elem_boxes or {}
    fields = fields or [f"bullets.{i}" for i in range(len(items))]

    if "bullets" in elem_boxes:
        bx, by, bw, bh = elem_boxes["bullets"]
        area = (int(bx * ctx.width), int(by * ctx.height), int(bw * ctx.width), int(bh * ctx.height))

    remaining_items = []
    for i, item in enumerate(items):
        key = f"bullets.{i}"
        if key in elem_boxes:
            ix, iy, iw, ih = elem_boxes[key]
            b_area = (int(ix * ctx.width), int(iy * ctx.height), int(iw * ctx.width), int(ih * ctx.height))
            b_box = slide.shapes.add_textbox(*b_area)
            b_frame = b_box.text_frame
            b_frame.word_wrap = True
            p = b_frame.paragraphs[0]
            if bullets:
                _set_bullet(p, ctx)
            _run(p, item, ctx, size=size)
            if spec is not None:
                _styled(p, spec, fields[i])
        else:
            remaining_items.append((i, item))

    if not remaining_items:
        return

    x, y, w, h = area
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    for idx, (orig_i, item) in enumerate(remaining_items):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        if bullets:
            _set_bullet(paragraph, ctx)
        _run(paragraph, item, ctx, size=size)
        if spec is not None:
            _styled(paragraph, spec, fields[orig_i])


def _two_columns(slide, spec: Slide, area, ctx: _Ctx) -> None:
    x, y, w, h = area
    elem_boxes = getattr(spec, "element_boxes", {}) or {}
    gap = int(w * 0.05)
    column_width = (w - gap) // 2
    for i, column in enumerate(spec.columns[:2]):
        col_key = f"columns.{i}"
        left = x + i * (column_width + gap)
        top = y
        cur_w = column_width
        cur_h = h
        if col_key in elem_boxes:
            cx, cy, cw, ch = elem_boxes[col_key]
            left, top, cur_w, cur_h = int(cx * ctx.width), int(cy * ctx.height), int(cw * ctx.width), int(ch * ctx.height)
        header = slide.shapes.add_textbox(left, top, cur_w, int(cur_h * 0.16))
        _write(header.text_frame, [column.title], ctx, size=22, bold=True, color=ctx.theme["accent"], heading=True)
        _styled_frame(header.text_frame, spec, f"columns.{i}.title")
        _bullets_box(
            slide,
            (left, top + int(cur_h * 0.18), cur_w, int(cur_h * 0.82)),
            column.bullets,
            ctx,
            size=18,
            spec=spec,
            fields=[f"columns.{i}.bullets.{j}" for j in range(len(column.bullets))],
        )


def _cards(slide, spec: Slide, area, ctx: _Ctx) -> None:
    """Сетка карточек: заголовок акцентным цветом и короткий текст на подложке."""
    if not spec.columns:
        return
    x, y, w, h = area
    elem_boxes = getattr(spec, "element_boxes", {}) or {}
    count = len(spec.columns)
    per_row = count if count <= 4 else math.ceil(count / 2)
    rows = math.ceil(count / per_row)
    gap = int(w * 0.025)
    card_width = (w - gap * (per_row - 1)) // per_row
    card_height = min((h - gap * (rows - 1)) // rows, int(ctx.height * 0.36))
    for i, card_spec in enumerate(spec.columns):
        row, col = divmod(i, per_row)
        card_x = x + col * (card_width + gap)
        card_y = y + row * (card_height + gap)
        card_w = card_width
        card_h = card_height
        card_key = f"columns.{i}"
        if card_key in elem_boxes:
            cx, cy, cw, ch = elem_boxes[card_key]
            card_x, card_y, card_w, card_h = int(cx * ctx.width), int(cy * ctx.height), int(cw * ctx.width), int(ch * ctx.height)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            card_x,
            card_y,
            card_w,
            card_h,
        )
        card.adjustments[0] = 0.08
        card.fill.solid()
        card.fill.fore_color.rgb = ctx.rgb(ctx.surface)
        card.line.fill.background()
        frame = card.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        _run(frame.paragraphs[0], card_spec.title, ctx, size=20, bold=True, color=ctx.theme["accent"], heading=True)
        _styled(frame.paragraphs[0], spec, f"columns.{i}.title")
        for j, bullet in enumerate(card_spec.bullets):
            paragraph = frame.add_paragraph()
            paragraph.space_before = Pt(6 * ctx.scale)
            _run(paragraph, bullet, ctx, size=15)
            _styled(paragraph, spec, f"columns.{i}.bullets.{j}")
        for paragraph in frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT


def _image_text(slide, spec: Slide, area, ctx: _Ctx) -> None:
    x, y, w, h = area
    elem_boxes = getattr(spec, "element_boxes", {}) or {}
    elem_crops = getattr(spec, "element_crops", {}) or {}
    elem_fonts = getattr(spec, "element_fonts", {}) or {}
    elem_scales = getattr(spec, "element_scales", {}) or {}

    text_width = int(w * 0.52)
    _bullets_box(
        slide,
        (x, y, text_width, h),
        spec.bullets,
        ctx,
        size=18,
        elem_boxes=elem_boxes,
        spec=spec,
    )
    left = x + int(w * 0.57)
    img_width = x + w - left
    img_height = h

    if "image" in elem_boxes:
        ix, iy, iw, ih = elem_boxes["image"]
        left, y, img_width, img_height = int(ix * ctx.width), int(iy * ctx.height), int(iw * ctx.width), int(ih * ctx.height)

    path = _slide_image(spec)
    if path is not None:
        try:
            pic = slide.shapes.add_picture(str(path), left, y, img_width, img_height)
            _apply_crop(pic, elem_crops.get("image"), path)
            _round_picture(pic, (spec.element_radii or {}).get("image"))
            return
        except Exception:
            logger.exception("Не удалось вставить картинку %s в слайд", path)

    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, y, img_width, img_height)
    frame.fill.solid()
    frame.fill.fore_color.rgb = ctx.rgb(ctx.surface)
    frame.line.fill.background()
    text = frame.text_frame
    text.word_wrap = True
    text.vertical_anchor = MSO_ANCHOR.MIDDLE
    prompt = spec.image.prompt if spec.image else "Иллюстрация"
    _write(text, ["Изображение будет сгенерировано", prompt], ctx, size=12, color=ctx.muted)
    for paragraph in text.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER


def _chart(slide, spec: Slide, area, ctx: _Ctx) -> None:
    chart_spec = spec.chart
    chart_style = ctx.design.get("chart_style") or {}
    x, y, w, h = area
    bars = chart_spec.type in ("bar", "horizontal_bar")
    labelled = chart_spec.show_values if chart_spec.show_values is not None else True
    # Подписанным столбцам ось значений не нужна — единица измерения уходит в подпись диаграммы.
    caption_text = chart_caption(chart_spec) if bars and labelled else chart_spec.caption
    if caption_text:
        caption = slide.shapes.add_textbox(x, y, w, int(h * 0.1))
        _write(caption.text_frame, [caption_text], ctx, size=ctx.role_size("caption", 14), color=ctx.muted)
        _styled_frame(caption.text_frame, spec, "chart.caption")
        y, h = y + int(h * 0.11), int(h * 0.89)

    data = CategoryChartData()
    data.categories = chart_spec.categories
    for series in chart_spec.series:
        data.add_series(series.name or "Серия", series.values)
    graphic = slide.shapes.add_chart(CHART_TYPES[chart_spec.type], x, y, w, h, data)
    chart = graphic.chart
    chart.has_title = False  # подпись диаграммы — отдельным текстом над ней
    chart.font.size = ctx.pt(ctx.absolute(chart_style.get("font_size_pt"), 14))
    chart.font.color.rgb = ctx.rgb(ctx.fg)
    if chart_spec.show_legend is not None:
        chart.has_legend = chart_spec.show_legend
    else:
        chart.has_legend = len(chart_spec.series) > 1 or chart_spec.type in ("pie", "donut")
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    plot = chart.plots[0]
    plot.has_data_labels = labelled
    if labelled:
        plot.data_labels.font.size = ctx.pt(12)
        plot.data_labels.number_format = number_format(chart_spec)
        plot.data_labels.number_format_is_linked = False
    accents = chart_spec.colors or chart_style.get("series_colors") or ctx.theme["accents"]
    if chart_spec.type in ("pie", "donut"):
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = ctx.rgb(accents[i % len(accents)])
        return

    if bars and labelled:
        chart.value_axis.visible = False
        chart.value_axis.has_major_gridlines = False
    elif chart_spec.unit:
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = chart_spec.unit
    if chart_spec.type == "horizontal_bar":
        # У горизонтальных столбцов PowerPoint рисует первую категорию снизу — порядок как в данных.
        chart.category_axis.reverse_order = True
    for i, series in enumerate(plot.series):
        color = accents[i % len(accents)]
        if chart_spec.type == "line":
            series.format.line.color.rgb = ctx.rgb(color)
            series.format.line.width = Pt(2.5)
            series.smooth = False
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = ctx.rgb(color)
    highlight_points(plot, chart_spec, ctx.background)


def _table(slide, spec: Slide, area, ctx: _Ctx) -> None:
    table_spec = spec.table
    x, y, w, h = area
    columns = max(len(table_spec.columns), max((len(r) for r in table_spec.rows), default=1))
    rows = len(table_spec.rows) + 1
    row_height = min(h // rows, int(Pt(40 * ctx.scale)))
    table = slide.shapes.add_table(rows, columns, x, y, w, row_height * rows).table
    # Стиль таблиц шаблона (tableStyles.xml) вместо стандартного стиля python-pptx.
    style = ctx.design.get("table_style") or {}
    if style.get("style_id"):
        table_props = table._tbl.tblPr
        style_id = table_props.find(qn("a:tableStyleId"))
        if style_id is None:
            style_id = etree.SubElement(table_props, qn("a:tableStyleId"))
        style_id.text = style["style_id"]
        table_props.set("firstRow", "1" if style.get("first_row", True) else "0")
        table_props.set("bandRow", "1" if style.get("band_row", True) else "0")
    size = ctx.absolute(style.get("font_size_pt"), 14)
    cells = [table_spec.columns, *table_spec.rows]
    numeric = numeric_columns(table_spec)
    for r, row in enumerate(cells):
        for c in range(columns):
            cell = table.cell(r, c)
            cell.text = ""
            text = row[c] if c < len(row) else ""
            paragraph = cell.text_frame.paragraphs[0]
            _run(paragraph, text, ctx, size=size, bold=r == 0, styled=True)
            _styled(paragraph, spec, f"table.columns.{c}" if r == 0 else f"table.rows.{r - 1}.{c}")
            if c in numeric:
                # Числа сравнивают по разрядам — столбец чисел выравнивается вправо вместе с шапкой.
                paragraph.alignment = PP_ALIGN.RIGHT


def _kpis(slide, spec: Slide, area, ctx: _Ctx) -> None:
    if not spec.kpis:
        return
    x, y, w, h = area
    count = len(spec.kpis)
    gap = int(w * 0.03)
    card_width = (w - gap * (count - 1)) // count
    card_height = min(h, int(ctx.height * 0.4))
    for i, kpi in enumerate(spec.kpis):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x + i * (card_width + gap), y, card_width, card_height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = ctx.rgb(ctx.surface)
        card.line.fill.background()
        frame = card.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _run(frame.paragraphs[0], kpi.value, ctx, size=44, bold=True, color=ctx.theme["accent"], heading=True)
        _styled(frame.paragraphs[0], spec, f"kpis.{i}.value")
        label = frame.add_paragraph()
        _run(label, kpi.label, ctx, size=16, color=ctx.fg)
        _styled(label, spec, f"kpis.{i}.label")
        for paragraph in frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT


def _timeline(slide, spec: Slide, area, ctx: _Ctx) -> None:
    if not spec.timeline:
        return
    x, y, w, h = area
    count = len(spec.timeline)
    line_y = y + int(h * 0.12)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, line_y, x + w, line_y)
    line.line.color.rgb = ctx.rgb(ctx.theme["accent"])
    line.line.width = Pt(2)
    step = w // count
    dot = int(Emu(Pt(16 * ctx.scale)))
    for i, item in enumerate(spec.timeline):
        left = x + i * step
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, line_y - dot // 2, dot, dot)
        marker.fill.solid()
        marker.fill.fore_color.rgb = ctx.rgb(ctx.theme["accent"])
        marker.line.fill.background()
        box = slide.shapes.add_textbox(left, line_y + dot, step - int(w * 0.02), int(h * 0.7))
        frame = box.text_frame
        frame.word_wrap = True
        _run(frame.paragraphs[0], item.date, ctx, size=14, bold=True, color=ctx.theme["accent"])
        _styled(frame.paragraphs[0], spec, f"timeline.{i}.date")
        title_paragraph = frame.add_paragraph()
        _run(title_paragraph, item.title, ctx, size=18, bold=True)
        _styled(title_paragraph, spec, f"timeline.{i}.title")
        if item.text:
            text_paragraph = frame.add_paragraph()
            _run(text_paragraph, item.text, ctx, size=14, color=ctx.muted)
            _styled(text_paragraph, spec, f"timeline.{i}.text")


def _changed(value, original, tolerance: float = 1e-3) -> bool:
    """Правка из редактора: координаты сравниваются с допуском (редактор округляет их до тысячных)."""
    if isinstance(value, (int, float)) and isinstance(original, (int, float)):
        return abs(value - original) > tolerance
    return (value or None) != (original or None)


def _shapes(slide, shapes: list, ctx: _Ctx, pattern_shapes: list | None = None) -> None:
    existing = {s.shape_id: s for s in slide.shapes}
    pattern_shapes = pattern_shapes or []
    # Фигуры слайда-примера в том виде, в каком их разобрал парсер: сравнивая с ними, сборка переносит
    # в клон только правки пользователя. Цвет парсер знает приблизительно (без прозрачности, оттенков
    # темы и карты цветов мастера) — перекрашивать нетронутую фигуру шаблона значит исказить её.
    originals = {ps.get("template_shape_id"): ps for ps in pattern_shapes if ps.get("template_shape_id")}

    # Если шаблонные фигуры были удалены пользователем из spec.shapes:
    if pattern_shapes:
        kept_tpl_ids = {
            getattr(s, "template_shape_id", None) or (s.get("template_shape_id") if isinstance(s, dict) else None)
            for s in (shapes or [])
        }
        for ps in pattern_shapes:
            pid = ps.get("template_shape_id")
            if pid and pid not in kept_tpl_ids and pid in existing:
                try:
                    el = existing[pid]._element
                    if el.getparent() is not None:
                        el.getparent().remove(el)
                except Exception:
                    pass

    if not shapes:
        return

    for s in shapes:
        kind = getattr(s, "kind", "circle") if not isinstance(s, dict) else s.get("kind", "circle")
        x = getattr(s, "x", 0.5) if not isinstance(s, dict) else s.get("x", 0.5)
        y = getattr(s, "y", 0.5) if not isinstance(s, dict) else s.get("y", 0.5)
        w = getattr(s, "width", 0.15) if not isinstance(s, dict) else s.get("width", 0.15)
        h = getattr(s, "height", 0.15) if not isinstance(s, dict) else s.get("height", 0.15)
        color = getattr(s, "color", None) if not isinstance(s, dict) else s.get("color")
        text = getattr(s, "text", None) if not isinstance(s, dict) else s.get("text")
        border = getattr(s, "border", False) if not isinstance(s, dict) else s.get("border", False)
        tpl_id = getattr(s, "template_shape_id", None) if not isinstance(s, dict) else s.get("template_shape_id")

        left = int(x * ctx.width)
        top = int(y * ctx.height)
        width = int(w * ctx.width)
        height = int(h * ctx.height)

        # Текстовый блок Студии — фигура без заливки: только текст цветом слайда.
        text_box = str(color or "").lower() == "transparent"
        fill_color = ctx.theme.get("accent", "#0077FF") if text_box or not color else color

        # Фигура шаблона, которой нет среди фигур слайда верхнего уровня, лежит внутри группы
        # или удалена сборкой вместе с лишней карточкой — пересоздавать её нельзя: будут дубли и «сироты».
        # Фигура чужого слайда-примера (после смены паттерна или на макете) к этому слайду не относится.
        if tpl_id:
            original = originals.get(tpl_id)
            if tpl_id not in existing or original is None:
                continue
            sp = existing[tpl_id]
            if any(_changed(value, original.get(key)) for key, value in (("x", x), ("y", y), ("width", w), ("height", h))):
                sp.left = left
                sp.top = top
                sp.width = width
                sp.height = height
            if color and str(color).upper() != str(original.get("color") or "").upper():
                try:
                    sp.fill.solid()
                    sp.fill.fore_color.rgb = ctx.rgb(fill_color)
                except Exception:
                    pass
            if text and sp.has_text_frame and _changed(text, original.get("text")):
                sp.text_frame.text = text
            continue

        if kind == "image":
            img_url = getattr(s, "image_url", None) if not isinstance(s, dict) else s.get("image_url")
            img_path = _resolve_image_url(img_url)
            crop_val = getattr(s, "crop", None) if not isinstance(s, dict) else s.get("crop")
            if img_path and img_path.is_file():
                try:
                    pic = slide.shapes.add_picture(str(img_path), left, top, width, height)
                    _apply_crop(pic, crop_val, img_path)
                    _round_picture(pic, getattr(s, "radius", None) if not isinstance(s, dict) else s.get("radius"))
                except Exception:
                    logger.exception("Не удалось вставить фигуру-картинку %s", img_path)
            continue

        if kind == "circle":
            sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
        elif kind == "line":
            sp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left, top, left + width, top)
            sp.line.color.rgb = ctx.rgb(fill_color)
            sp.line.width = Pt(3)
            continue
        elif kind in ("badge", "pill", "rect"):
            sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
            if kind in ("badge", "pill"):
                sp.adjustments[0] = 0.5
            else:
                sp.adjustments[0] = 0.08
        else:
            sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)

        if text_box:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = ctx.rgb(fill_color)
        if border:
            sp.line.color.rgb = ctx.rgb(ctx.fg)
            sp.line.width = Pt(1.5)
        else:
            sp.line.fill.background()

        if text:
            tf = sp.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if text_box else PP_ALIGN.CENTER
            shape_font = getattr(s, "font_family", None) if not isinstance(s, dict) else s.get("font_family")
            _run(
                p,
                text,
                ctx,
                size=20 if text_box else 14,
                bold=not text_box,
                color=ctx.fg if text_box else ("#FFFFFF" if text_on(fill_color, "#FFFFFF") == "#FFFFFF" else "#000000"),
                font_name=shape_font,
            )

