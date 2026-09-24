"""Извлечение текста из файлов контент-пакета, выдержки для модели и разбиение брифа на тезисы."""

import json
import logging
import re
import zipfile
from pathlib import Path

import pymupdf
from lxml import etree
from pptx import Presentation

logger = logging.getLogger(__name__)

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)

_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".svg"}
_DATA_EXTENSIONS = {".json", ".yaml", ".yml"}


def extract_text(path: Path, filename: str) -> str | None:
    """Текст файла или None, если формат не поддержан."""
    ext = Path(filename).suffix.lower()
    try:
        if ext in {".txt", ".md", ".csv"}:
            return _read_text(path)
        if ext == ".pdf":
            with pymupdf.open(path) as doc:
                return "\n".join(page.get_text() for page in doc)
        if ext == ".docx":
            return _docx_text(path)
        if ext == ".pptx":
            return _pptx_text(path)
        if ext == ".json":
            return _json_text(path)
        if ext in {".yaml", ".yml"}:
            return _yaml_text(path)
        if ext in _IMAGE_EXTENSIONS:
            return _image_description(path, filename)
    except Exception:
        logger.exception("Не удалось извлечь текст из %s", filename)
    return None


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        root = etree.fromstring(zf.read("word/document.xml"), _XML_PARSER)
    paragraphs = ("".join(t.text or "" for t in p.iter(f"{{{NS_W}}}t")) for p in root.iter(f"{{{NS_W}}}p"))
    return "\n".join(p for p in paragraphs if p.strip())


def _pptx_text(path: Path) -> str:
    lines = []
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
    return "\n".join(lines)


def _flatten(obj, *, prefix: str = "", depth: int = 0) -> list[str]:
    """Рекурсивно разворачивает JSON/YAML объект в список «ключ: значение»."""
    lines: list[str] = []
    if depth > 6:
        lines.append(f"{prefix}: [...]")
        return lines
    if isinstance(obj, dict):
        for k, v in obj.items():
            full_key = f"{prefix}.{k}" if prefix else str(k)
            if isinstance(v, (dict, list)):
                lines.extend(_flatten(v, prefix=full_key, depth=depth + 1))
            else:
                val = str(v) if v is not None else "null"
                if len(val) > 200:
                    val = val[:200] + "…"
                lines.append(f"{full_key}: {val}")
    elif isinstance(obj, list):
        for i, item in enumerate(obj[:50]):  # ограничение списков
            if isinstance(item, (dict, list)):
                lines.extend(_flatten(item, prefix=f"{prefix}[{i}]", depth=depth + 1))
            else:
                val = str(item) if item is not None else "null"
                lines.append(f"{prefix}[{i}]: {val}")
        if len(obj) > 50:
            lines.append(f"{prefix}: … ещё {len(obj) - 50} элементов")
    else:
        val = str(obj) if obj is not None else "null"
        lines.append(f"{prefix}: {val}")
    return lines


def _json_text(path: Path) -> str:
    raw = _read_text(path)
    try:
        data = json.loads(raw)
        return "\n".join(_flatten(data))
    except json.JSONDecodeError:
        # Если невалидный JSON — вернём как есть (может быть JSONL)
        return raw


def _yaml_text(path: Path) -> str:
    try:
        import yaml  # type: ignore[import-untyped]

        raw = _read_text(path)
        data = yaml.safe_load(raw)
        if data is None:
            return ""
        return "\n".join(_flatten(data))
    except ImportError:
        # PyYAML не установлен — вернём сырой текст
        return _read_text(path)
    except Exception:
        return _read_text(path)


def _image_description(path: Path, filename: str) -> str:
    """Метаданные картинки: размер, формат. Если Pillow доступен — добавляем точные размеры."""
    ext = path.suffix.lower()
    meta_parts = [f"Изображение: {filename}", f"Формат: {ext.lstrip('.').upper()}"]
    try:
        from PIL import Image  # type: ignore[import-untyped]

        with Image.open(path) as img:
            w, h = img.size
            mode = img.mode
            meta_parts.append(f"Размер: {w}×{h} пикселей")
            meta_parts.append(f"Цветовое пространство: {mode}")
    except Exception:
        size_bytes = path.stat().st_size if path.exists() else 0
        meta_parts.append(f"Размер файла: {size_bytes} байт")
    return "\n".join(meta_parts)


def sources_digest(files: list[tuple[str, str]], *, per_file: int, total: int) -> str:
    """Материалы для промпта: имя файла и выдержка из его текста.

    Целиком контент-пакет в запрос не влезет и не нужен: модели хватает начала каждого
    файла, где обычно и лежат цифры и тезисы. Границы заданы LLM_SOURCE_CHARS и
    LLM_SOURCES_TOTAL_CHARS, обрезка отмечается явно — чтобы модель не додумывала хвост."""
    blocks: list[str] = []
    left = total
    for filename, text in files:
        text = re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", (text or "").strip()))
        if not text or left <= 0:
            continue
        limit = min(per_file, left)
        excerpt = text[:limit]
        if len(text) > limit:
            excerpt += "\n[…документ обрезан…]"
        left -= len(excerpt)
        blocks.append(f"### {filename}\n{excerpt}")
    return "\n\n".join(blocks)


def split_sentences(text: str, limit: int = 12) -> list[str]:
    """Делит бриф на тезисы: по предложениям и строкам (в том числе маркированным)."""
    theses: list[str] = []
    for part in re.split(r"(?<=[.!?…])\s+|\n+", text or ""):
        thesis = part.strip().lstrip("-•*—–·").strip().rstrip(".")
        if len(thesis.split()) >= 3:
            theses.append(thesis)
        if len(theses) >= limit:
            break
    return theses
