"""Рендер шаблона через LibreOffice: слайды-примеры и фоны макетов — одним прогоном.

Рендеры слайдов-примеров показывают паттерны в интерфейсе и уходят в VLM
для разметки. Рендеры макетов (всё, что дают макет и мастер, без плейсхолдеров)
служат подложкой слайдов в превью и HTML-экспорте.
"""

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import pymupdf
from pptx import Presentation

from app.pipeline.export.render import pptx_to_pdf

LAYOUT_WIDTH_PX = 1280
SLIDE_WIDTH_PX = 960
JPEG_QUALITY = 82


def layout_image_name(key: str) -> str:
    return key.replace(":", "-") + ".jpg"


def slide_image_name(index: int) -> str:
    return f"{index + 1}.jpg"


@dataclass
class TemplateRenders:
    slides: dict[int, str] = field(default_factory=dict)  # индекс слайда → файл в slides_dir
    layouts: dict[str, str] = field(default_factory=dict)  # ключ макета → файл в layouts_dir


def _save(page, path: Path, width_px: int) -> None:
    zoom = width_px / page.rect.width
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
    path.write_bytes(pixmap.tobytes("jpg", jpg_quality=JPEG_QUALITY))


def render_template(template_path: str | Path, slides_dir: Path, layouts_dir: Path) -> TemplateRenders:
    """Слайды шаблона как есть + по слайду на каждый макет без плейсхолдеров → JPEG."""
    prs = Presentation(str(template_path))
    slides_count = len(prs.slides)
    keys: list[str] = []
    for mi, master in enumerate(prs.slide_masters):
        for li, layout in enumerate(master.slide_layouts):
            slide = prs.slides.add_slide(layout)
            for placeholder in list(slide.placeholders):
                placeholder._element.getparent().remove(placeholder._element)
            keys.append(f"{mi}:{li}")

    buffer = BytesIO()
    prs.save(buffer)
    pdf = pptx_to_pdf(buffer.getvalue())

    slides_dir.mkdir(parents=True, exist_ok=True)
    layouts_dir.mkdir(parents=True, exist_ok=True)
    renders = TemplateRenders()
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        for number, page in enumerate(doc):
            if number < slides_count:
                name = slide_image_name(number)
                _save(page, slides_dir / name, SLIDE_WIDTH_PX)
                renders.slides[number] = name
            elif number - slides_count < len(keys):
                key = keys[number - slides_count]
                name = layout_image_name(key)
                _save(page, layouts_dir / name, LAYOUT_WIDTH_PX)
                renders.layouts[key] = name
    return renders
