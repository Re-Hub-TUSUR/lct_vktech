"""Паттерны слайдов: разбор слайдов-примеров шаблона на слоты для генерации.

Дизайн корпоративных шаблонов обычно живёт не в макетах, а в слайдах-примерах:
карточки из нескольких фигур, KPI с крупными числами, таймлайны, сетки. Этот модуль
превращает каждый слайд-пример в паттерн: тип, слоты с ролями и ёмкостью,
повторяющиеся элементы (карточки) со всеми их фигурами. Генерация клонирует
слайд-пример и заполняет слоты — так результат выглядит как сам шаблон.

Всё детерминированно и не зависит от конкретного шаблона: роли определяются по
геометрии, кеглю, повторяемости и характеру текста-заглушки. Семантику можно
уточнить VLM (см. app/pipeline/ai/vlm.py) — модель лишь размечает найденные фигуры.
"""

import re
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Any

from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
EMU_PER_PT = 12700

# Текст-заглушка шаблонов: такие фигуры — слоты, а не содержательный декор.
FILLER_RX = re.compile(
    r"^(заголовок|подзаголовок|основной текст|текстовый блок|текст|описание|дата|показатель|пункт|"
    r"имя фамилия|имя|фамилия|должность|подпись|название|заметка|примечание|кейс|содержание|"
    r"описание показателя|данные показателя|текст описания|вставить|вставьте|иллюстрация|изображение|"
    r"фото|картинка|ссылка|qr|lorem ipsum|title|subtitle|text|description|date|name|caption|insert|"
    r"image|picture|photo|link)\b",
    re.IGNORECASE,
)
# «ххх%», «XX» — метки для цифр в тексте-заглушке.
MASK_RX = re.compile(r"(?<!\w)[хx]{2,}(?!\w)", re.IGNORECASE)
NUMERIC_RX = re.compile(
    r"^[\s×x~<>+\-–≈]*(\d[\d\s.,]*|[хx]{2,})\s*(%|п\.\s?п\.|млн|млрд|тыс|ч|мин|₽|\$|раз)?[\s\w.]{0,12}\*?$",
    re.IGNORECASE,
)
NUMBERING_RX = re.compile(r"^0?\d{1,2}[.)]?$")
TITLE_LIKE_RX = re.compile(r"^(заголовок|название|title|heading)", re.IGNORECASE)
AGENDA_RX = re.compile(r"оглавлен|содержан|agenda|contents|план", re.IGNORECASE)
DATE_RX = re.compile(r"^(дата|date|q[1-4]|\d{4}|шаг|этап|step|\d{1,2}\.\d{2})", re.IGNORECASE)

# Назначение слайда-примера: под контент идут только category=content.
CLOSING_RX = re.compile(r"спасибо|благодар|thank|вопрос|q\s*&\s*a|call to action|контакт", re.IGNORECASE)
COVER_RX = re.compile(r"название презентации|тем[аyу] презентации|presentation title", re.IGNORECASE)
SECTION_RX = re.compile(r"разделител|раздел\b|section", re.IGNORECASE)
PROFILE_RX = re.compile(r"имя\s+фамили|имя спикера|должност|визитк|спикер|speaker", re.IGNORECASE)
QR_RX = re.compile(r"\bqr\b", re.IGNORECASE)
CODE_RX = re.compile(r"[{}]|;\s*(\||$)|</?[a-z]+>")
QUOTE_LAYOUT_RX = re.compile(r"цитат|quote", re.IGNORECASE)
# Подпись места под картинку: «Иллюстрация», «Вставить фото», «QR-code».
IMAGE_HINT_RX = re.compile(
    r"^(иллюстрац|изображени|фото|картинк|скриншот|вставить|вставьте|insert|image|picture|photo|screenshot|qr)",
    re.IGNORECASE,
)

MAX_ITEMS = 8
MAX_TEXT_SHAPES = 30
# Больше стольких мелких свободных картинок и фигур — это витрина иконок или коллаж, а не макет слайда.
MAX_LOOSE_DECOR = 12

TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
# Колонтитулы, номер слайда и дата — оформление шаблона, а не слоты: сборка их не трогает.
CHROME_PLACEHOLDERS = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}

# Типы паттернов, которые умеет заполнять генерация.
USABLE_TYPES = {"text", "cards", "two_columns", "kpi", "timeline", "table", "image_text", "agenda", "chart"}
CATEGORIES = ("content", "cover", "section", "closing", "profile", "code", "quote", "library")
# Служебные слайды-примеры (обложка, разделитель, финал) годятся для слайдов своего типа.
SERVICE_CATEGORIES = {"cover", "section", "closing"}
# Тип диаграммы шаблона → тип диаграммы слайда (ChartData.type). Остальные типы не заполняем.
CHART_KINDS = {
    XL_CHART_TYPE.COLUMN_CLUSTERED: "bar",
    XL_CHART_TYPE.COLUMN_STACKED: "bar",
    XL_CHART_TYPE.BAR_CLUSTERED: "horizontal_bar",
    XL_CHART_TYPE.BAR_STACKED: "horizontal_bar",
    XL_CHART_TYPE.LINE: "line",
    XL_CHART_TYPE.LINE_MARKERS: "line",
    XL_CHART_TYPE.PIE: "pie",
    XL_CHART_TYPE.PIE_EXPLODED: "pie",
    XL_CHART_TYPE.DOUGHNUT: "donut",
    XL_CHART_TYPE.DOUGHNUT_EXPLODED: "donut",
}


@dataclass
class Leaf:
    id: int
    kind: str  # title | text | picture | table | chart | line | shape | media
    bbox: tuple[float, float, float, float]  # доли слайда: x, y, w, h
    text: str = ""
    paragraphs: list[str] = field(default_factory=list)
    size: float | None = None
    bold: bool = False
    top: int = 0  # фигура верхнего уровня (сама фигура или её группа)
    grouped: bool = False
    shape_type: str = "rect"  # circle | rect | badge | pill | line
    color: str | None = None
    border: bool = False

    @property
    def cx(self) -> float:
        return self.bbox[0] + self.bbox[2] / 2

    @property
    def cy(self) -> float:
        return self.bbox[1] + self.bbox[3] / 2

    @property
    def area(self) -> float:
        return self.bbox[2] * self.bbox[3]

    @property
    def numeric(self) -> bool:
        return bool(self.text) and bool(NUMERIC_RX.match(self.text.split("|")[0].strip()))

    @property
    def filler(self) -> bool:
        text = self.text.strip()
        return bool(FILLER_RX.match(text) or MASK_RX.search(text)) or "lorem ipsum" in text.lower()


def _font_size(shape) -> float | None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                return round(run.font.size.pt, 1)
        props = paragraph._p.find(f"{{{NS_A}}}pPr/{{{NS_A}}}defRPr")
        if props is not None and (size := props.get("sz")):
            return int(size) / 100
    return None


def _bold(shape) -> bool:
    paragraphs = shape.text_frame.paragraphs
    return bool(paragraphs) and any(run.font.bold for run in paragraphs[0].runs)


def _kind(shape) -> str:
    if shape.is_placeholder and shape.placeholder_format.type in TITLE_PLACEHOLDERS:
        return "title"
    if shape.is_placeholder and shape.placeholder_format.type in CHROME_PLACEHOLDERS:
        return "chrome"
    shape_type = shape.shape_type
    if shape_type == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE):
        return "picture"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if shape_type in (MSO_SHAPE_TYPE.LINE,) or shape._element.tag.endswith("}cxnSp"):
        return "line"
    if shape_type in (MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
        return "media"
    if shape.has_text_frame and shape.text_frame.text.strip():
        return "text"
    return "shape"


def flatten(
    shapes,
    width: int,
    height: int,
    transform=(0.0, 0.0, 1.0, 1.0),
    top: int | None = None,
    theme_colors: dict[str, str] | None = None,
):
    """Все фигуры слайда с абсолютной геометрией (группы раскрываются)."""
    ox, oy, sx, sy = transform
    for shape in shapes:
        if shape.left is None or shape.top is None:
            continue
        left, top_px = ox + shape.left * sx, oy + shape.top * sy
        w, h = (shape.width or 0) * sx, (shape.height or 0) * sy
        owner = top if top is not None else shape.shape_id
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.grpSpPr.find(f"{{{NS_A}}}xfrm")
            ch_off = xfrm.find(f"{{{NS_A}}}chOff") if xfrm is not None else None
            ch_ext = xfrm.find(f"{{{NS_A}}}chExt") if xfrm is not None else None
            if ch_off is None or ch_ext is None:
                continue
            cw = int(ch_ext.get("cx")) or 1
            ch = int(ch_ext.get("cy")) or 1
            gsx, gsy = w / cw, h / ch
            inner = (left - int(ch_off.get("x")) * gsx, top_px - int(ch_off.get("y")) * gsy, gsx, gsy)
            yield from flatten(shape.shapes, width, height, inner, owner, theme_colors=theme_colors)
            continue
        kind = _kind(shape)
        if kind == "table":
            # Рамка graphicFrame в шаблонах бывает уже самой таблицы: размер — по столбцам и строкам.
            grid_width = sum(int(col.get("w", "0")) for col in shape._element.iter(f"{{{NS_A}}}gridCol"))
            rows_height = sum(int(row.get("h", "0")) for row in shape._element.iter(f"{{{NS_A}}}tr"))
            w, h = max(w, grid_width * sx), max(h, rows_height * sy)

        shape_type = "rect"
        ast = None
        try:
            ast = shape.auto_shape_type
        except Exception:
            pass

        if ast is not None:
            ast_name = str(ast).upper()
            if "OVAL" in ast_name:
                shape_type = "circle"
            elif "ROUNDED" in ast_name:
                shape_type = "badge"
            elif "RECT" in ast_name:
                shape_type = "rect"
            elif "LINE" in ast_name or "ARROW" in ast_name:
                shape_type = "line"
        elif kind == "line" or shape.shape_type == MSO_SHAPE_TYPE.LINE or getattr(shape._element, "tag", "").endswith("}cxnSp"):
            shape_type = "line"

        color = None
        border = False
        spPr = getattr(shape._element, "spPr", None)
        if spPr is not None:
            fill = spPr.find(f"{{{NS_A}}}solidFill")
            if fill is not None and theme_colors:
                for child in fill:
                    from app.pipeline.parsing.template_parser import _resolve_color
                    c = _resolve_color(child, theme_colors, {})
                    if c:
                        color = c
                        break
            ln = spPr.find(f"{{{NS_A}}}ln")
            if ln is not None and ln.find(f"{{{NS_A}}}noFill") is None:
                border = True

        leaf = Leaf(
            id=shape.shape_id,
            kind=kind,
            bbox=(left / width, top_px / height, w / width, h / height),
            top=owner,
            grouped=top is not None,
            shape_type=shape_type,
            color=color,
            border=border,
        )
        if shape.has_text_frame:
            leaf.paragraphs = [p.text.replace("\x0b", " ").strip() for p in shape.text_frame.paragraphs]
            leaf.paragraphs = [p for p in leaf.paragraphs if p]
            leaf.text = " | ".join(leaf.paragraphs)
            leaf.size = _font_size(shape)
            leaf.bold = _bold(shape)
        yield leaf


def _capacity(leaf: Leaf, width: int, height: int) -> int:
    """Сколько знаков помещается в рамку при кегле слота."""
    size_emu = (leaf.size or 18) * EMU_PER_PT
    area = leaf.bbox[2] * width * leaf.bbox[3] * height
    return max(8, int(area / (size_emu * size_emu * 0.55 * 1.25)))


def _slot(leaf: Leaf, role: str, width: int, height: int) -> dict[str, Any]:
    return {
        "id": leaf.id,
        "role": role,
        "bbox": [round(v, 4) for v in leaf.bbox],
        "size": leaf.size,
        "chars": _capacity(leaf, width, height),
        "paragraphs": len(leaf.paragraphs),
        "sample": leaf.text[:60],
    }


def _family_key(leaf: Leaf) -> tuple:
    return (round(leaf.size or 0), leaf.bold, leaf.numeric, len(leaf.paragraphs) > 1)


def _shape_key(leaf: Leaf) -> tuple:
    # Подложка может быть и фигурой, и картинкой (в одной серии встречается и то и другое).
    return ("line" if leaf.kind == "line" else "visual",)


def _similar(leaf: Leaf, seed: Leaf, height_tolerance: float, width: bool) -> bool:
    """Близкий размер рамки. Текст сравниваем по высоте с допуском (она зависит от длины
    текста-примера), ширину не учитываем: у чисел «7» и «10» и у крайней карточки она своя."""
    h, h0 = leaf.bbox[3], seed.bbox[3]
    if abs(h - h0) > max(0.004, h0 * height_tolerance):
        return False
    return not width or abs(leaf.bbox[2] - seed.bbox[2]) <= max(0.012, seed.bbox[2] * 0.08)


def _reading_order(leaves: list[Leaf]) -> list[Leaf]:
    return sorted(leaves, key=lambda leaf: (round(leaf.bbox[1] * 12), leaf.bbox[0]))


def _families(leaves: list[Leaf], key, height_tolerance: float = 0.25, width: bool = False) -> list[list[Leaf]]:
    """Серии похожих фигур: общий ключ (кегль, начертание, вид) и близкий размер рамки."""
    groups: dict[tuple, list[Leaf]] = defaultdict(list)
    for leaf in leaves:
        groups[key(leaf)].append(leaf)
    families: list[list[Leaf]] = []
    for members in groups.values():
        clusters: list[list[Leaf]] = []
        for leaf in sorted(members, key=lambda leaf: (leaf.bbox[3], leaf.bbox[2])):
            cluster = next(
                (cluster for cluster in clusters if _similar(leaf, cluster[0], height_tolerance, width)), None
            )
            if cluster is None:
                clusters.append([leaf])
            else:
                cluster.append(leaf)
        families.extend(cluster for cluster in clusters if len(cluster) >= 2)
    return families


def _sequence(family: list[Leaf]) -> bool:
    """Нумерация пунктов: 1, 2, 3… (или 01, 02…) в порядке чтения."""
    texts = [leaf.text.strip().rstrip(".)") for leaf in _reading_order(family)]
    if not all(text.isdigit() for text in texts):
        return False
    numbers = [int(text) for text in texts]
    return numbers[0] in (0, 1) and numbers == list(range(numbers[0], numbers[0] + len(numbers)))


def _attached_decor(leaf: Leaf, decor: list[Leaf], texts: list[Leaf], *, image: bool = False) -> list[int]:
    """Оформление одиночного текста, которое уходит вместе с ним: маркер строки, иконка внутри,
    подложка-плашка под текстом (для места под картинку — рамка любого размера) и всё, что на ней."""
    x, y, w, h = leaf.bbox
    links: list[int] = []
    containers: list[Leaf] = []
    for other in decor:
        ox, oy, ow, oh = other.bbox
        inside = x - 0.005 <= ox and ox + ow <= x + w + 0.005 and y - 0.005 <= oy and oy + oh <= y + h + 0.005
        marker = other.area < 0.004 and y - 0.01 <= other.cy <= y + h + 0.01 and x - 0.06 <= other.cx <= x + 0.01
        if inside or marker:
            links.append(other.id)
            continue
        covers = ox <= leaf.cx <= ox + ow and oy <= leaf.cy <= oy + oh
        # Подложка только этого текста: не шире трёх его площадей и без других текстов на ней.
        if covers and other.kind != "line" and (image or other.area <= max(3 * leaf.area, 0.01)):
            shared = any(
                text is not leaf and ox <= text.cx <= ox + ow and oy <= text.cy <= oy + oh for text in texts
            )
            if not shared:
                containers.append(other)
    for box in containers:
        bx, by, bw, bh = box.bbox
        links.append(box.id)
        links += [
            other.id
            for other in decor
            if other.id not in links and other.area < box.area and bx <= other.cx <= bx + bw and by <= other.cy <= by + bh
        ]
    return links


def _attach(anchors: list[Leaf], family: list[Leaf]) -> list[list[Leaf]] | None:
    """Каждому якорю — k ближайших свободных членов семейства (k = размер семейства / число якорей)."""
    if not anchors or len(family) % len(anchors):
        return None
    per_item = len(family) // len(anchors)
    if per_item > 4:
        return None
    free = list(family)
    result: list[list[Leaf]] = [[] for _ in anchors]
    for _ in range(per_item):
        for i, anchor in enumerate(anchors):
            best = min(free, key=lambda leaf: (leaf.cx - anchor.cx) ** 2 + (leaf.cy - anchor.cy) ** 2)
            free.remove(best)
            result[i].append(best)
    return [sorted(group, key=lambda leaf: (leaf.bbox[1], leaf.bbox[0])) for group in result]


def _companions(texts: list[Leaf], markers: list[Leaf]) -> dict[int, list[int]]:
    """Маркеры списка и иконки строк: мелкие фигуры слева от текста на его высоте. Удаляются вместе со строкой."""
    result: dict[int, list[int]] = defaultdict(list)
    for marker in markers:
        near = [
            leaf
            for leaf in texts
            if leaf.bbox[1] - 0.01 <= marker.cy <= leaf.bbox[1] + leaf.bbox[3] + 0.01
            and leaf.bbox[0] - 0.06 <= marker.cx <= leaf.bbox[0] + 0.01
        ]
        if near:
            owner = min(near, key=lambda leaf: abs(leaf.cy - marker.cy))
            result[owner.id].append(marker.id)
    return result


def _item_role(family: list[Leaf], largest_size: float, per_item: int = 1) -> str:
    sample = family[0]
    texts = [leaf.text for leaf in family]
    if per_item > 1:
        return "lines"
    # Крупные числа — показатели (KPI), «1, 2, 3» по порядку и мелкие «01», «02» — нумерация пунктов.
    if _sequence(family) or (all(NUMBERING_RX.match(text) for text in texts) and (sample.size or 0) < 30):
        return "number"
    if sample.numeric:
        return "value"
    if all(DATE_RX.match(text) for text in texts):
        return "date"
    if len(sample.paragraphs) > 1:
        return "title_text"
    if (sample.size or 0) >= largest_size:
        return "title"
    return "text"


def _category(index: int, title: Leaf | None, leaves: list[Leaf], layout_name: str | None) -> str:
    """Назначение слайда-примера. Обложки, финалы, визитки, код и цитаты не годятся под обычный контент."""
    title_text = title.text if title is not None else ""
    texts = [leaf.text for leaf in leaves if leaf.text]
    if index == 0 or COVER_RX.search(title_text):
        return "cover"
    if CLOSING_RX.search(title_text) or any(QR_RX.search(text) for text in texts):
        return "closing"
    if any(PROFILE_RX.search(text) for text in texts):
        return "profile"
    if any(CODE_RX.search(text) for text in texts):
        return "code"
    if layout_name and QUOTE_LAYOUT_RX.search(layout_name):
        return "quote"
    if SECTION_RX.search(title_text):
        return "section"
    return "content"


def shape_type(shape):
    """Тип фигуры или None: на экзотических фигурах python-pptx бросает NotImplementedError."""
    try:
        return shape.shape_type
    except Exception:  # noqa: BLE001
        return None


def _walk(shapes):
    """Все фигуры слайда, включая вложенные в группы."""
    for shape in shapes:
        yield shape
        if shape_type(shape) == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)


def chart_meta(shape) -> dict[str, Any]:
    """Что за диаграмма в слайде-примере: тип, число рядов и категорий, легенда, заголовок."""
    chart = shape.chart
    try:
        chart_type = chart.chart_type
    except Exception:  # noqa: BLE001 — экзотические типы python-pptx не распознаёт
        chart_type = None
    plots = list(chart.plots)
    series = sum(len(list(plot.series)) for plot in plots)
    categories = len(list(plots[0].categories)) if plots else 0
    return {
        "kind": CHART_KINDS.get(chart_type, "other"),
        "plots": len(plots),
        "series": series,
        "categories": categories,
        "has_legend": bool(chart.has_legend),
        "has_title": bool(chart.has_title),
    }


def is_usable(pattern: dict[str, Any]) -> bool:
    if pattern["type"] == "chart":
        chart = pattern["frames"].get("chart") or {}
        # Диаграмму подменяем данными только в «чистом» примере: одна диаграмма одного типа и без карточек.
        simple = (
            chart.get("kind", "other") != "other"
            and chart.get("plots") == 1
            and pattern.get("charts_count") == 1
        )
        if not simple or pattern.get("items_count"):
            return False
    return (
        pattern["type"] in USABLE_TYPES
        and pattern.get("title") is not None
        and pattern.get("category", "content") == "content"
    )


def is_service(pattern: dict[str, Any]) -> bool:
    """Обложка, разделитель или финал, которые можно клонировать для слайдов этих типов."""
    return (
        pattern.get("category") in SERVICE_CATEGORIES
        and pattern.get("title") is not None
        and (pattern["title"].get("chars") or 0) >= 16
        and not pattern.get("items")
        and not pattern.get("frames")
        and not any(slot["role"] == "image_hint" for slot in pattern.get("slots") or [])
        and pattern.get("type") in {"title_only", "text"}
    )


def _min_items(count: int) -> int:
    """Сколько элементов минимум нужно паттерну: лишних не больше половины (как в подборе паттерна)."""
    return next((n for n in range(1, count + 1) if count - n <= max(1, n // 2)), count)


def pattern_schema(pattern: dict[str, Any]) -> dict[str, Any]:
    """Схема слотов для генерации текста: что писать и сколько знаков помещается. Её получит LLM,
    чтобы сразу писать текст под выбранный паттерн, а не подгонять его потом."""
    schema: dict[str, Any] = {"type": pattern["type"]}
    if pattern.get("title"):
        schema["title"] = pattern["title"]["chars"]
    for slot in pattern.get("slots") or []:
        if slot["role"] in {"body", "subtitle", "value", "value_label"} and slot["role"] not in schema:
            schema[slot["role"]] = {"chars": slot["chars"], "paragraphs": max(1, slot.get("paragraphs") or 1)}
    items = pattern.get("items") or []
    if items:
        fields: dict[str, int] = {}
        for slot in items[0]["slots"]:
            role = "text" if slot["role"] == "lines" else slot["role"]
            fields[role] = fields.get(role, 0) + slot["chars"]
        schema["items"] = {"min": _min_items(len(items)), "max": len(items), "fields": fields}
    chart = pattern.get("frames", {}).get("chart")
    if chart:
        schema["chart"] = {"kind": chart.get("kind"), "series": chart.get("series"), "categories": chart.get("categories")}
    if pattern.get("frames", {}).get("table"):
        schema["table"] = True
    return schema


def analyze_slide(
    slide,
    index: int,
    layout_key: str | None,
    width: int,
    height: int,
    layout_name: str | None = None,
    theme_colors: dict[str, str] | None = None,
) -> dict[str, Any]:
    leaves = list(flatten(slide.shapes, width, height, theme_colors=theme_colors))
    texts = [leaf for leaf in leaves if leaf.kind == "text"]
    pattern: dict[str, Any] = {
        "id": f"s{index + 1}",
        "slide_index": index,
        "layout_key": layout_key,
        "type": "complex",
        "category": "content",
        "usable": False,
        "items_count": 0,
        "title": None,
        "slots": [],
        "items": [],
        "shapes": [],
        "frames": {},
        "images": [],
        "source": "rules",
        # Значения на случай раннего выхода (слишком сложный слайд): дальше они уточняются.
        "service": False,
        "schema": {"type": "complex"},
    }

    title = next((leaf for leaf in leaves if leaf.kind == "title"), None)
    if title is None:
        # Заголовок, нарисованный текстовым полем: крупный текст в верхней трети.
        top_texts = [leaf for leaf in texts if leaf.bbox[1] < 0.3 and not leaf.numeric and (leaf.size or 0) >= 20]
        title = max(top_texts, key=lambda leaf: leaf.size or 0, default=None)
        if title is not None:
            texts.remove(title)
    if title is not None:
        pattern["title"] = _slot(title, "title", width, height)
    pattern["category"] = _category(index, title, leaves, layout_name)

    shapes_by_id = {shape.shape_id: shape for shape in _walk(slide.shapes)}
    for leaf in leaves:
        if leaf.kind in {"table", "chart"}:
            frame: dict[str, Any] = {"id": leaf.id, "bbox": [round(v, 4) for v in leaf.bbox]}
            if leaf.kind == "chart" and leaf.id in shapes_by_id:
                try:
                    frame.update(chart_meta(shapes_by_id[leaf.id]))
                except Exception:  # noqa: BLE001 — повреждённая диаграмма не должна ронять разбор
                    frame["kind"] = "other"
            pattern["frames"][leaf.kind] = frame
    pattern["charts_count"] = sum(1 for leaf in leaves if leaf.kind == "chart")
    if len(texts) > MAX_TEXT_SHAPES or any(leaf.kind == "media" for leaf in leaves):
        return pattern

    # --- Повторяющиеся элементы (карточки, колонки, шаги, KPI) ---
    families = [f for f in _families(texts, _family_key) if len(f) <= MAX_ITEMS * 4]
    words = [
        f
        for f in families
        if len(f) <= MAX_ITEMS and not f[0].numeric and not all(NUMBERING_RX.match(leaf.text) for leaf in f)
    ]
    # Якорь элементов — семейство заголовков карточек: «Заголовок»-заглушка, затем крупнее кегль, затем больше членов.
    anchor_family = max(
        words or [f for f in families if len(f) <= MAX_ITEMS],
        key=lambda f: (
            all(TITLE_LIKE_RX.match(leaf.text) for leaf in f),
            max(leaf.size or 0 for leaf in f),
            len(f),
        ),
        default=None,
    )
    items: list[dict[str, Any]] = []
    used: set[int] = set()
    if anchor_family is not None:
        anchors = _reading_order(anchor_family)
        item_families = [anchor_family] + [
            f
            for f in families
            if f is not anchor_family and len(f) % len(anchors) == 0 and len(f) // len(anchors) <= 4
        ]
        largest = max(leaf.size or 0 for f in item_families for leaf in f)
        members: list[list[tuple[Leaf, str]]] = [[] for _ in anchors]
        for family in item_families:
            attached = [[leaf] for leaf in anchors] if family is anchor_family else _attach(anchors, family)
            if attached is None:
                continue
            role = _item_role(family, largest, len(family) // len(anchors))
            for i, group in enumerate(attached):
                for leaf in group:
                    members[i].append((leaf, role))
                    used.add(leaf.id)
        # Декор элементов: фон карточки, иконки, точки таймлайна — фигуры, повторяющиеся столько же раз.
        others = [leaf for leaf in leaves if leaf.kind in {"shape", "picture", "line"}]
        decor_members: list[list[Leaf]] = [[] for _ in anchors]
        for family in _families(others, _shape_key, height_tolerance=0.08, width=True):
            attached = _attach(anchors, family)
            if attached is None:
                continue
            for i, group in enumerate(attached):
                decor_members[i].extend(group)
                used.update(leaf.id for leaf in group)
        for i, anchor in enumerate(anchors):
            slots = members[i]
            shapes = [leaf for leaf, _ in slots] + decor_members[i]
            tops = {leaf.top for leaf in shapes if leaf.grouped}
            box = _union([leaf.bbox for leaf in shapes])
            photo = next((leaf for leaf in decor_members[i] if leaf.kind == "picture" and leaf.area > 0.015), None)
            markers = [leaf for leaf in decor_members[i] if leaf.area < 0.004]
            companions = _companions([leaf for leaf, _ in slots], markers)
            item_slots = [_slot(leaf, role, width, height) for leaf, role in slots]
            for slot in item_slots:
                slot["companions"] = companions.get(slot["id"], [])
            items.append(
                {
                    "bbox": box,
                    # Если элемент целиком лежит в своей группе — удаляем группу целиком.
                    "group": tops.pop() if len(tops) == 1 and all(leaf.grouped for leaf in shapes) else None,
                    "shapes": sorted({leaf.id for leaf in shapes}),
                    "slots": item_slots,
                    "image": photo.id if photo else None,
                }
            )

    # --- Одиночные слоты ---
    singles = [leaf for leaf in texts if leaf.id not in used]
    hints = [leaf for leaf in singles if IMAGE_HINT_RX.match(leaf.text)]
    singles = [leaf for leaf in singles if leaf not in hints]
    roles: list[tuple[Leaf, str]] = []
    values = [leaf for leaf in singles if leaf.numeric and (leaf.size or 0) >= 36]
    for value in values:
        roles.append((value, "value"))
        below = [
            leaf
            for leaf in singles
            if leaf not in values and leaf.bbox[1] >= value.bbox[1] and abs(leaf.bbox[0] - value.bbox[0]) < 0.08
        ]
        if below:
            label = min(below, key=lambda leaf: leaf.bbox[1])
            roles.append((label, "value_label"))
            singles.remove(label)
    rest = [leaf for leaf in singles if leaf not in values]
    # Основной текст — самая большая рамка обычного кегля (54 pt — это уже вынос, а не абзац).
    body = max(
        (leaf for leaf in rest if leaf.area >= 0.03 and (leaf.size or 18) <= 28), key=lambda leaf: leaf.area, default=None
    )
    if body is not None:
        roles.append((body, "body"))
        rest.remove(body)
    if title is not None:
        title_bottom = title.bbox[1] + title.bbox[3]
        near = [
            leaf
            for leaf in rest
            if 0 <= leaf.bbox[1] - title_bottom < 0.12 and (leaf.size or 0) < (title.size or 99)
        ]
        if near:
            subtitle = min(near, key=lambda leaf: leaf.bbox[1])
            roles.append((subtitle, "subtitle"))
            rest.remove(subtitle)
    for leaf in rest:
        # Оставшиеся тексты при генерации удаляются: заглушки — как extra, текст-пример — как decor_text.
        roles.append((leaf, "extra" if leaf.filler else "decor_text"))
    roles += [(leaf, "image_hint") for leaf in hints]

    free_decor = [
        leaf for leaf in leaves if leaf.kind in {"shape", "picture", "line"} and leaf.id not in used and leaf.area < 0.35
    ]
    all_texts = texts + ([title] if title is not None else [])
    slots: list[dict[str, Any]] = []
    for leaf, role in roles:
        slot = _slot(leaf, role, width, height)
        slot["companions"] = _attached_decor(leaf, free_decor, all_texts, image=role == "image_hint")
        slots.append(slot)

    loose = [leaf for leaf in leaves if leaf.kind in {"picture", "shape"} and leaf.id not in used and leaf.area < 0.01]
    if pattern["category"] == "content" and len(loose) > MAX_LOOSE_DECOR:
        pattern["category"] = "library"
    images = [leaf for leaf in leaves if leaf.kind == "picture" and leaf.id not in used and leaf.area >= 0.08]
    pattern["images"] = [{"id": leaf.id, "bbox": [round(v, 4) for v in leaf.bbox]} for leaf in images]
    pattern["slots"] = slots
    pattern["items"] = items
    pattern["items_count"] = len(items)

    shapes: list[dict[str, Any]] = []
    for leaf in leaves:
        if leaf.kind in {"shape", "line"}:
            if leaf.area >= 0.85:
                continue
            shapes.append({
                "id": f"tpl-{leaf.id}",
                "template_shape_id": leaf.id,
                "kind": leaf.shape_type,
                "x": round(leaf.bbox[0], 4),
                "y": round(leaf.bbox[1], 4),
                "width": round(leaf.bbox[2], 4),
                "height": round(leaf.bbox[3], 4),
                "color": leaf.color,
                "border": leaf.border,
                "text": leaf.text if leaf.text else None,
            })
    pattern["shapes"] = shapes
    pattern["type"] = _classify(pattern, leaves, items, used, title)
    pattern["usable"] = is_usable(pattern)
    pattern["service"] = is_service(pattern)
    pattern["schema"] = pattern_schema(pattern)
    return pattern


def _union(boxes: list[tuple[float, float, float, float]]) -> list[float]:
    x0 = min(b[0] for b in boxes)
    y0 = min(b[1] for b in boxes)
    x1 = max(b[0] + b[2] for b in boxes)
    y1 = max(b[1] + b[3] for b in boxes)
    return [round(x0, 4), round(y0, 4), round(x1 - x0, 4), round(y1 - y0, 4)]


def _classify(
    pattern: dict[str, Any], leaves: list[Leaf], items: list[dict[str, Any]], used: set[int], title: Leaf | None
) -> str:
    roles = {slot["role"] for item in items for slot in item["slots"]}
    single_roles = {slot["role"] for slot in pattern["slots"]}
    if pattern.get("charts_count", 0) > 1:
        return "complex"
    if "table" in pattern["frames"]:
        return "table"
    if "chart" in pattern["frames"]:
        return "chart"
    # Схема: блоки, соединённые линиями, — под неё нет содержательной модели.
    connectors = [leaf for leaf in leaves if leaf.kind == "line" and leaf.id not in used]
    if len(connectors) >= 4:
        return "diagram"
    if len(items) >= 2:
        # Оглавление узнаём по заголовку слайда («Содержание», «Оглавление», «План»), а не по цифрам.
        if title is not None and AGENDA_RX.search(title.text):
            return "agenda"
        if "value" in roles:
            return "kpi"
        column = max(item["bbox"][0] for item in items) - min(item["bbox"][0] for item in items) < 0.05
        if "number" in roles and column:
            return "agenda"
        spans = [
            leaf
            for leaf in leaves
            if leaf.kind in {"line", "shape"} and leaf.bbox[2] >= 0.6 and leaf.bbox[3] <= 0.03
        ]
        if "date" in roles or ("number" in roles and len(items) >= 3) or spans:
            return "timeline"
        if any(item["image"] for item in items):
            return "people"
        if not roles & {"title", "title_text", "text"}:
            return "complex"
        return "two_columns" if len(items) == 2 else "cards"
    if "value" in single_roles:
        return "kpi"
    if "body" in single_roles:
        return "image_text" if pattern["images"] else "text"
    return "title_only"


def extract_patterns(
    prs,
    slide_layout_keys: list[str | None],
    layout_names: dict[str, str] | None = None,
    theme_colors: dict[str, str] | None = None,
) -> list[dict[str, Any]]:
    width, height = int(prs.slide_width), int(prs.slide_height)
    names = layout_names or {}
    patterns = []
    for index, (slide, key) in enumerate(zip(prs.slides, slide_layout_keys)):
        try:
            patterns.append(
                analyze_slide(slide, index, key, width, height, names.get(key or ""), theme_colors=theme_colors)
            )
        except Exception as exc:  # noqa: BLE001 — один странный слайд не должен ронять разбор чужого шаблона
            patterns.append(broken_pattern(index, key, exc))
    return patterns


def broken_pattern(index: int, layout_key: str | None, exc: Exception) -> dict[str, Any]:
    return {
        "id": f"s{index + 1}",
        "slide_index": index,
        "layout_key": layout_key,
        "type": "complex",
        "category": "content",
        "usable": False,
        "service": False,
        "items_count": 0,
        "title": None,
        "slots": [],
        "items": [],
        "shapes": [],
        "frames": {},
        "images": [],
        "source": "rules",
        "schema": {"type": "complex"},
        "error": f"{type(exc).__name__}: {exc}"[:200],
    }


def pattern_summary(patterns: list[dict[str, Any]]) -> dict[str, int]:
    summary: dict[str, int] = defaultdict(int)
    for pattern in patterns:
        if pattern["usable"]:
            summary[pattern["type"]] += 1
    return dict(summary)

