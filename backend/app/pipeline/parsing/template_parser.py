"""Декомпозиция .pptx-шаблона: тема, палитра, шрифты, кегли, макеты и их роли,
паттерны слайдов-примеров и макетов, токены дизайн-системы.

Парсер не знает о конкретных шаблонах. Роли макетов определяются по тому, как
макеты используются в слайдах шаблона, по плейсхолдерам, их геометрии и
ключевым словам в названиях (паттерны — в skills/template_parser.yaml).
Шаблоны .potx, .pptm, .ppt, .odp и т. п. приводятся к .pptx до разбора (parsing/formats.py).
"""

import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from app.pipeline.geometry import content_frame, to_fractions
from app.pipeline.parsing import design_tokens
from app.pipeline.parsing.layout_patterns import extract_layout_patterns
from app.pipeline.parsing.patterns import extract_patterns, pattern_summary
from app.pipeline.skills import get_skill

PARSER_VERSION = "0.5.2"

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
URI_CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"
URI_TABLE = "http://schemas.openxmlformats.org/drawingml/2006/table"
URI_DIAGRAM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"

SCHEME_ALIASES = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
TITLE_TYPES = {"TITLE", "CENTER_TITLE"}
BODY_TYPES = {"BODY", "OBJECT", "SUBTITLE"}
PART_RX = re.compile(r"^ppt/(slides|slideLayouts|slideMasters)/[^/]+\.xml$")
EMU_PER_INCH = 914400
# С этого кегля (в сотых пункта) прогон считается заголовочным при определении гарнитуры заголовков.
HEADING_SZ = 2400

_XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, remove_comments=True)


@dataclass
class ParsedTemplate:
    slide_width: int
    slide_height: int
    slides_count: int
    tokens: dict[str, Any]
    layouts: list[dict[str, Any]]
    roles: dict[str, str | None]
    # Паттерны слайдов-примеров и макетов (см. parsing/patterns.py, parsing/layout_patterns.py).
    patterns: list[dict[str, Any]]

    @property
    def layouts_count(self) -> int:
        return len(self.layouts)


def parse_template(path: str | Path) -> ParsedTemplate:
    path = Path(path)
    params = get_skill("template_parser").params
    prs = Presentation(str(path))
    width, height = int(prs.slide_width), int(prs.slide_height)

    theme = extract_theme(prs)
    layouts, partname_to_key = _layouts(prs, theme, params)
    slide_layout_keys = [partname_to_key.get(str(s.slide_layout.part.partname)) for s in prs.slides]
    roles = _assign_roles(layouts, slide_layout_keys, width, params)
    for layout in layouts:
        layout["assigned_roles"] = sorted(role for role, key in roles.items() if key == layout["key"])

    usage = _usage(path, theme, params)
    patterns = extract_patterns(
        prs, slide_layout_keys, {layout["key"]: layout["name"] for layout in layouts}, theme_colors=theme.get("colors") or {}
    )
    # Паттерны из макетов: шаблон почти без слайдов-примеров (.potx) тоже даёт колонки, карточки, текст с картинкой.
    patterns += design_tokens.safe("layout_patterns", extract_layout_patterns, prs, layouts, roles) or []
    font_names = [font["name"] for font in usage["fonts"]] + [
        name for name in (theme.get("fonts") or {}).values() if name
    ]
    tokens = {
        "slide_size": {
            "width_emu": width,
            "height_emu": height,
            "width_in": round(width / EMU_PER_INCH, 3),
            "height_in": round(height / EMU_PER_INCH, 3),
            "aspect": _aspect_label(width, height),
        },
        "theme": theme,
        # Фон макетов по ролям — по нему выбирается цвет текста (тёмные шаблоны).
        "backgrounds": {
            role: next((layout["background"] for layout in layouts if layout["key"] == key), None)
            for role, key in roles.items()
            if key
        },
        "palette": usage["palette"],
        "fonts": usage["fonts"],
        "font_roles": usage["font_roles"],
        "type_scale": usage["type_scale"],
        # Токены дизайн-системы сверх палитры и шрифтов (parsing/design_tokens.py).
        "typography": design_tokens.safe("typography", design_tokens.typography, prs, theme, usage["type_scale"]),
        "grid": design_tokens.safe("grid", design_tokens.grid, prs, layouts, roles),
        "chrome": design_tokens.safe("chrome", design_tokens.chrome, prs),
        "table_style": design_tokens.safe("table_style", design_tokens.table_style, prs, theme),
        "chart_style": design_tokens.safe("chart_style", design_tokens.chart_style, prs, theme),
        "fonts_check": design_tokens.safe("fonts_check", design_tokens.fonts_check, font_names),
        "stats": {
            **usage["stats"],
            "slides": len(prs.slides),
            "layouts": len(layouts),
            "masters": len(prs.slide_masters),
            "patterns": sum(1 for pattern in patterns if pattern["usable"]),
            "layout_patterns": sum(1 for pattern in patterns if pattern.get("origin") == "layout"),
            "service_patterns": sum(1 for pattern in patterns if pattern.get("service")),
            "broken_slides": sum(1 for pattern in patterns if pattern.get("error")),
        },
        "patterns": pattern_summary(patterns),
    }
    return ParsedTemplate(
        slide_width=width,
        slide_height=height,
        slides_count=len(prs.slides),
        tokens=tokens,
        layouts=layouts,
        roles=roles,
        patterns=patterns,
    )


def _aspect_label(width: int, height: int) -> str:
    ratio = width / height
    for label, value in (("16:9", 16 / 9), ("16:10", 16 / 10), ("4:3", 4 / 3), ("1:1", 1.0)):
        if abs(ratio - value) < 0.02:
            return label
    return f"{ratio:.2f}:1"


def extract_theme(prs) -> dict[str, Any]:
    empty = {"name": None, "colors": {}, "fonts": {"major": None, "minor": None}}
    if not prs.slide_masters:
        return empty
    try:
        part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    except KeyError:
        return empty
    root = etree.fromstring(part.blob, _XML_PARSER)

    colors: dict[str, str] = {}
    scheme = root.find(f".//{{{NS_A}}}clrScheme")
    if scheme is not None:
        for el in scheme:
            if not isinstance(el.tag, str) or len(el) == 0:
                continue
            value_el = el[0]
            local = etree.QName(value_el).localname
            value = value_el.get("val") if local == "srgbClr" else value_el.get("lastClr")
            if value:
                colors[etree.QName(el).localname] = "#" + value.upper()

    major = root.find(f".//{{{NS_A}}}majorFont/{{{NS_A}}}latin")
    minor = root.find(f".//{{{NS_A}}}minorFont/{{{NS_A}}}latin")
    return {
        "name": root.get("name"),
        "colors": colors,
        "fonts": {
            "major": major.get("typeface") if major is not None else None,
            "minor": minor.get("typeface") if minor is not None else None,
        },
    }


def _resolve_color(element, theme_colors: dict[str, str], clr_map: dict[str, str]) -> str | None:
    local = etree.QName(element).localname
    if local == "srgbClr" and element.get("val"):
        return "#" + element.get("val").upper()
    if local == "sysClr":
        return "#" + (element.get("lastClr") or "FFFFFF").upper()
    if local == "schemeClr":
        value = element.get("val", "")
        # bg1/tx1/… отображаются на цвета темы через clrMap мастера.
        return theme_colors.get(clr_map.get(value) or SCHEME_ALIASES.get(value, value))
    return None


def _background(element, theme_colors: dict[str, str], clr_map: dict[str, str]) -> str | None:
    """Цвет фона слайда/макета/мастера: заливка, ссылка на стиль или первая точка градиента."""
    bg = element.find(f"{{{NS_P}}}cSld/{{{NS_P}}}bg")
    if bg is None:
        return None
    holders = (
        bg.find(f"{{{NS_P}}}bgPr/{{{NS_A}}}solidFill"),
        bg.find(f"{{{NS_P}}}bgRef"),
        bg.find(f"{{{NS_P}}}bgPr/{{{NS_A}}}gradFill/{{{NS_A}}}gsLst/{{{NS_A}}}gs"),
    )
    for holder in holders:
        if holder is None:
            continue
        for child in holder:
            if isinstance(child.tag, str) and (color := _resolve_color(child, theme_colors, clr_map)):
                return color
    return None


def layout_background(layout, theme_colors: dict[str, str]) -> str | None:
    master = layout.slide_master
    clr_map_el = master._element.find(f"{{{NS_P}}}clrMap")
    clr_map = dict(clr_map_el.attrib) if clr_map_el is not None else {}
    return (
        _background(layout._element, theme_colors, clr_map)
        or _background(master._element, theme_colors, clr_map)
        or theme_colors.get(clr_map.get("bg1", "lt1"))
    )


def _classify(name: str, placeholders: list[dict[str, Any]], params: dict[str, Any]) -> str:
    for role, pattern in (params.get("role_patterns") or {}).items():
        if re.search(pattern, name, re.IGNORECASE):
            return role
    return "content" if placeholders else "blank"


def _length(value: Any) -> int | None:
    return int(value) if value is not None else None


def _layouts(prs, theme: dict[str, Any], params: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, str]]:
    usage = Counter(str(slide.slide_layout.part.partname) for slide in prs.slides)
    layouts: list[dict[str, Any]] = []
    partname_to_key: dict[str, str] = {}
    for mi, master in enumerate(prs.slide_masters):
        for li, layout in enumerate(master.slide_layouts):
            key = f"{mi}:{li}"
            partname = str(layout.part.partname)
            partname_to_key[partname] = key
            placeholders = []
            for ph in layout.placeholders:
                try:
                    fmt = ph.placeholder_format
                    placeholders.append(
                        {
                            "type": fmt.type.name if fmt.type is not None else "OBJECT",
                            "idx": fmt.idx,
                            "name": ph.name,
                            "left": _length(ph.left),
                            "top": _length(ph.top),
                            "width": _length(ph.width),
                            "height": _length(ph.height),
                        }
                    )
                except Exception:  # noqa: BLE001 — нестандартный плейсхолдер пропускаем, а не роняем разбор
                    continue
            name = layout.name or f"Макет {li + 1}"
            try:
                background = layout_background(layout, theme["colors"])
            except Exception:  # noqa: BLE001
                background = None
            try:
                title_style = _title_style(layout, theme["colors"])
            except Exception:  # noqa: BLE001
                title_style = {"size_pt": None, "color": None, "anchor": None, "align": None, "bold": None}
            layouts.append(
                {
                    "key": key,
                    "master": mi,
                    "index": li,
                    "name": name,
                    "used_by": usage.get(partname, 0),
                    "role": _classify(name, placeholders, params),
                    "background": background,
                    "frames": _frames(placeholders, int(prs.slide_width), int(prs.slide_height)),
                    "title_style": title_style,
                    "placeholders": placeholders,
                }
            )
    return layouts, partname_to_key


def _rect(placeholder: dict[str, Any]) -> tuple[int, int, int, int] | None:
    dims = (placeholder["left"], placeholder["top"], placeholder["width"], placeholder["height"])
    return None if None in dims else dims  # type: ignore[return-value]


def _frames(placeholders: list[dict[str, Any]], width: int, height: int) -> dict[str, list[float] | None]:
    """Области макета в долях слайда: заголовок, подзаголовок (первое тело) и контент под заголовком."""
    title = next((r for p in placeholders if p["type"] in TITLE_TYPES and (r := _rect(p))), None)
    bodies = sorted((p for p in placeholders if p["type"] in BODY_TYPES and _rect(p)), key=lambda p: p.get("idx") or 0)
    return {
        "title": to_fractions(title, width, height),
        "subtitle": to_fractions(_rect(bodies[0]), width, height) if bodies else None,
        "body": to_fractions(content_frame(title, width, height), width, height),
    }


def _placeholder_sp(element, types: set[str]):
    for sp in element.iter(f"{{{NS_P}}}sp"):
        ph = sp.find(f"{{{NS_P}}}nvSpPr/{{{NS_P}}}nvPr/{{{NS_P}}}ph")
        if ph is not None and ph.get("type", "obj") in types:
            return sp
    return None


def _fill_color(props, theme_colors: dict[str, str], clr_map: dict[str, str]) -> str | None:
    fill = props.find(f"{{{NS_A}}}solidFill") if props is not None else None
    if fill is None:
        return None
    return next(
        (color for child in fill if isinstance(child.tag, str) and (color := _resolve_color(child, theme_colors, clr_map))),
        None,
    )


def _title_style(layout, theme_colors: dict[str, str]) -> dict[str, Any]:
    """Эффективный стиль заголовка по цепочке наследования: плейсхолдер макета → мастер → titleStyle."""
    master = layout.slide_master
    clr_map_el = master._element.find(f"{{{NS_P}}}clrMap")
    clr_map = dict(clr_map_el.attrib) if clr_map_el is not None else {}
    style: dict[str, Any] = {"size_pt": None, "color": None, "anchor": None, "align": None, "bold": None}

    levels = []
    for sp in (_placeholder_sp(layout._element, {"title", "ctrTitle"}), _placeholder_sp(master._element, {"title"})):
        if sp is None:
            continue
        body_pr = sp.find(f"{{{NS_P}}}txBody/{{{NS_A}}}bodyPr")
        if style["anchor"] is None and body_pr is not None:
            style["anchor"] = body_pr.get("anchor")
        levels.append(sp.find(f"{{{NS_P}}}txBody/{{{NS_A}}}lstStyle/{{{NS_A}}}lvl1pPr"))
    levels.append(master._element.find(f"{{{NS_P}}}txStyles/{{{NS_P}}}titleStyle/{{{NS_A}}}lvl1pPr"))

    for level in levels:
        if level is None:
            continue
        style["align"] = style["align"] or level.get("algn")
        props = level.find(f"{{{NS_A}}}defRPr")
        if props is None:
            continue
        if style["size_pt"] is None and (size := props.get("sz")) and size.isdigit():
            style["size_pt"] = int(size) / 100
        if style["bold"] is None and props.get("b") is not None:
            style["bold"] = props.get("b") in {"1", "true"}
        style["color"] = style["color"] or _fill_color(props, theme_colors, clr_map)
    return style


def _has_title(layout: dict[str, Any]) -> bool:
    return any(p["type"] in TITLE_TYPES for p in layout["placeholders"])


def _bodies(layout: dict[str, Any]) -> list[dict[str, Any]]:
    return [p for p in layout["placeholders"] if p["type"] in BODY_TYPES]


def _side_by_side(layout: dict[str, Any], slide_width: int) -> bool:
    """Два тела примерно одной высоты и ширины, стоящие рядом, — признак макета «две колонки»."""
    bodies = [p for p in _bodies(layout) if None not in (p["left"], p["top"], p["width"])]
    for i, a in enumerate(bodies):
        for b in bodies[i + 1 :]:
            narrow = all(0.2 * slide_width <= p["width"] <= 0.6 * slide_width for p in (a, b))
            aligned = abs(a["top"] - b["top"]) <= 0.1 * slide_width
            apart = a["left"] + a["width"] <= b["left"] + 0.02 * slide_width or (
                b["left"] + b["width"] <= a["left"] + 0.02 * slide_width
            )
            if narrow and aligned and apart:
                return True
    return False


def _pick(candidates, key=None) -> dict[str, Any] | None:
    candidates = list(candidates)
    if not candidates:
        return None
    return max(candidates, key=key or (lambda layout: layout["used_by"]))


def _assign_roles(
    layouts: list[dict[str, Any]],
    slide_layout_keys: list[str | None],
    slide_width: int,
    params: dict[str, Any],
) -> dict[str, str | None]:
    if not layouts:
        return {}
    by_key = {layout["key"]: layout for layout in layouts}
    first = slide_layout_keys[0] if slide_layout_keys else None
    last = slide_layout_keys[-1] if slide_layout_keys else None

    def key_of(layout: dict[str, Any] | None) -> str | None:
        return layout["key"] if layout else None

    # Обложка — макет первого слайда шаблона.
    title = first or key_of(_pick(l for l in layouts if l["role"] == "title")) or layouts[0]["key"]
    # Финал — макет последнего слайда, если он похож на финальный, иначе макет с «финальным» названием.
    if last and by_key[last]["role"] == "closing":
        closing = last
    else:
        closing = key_of(_pick(l for l in layouts if l["role"] == "closing"))

    section = key_of(_pick(l for l in layouts if l["role"] == "section"))
    agenda = key_of(_pick(l for l in layouts if l["role"] == "agenda"))
    reserved = {title, closing, section, agenda}

    # Основной контентный макет — самый используемый макет с заголовком.
    # Названия ненадёжны (в шаблонах бывают «N_Титульный слайд» с контентом), поэтому решает частота.
    free = [l for l in layouts if l["key"] not in reserved and _has_title(l)]
    content_layout = _pick(free, key=lambda l: (l["used_by"], l["role"] == "content"))
    content = key_of(content_layout) or title

    # Финальных названий нет — берём макет последнего слайда, если это не контент и не обложка.
    if closing is None:
        closing = last if last and last not in {content, title} else title
    free = [l for l in free if l["key"] != closing]

    columns_rx = re.compile(params.get("two_columns_hint") or "$^", re.IGNORECASE)
    two_columns = key_of(
        _pick(
            (l for l in free if _side_by_side(l, slide_width)),
            key=lambda l: (bool(columns_rx.search(l["name"])), l["used_by"]),
        )
    )

    image_rx = re.compile(params.get("image_hint") or "$^", re.IGNORECASE)
    image = key_of(
        _pick(
            (l for l in free if any(p["type"] == "PICTURE" for p in l["placeholders"])),
            key=lambda l: (bool(image_rx.search(l["name"])), l["used_by"]),
        )
    )

    return {
        "title": title,
        "section": section,
        "agenda": agenda,
        "content": content,
        "two_columns": two_columns,
        "image": image,
        "closing": closing,
    }


def _usage(path: Path, theme: dict[str, Any], params: dict[str, Any]) -> dict[str, Any]:
    """Фактическое использование цветов, гарнитур и кеглей в слайдах, макетах и мастерах."""
    theme_colors = theme.get("colors") or {}
    theme_fonts = theme.get("fonts") or {}
    colors: Counter[str] = Counter()
    color_sources: dict[str, set[str]] = defaultdict(set)
    fonts: Counter[str] = Counter()
    # Роли гарнитур отдельно по слайдам и по макетам/мастерам: в макетах бывают служебные плейсхолдеры
    # другим шрифтом (Calibri), а решает то, чем набраны сами слайды; макеты — если слайдов нет (.potx).
    role_fonts: dict[str, dict[str, Counter[str]]] = {
        part: {"heading": Counter(), "body": Counter()} for part in ("slides", "layouts")
    }
    sizes: Counter[float] = Counter()
    stats: Counter[str] = Counter()

    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        for name in names:
            match = PART_RX.match(name)
            if not match:
                continue
            root = etree.fromstring(zf.read(name), _XML_PARSER)

            for el in root.iter(f"{{{NS_A}}}srgbClr"):
                value = el.get("val")
                if value:
                    color = "#" + value.upper()
                    colors[color] += 1
                    color_sources[color].add("srgb")
            for el in root.iter(f"{{{NS_A}}}schemeClr"):
                scheme_key = SCHEME_ALIASES.get(el.get("val", ""), el.get("val", ""))
                color = theme_colors.get(scheme_key)
                if color:
                    colors[color] += 1
                    color_sources[color].add(scheme_key)

            for el in root.iter(f"{{{NS_A}}}latin"):
                face = el.get("typeface") or ""
                if face.startswith("+mj"):
                    face = theme_fonts.get("major") or ""
                elif face.startswith("+mn"):
                    face = theme_fonts.get("minor") or ""
                if face:
                    fonts[face] += 1
                    # Роль гарнитуры — по кеглю прогона: шрифты темы в шаблонах часто пустая формальность
                    # (Arial), а весь текст слайдов набран фирменной гарнитурой прямо в прогонах.
                    size = el.getparent().get("sz") or ""
                    if size.isdigit():
                        part = "slides" if match.group(1) == "slides" else "layouts"
                        role_fonts[part]["heading" if int(size) >= HEADING_SZ else "body"][face] += 1

            for tag in ("rPr", "defRPr", "endParaRPr"):
                for el in root.iter(f"{{{NS_A}}}{tag}"):
                    size = el.get("sz")
                    if size and size.isdigit():
                        sizes[int(size) / 100] += 1

            if match.group(1) == "slides":
                stats["pictures"] += sum(1 for _ in root.iter(f"{{{NS_P}}}pic"))
                stats["groups"] += sum(1 for _ in root.iter(f"{{{NS_P}}}grpSp"))
                for el in root.iter(f"{{{NS_A}}}graphicData"):
                    uri = el.get("uri", "")
                    if uri == URI_CHART:
                        stats["charts"] += 1
                    elif uri == URI_TABLE:
                        stats["tables"] += 1
                    elif uri == URI_DIAGRAM:
                        stats["smartart"] += 1
        stats["media"] = sum(1 for name in names if name.startswith("ppt/media/"))

    palette = [
        {"hex": color, "count": count, "sources": sorted(color_sources[color])}
        for color, count in colors.most_common(int(params.get("palette_size", 12)))
    ]
    font_list = [{"name": face, "count": count} for face, count in fonts.most_common(int(params.get("fonts_limit", 8)))]
    frequent_sizes = [item for item in sizes.most_common(int(params.get("type_scale_limit", 14))) if item[1] >= 2]
    type_scale = [
        {"size_pt": size, "count": count} for size, count in sorted(frequent_sizes, key=lambda item: -item[0])
    ]
    return {
        "palette": palette,
        "fonts": font_list,
        "font_roles": {
            role: (role_fonts["slides"][role] or role_fonts["layouts"][role]).most_common(1)[0][0]
            for role in ("heading", "body")
            if role_fonts["slides"][role] or role_fonts["layouts"][role]
        },
        "type_scale": type_scale,
        "stats": {key: stats.get(key, 0) for key in ("pictures", "groups", "charts", "tables", "smartart", "media")},
    }
