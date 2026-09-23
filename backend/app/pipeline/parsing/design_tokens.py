"""Токены дизайн-системы шаблона сверх палитры и шрифтов (см. template_parser.py).

- типографика по ролям — заголовок, подзаголовок, текст и его уровни, подпись: гарнитура,
  кегль, начертание, цвет и интервал из стилей мастера и макетов;
- сетка — поля слайда, направляющие PowerPoint, края колонок: по ним верстается
  контент на макетах и проверяются «блоки не по направляющим» и «контент в полях»;
- оформление-обвязка — номер слайда, колонтитул, дата, логотипы: что и где стоит,
  используют ли их слайды шаблона (сборка переносит их на новые слайды, аудит проверяет);
- стиль таблиц и диаграмм — стиль таблицы шаблона и цвета рядов его диаграмм, чтобы
  новые таблицы и графики не были «чужими»;
- доступность шрифтов — какие гарнитуры шаблона подменит LibreOffice при рендере;
- библиотека пиктограмм — мелкие картинки шаблона (иконки) для будущих слайдов.

Каждый блок считается независимо: ошибка в одном не мешает остальным.
"""

import hashlib
import logging
import shutil
import subprocess
from collections import Counter
from pathlib import Path
from typing import Any

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from app.pipeline.parsing.patterns import shape_type

logger = logging.getLogger(__name__)

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
NS_P15 = "http://schemas.microsoft.com/office/powerpoint/2012/main"
SCHEME_ALIASES = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
CHROME_TYPES = {PP_PLACEHOLDER.SLIDE_NUMBER: "slide_number", PP_PLACEHOLDER.FOOTER: "footer", PP_PLACEHOLDER.DATE: "date"}
MAX_ICONS = 150


def safe(name: str, fn, *args, **kwargs):
    """Считает блок токенов; при ошибке — None и запись в лог, разбор шаблона продолжается."""
    try:
        return fn(*args, **kwargs)
    except Exception:  # noqa: BLE001
        logger.exception("Токены «%s» не извлечены", name)
        return None


def _color(element, theme_colors: dict[str, str], clr_map: dict[str, str]) -> str | None:
    if element is None:
        return None
    for child in element:
        if not isinstance(child.tag, str):
            continue
        local = etree.QName(child).localname
        if local == "srgbClr" and child.get("val"):
            return "#" + child.get("val").upper()
        if local == "schemeClr":
            value = child.get("val", "")
            return theme_colors.get(clr_map.get(value) or SCHEME_ALIASES.get(value, value))
        if local == "sysClr":
            return "#" + (child.get("lastClr") or "000000").upper()
    return None


def _clr_map(master) -> dict[str, str]:
    element = master._element.find(f"{{{NS_P}}}clrMap")
    return dict(element.attrib) if element is not None else {}


def _level_style(level, theme: dict[str, Any], clr_map: dict[str, str], default_face: str | None) -> dict[str, Any]:
    style: dict[str, Any] = {"font": None, "size_pt": None, "bold": None, "color": None, "line_spacing": None}
    if level is None:
        style["font"] = default_face
        return style
    props = level.find(f"{{{NS_A}}}defRPr")
    if props is not None:
        if (size := props.get("sz", "")).isdigit():
            style["size_pt"] = int(size) / 100
        if props.get("b") is not None:
            style["bold"] = props.get("b") in {"1", "true"}
        latin = props.find(f"{{{NS_A}}}latin")
        face = latin.get("typeface") if latin is not None else None
        fonts = theme.get("fonts") or {}
        if face and face.startswith("+mj"):
            face = fonts.get("major")
        elif face and face.startswith("+mn"):
            face = fonts.get("minor")
        style["font"] = face or default_face
        style["color"] = _color(props.find(f"{{{NS_A}}}solidFill"), theme.get("colors") or {}, clr_map)
    else:
        style["font"] = default_face
    spacing = level.find(f"{{{NS_A}}}lnSpc/{{{NS_A}}}spcPct")
    if spacing is not None and (value := spacing.get("val", "")).isdigit():
        style["line_spacing"] = int(value) / 100000
    return style


def typography(prs, theme: dict[str, Any], type_scale: list[dict[str, Any]]) -> dict[str, Any]:
    """Роли текста по стилям первого мастера: заголовок, подзаголовок, текст (3 уровня), подпись."""
    master = prs.slide_masters[0]
    clr_map = _clr_map(master)
    fonts = theme.get("fonts") or {}
    styles = master._element.find(f"{{{NS_P}}}txStyles")

    def level(style: str, number: int):
        return styles.find(f"{{{NS_P}}}{style}/{{{NS_A}}}lvl{number}pPr") if styles is not None else None

    result = {
        "title": _level_style(level("titleStyle", 1), theme, clr_map, fonts.get("major")),
        "body": _level_style(level("bodyStyle", 1), theme, clr_map, fonts.get("minor")),
        "body_2": _level_style(level("bodyStyle", 2), theme, clr_map, fonts.get("minor")),
        "body_3": _level_style(level("bodyStyle", 3), theme, clr_map, fonts.get("minor")),
    }
    # Подзаголовок — стиль плейсхолдера subTitle на макете обложки, иначе как текст.
    subtitle = None
    for layout in master.slide_layouts:
        for ph in layout.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
                subtitle = ph._element.find(f"{{{NS_P}}}txBody/{{{NS_A}}}lstStyle/{{{NS_A}}}lvl1pPr")
                break
        if subtitle is not None:
            break
    result["subtitle"] = _level_style(subtitle, theme, clr_map, fonts.get("minor")) if subtitle is not None else dict(result["body"])
    # Подпись — самый мелкий из частых кеглей шаблона (не мельче 8 pt).
    sizes = sorted(item["size_pt"] for item in type_scale or [] if item.get("size_pt") and item["size_pt"] >= 8)
    body_size = result["body"]["size_pt"] or 18
    caption = sizes[0] if sizes else round(body_size * 0.7, 1)
    result["caption"] = {**result["body"], "size_pt": min(caption, body_size), "bold": False}
    return result


FULL_BLEED_AREA = 0.85
# Поля считаются по текстовым областям: фотографии в корпоративных макетах часто уходят под обрез.
TEXT_PLACEHOLDERS = {"TITLE", "CENTER_TITLE", "SUBTITLE", "BODY", "OBJECT", "VERTICAL_TITLE", "VERTICAL_BODY"}


def _mode(values: list[float], step: float = 0.005) -> float | None:
    if not values:
        return None
    counts = Counter(round(value / step) * step for value in values)
    return round(counts.most_common(1)[0][0], 4)


def _guides(element, width: int, height: int) -> tuple[list[float], list[float]]:
    """Направляющие PowerPoint (p15:sldGuideLst): pos — в 1/8 пункта."""
    vertical, horizontal = [], []
    for guide in element.iter(f"{{{NS_P15}}}guide"):
        pos = guide.get("pos", "")
        if not pos.lstrip("-").isdigit():
            continue
        emu = int(pos) * 12700 / 8
        if guide.get("orient") == "horz":
            horizontal.append(round(emu / height, 4))
        else:
            vertical.append(round(emu / width, 4))
    return sorted(set(vertical)), sorted(set(horizontal))


def grid(prs, layouts: list[dict[str, Any]], roles: dict[str, Any]) -> dict[str, Any]:
    """Поля, направляющие и края колонок по плейсхолдерам контентных макетов."""
    width, height = int(prs.slide_width), int(prs.slide_height)
    service = {roles.get(role) for role in ("title", "section", "closing")} - {roles.get("content")}
    lefts, tops, rights, bottoms, x_edges, y_edges = [], [], [], [], [], []
    for layout in layouts:
        if layout["key"] in service:
            continue
        boxes: list[tuple[float, float, float, float]] = []
        for ph in layout.get("placeholders") or []:
            if ph["type"] in {"SLIDE_NUMBER", "FOOTER", "DATE"} or None in (ph["left"], ph["top"], ph["width"], ph["height"]):
                continue
            x, y, w, h = ph["left"] / width, ph["top"] / height, ph["width"] / width, ph["height"] / height
            if w <= 0 or h <= 0 or x < 0 or y < 0 or x + w > 1.001 or y + h > 1.001:
                continue
            if w * h >= FULL_BLEED_AREA:
                continue  # картинка во весь слайд — это фон макета, а не область контента
            x_edges += [round(x, 3), round(x + w, 3)]
            y_edges += [round(y, 3), round(y + h, 3)]
            if ph["type"] in TEXT_PLACEHOLDERS:
                boxes.append((x, y, w, h))
        if not boxes:
            continue
        # Поле макета — расстояние от края слайда до крайней текстовой области, а не до каждого
        # плейсхолдера: иначе у макета «только заголовок» нижним полем станет вся пустота под ним.
        lefts.append(min(x for x, _, _, _ in boxes))
        tops.append(min(y for _, y, _, _ in boxes))
        rights.append(1 - max(x + w for x, _, w, _ in boxes))
        bottoms.append(1 - max(y + h for _, y, _, h in boxes))
    # Поля шаблона — самая крайняя область контента среди макетов: в эти границы помещается
    # всё, что шаблон вообще размещает на слайде. Среднее здесь врёт: макетов «только заголовок»
    # в корпоративных шаблонах больше, чем макетов с текстом во всю высоту.
    margins = {
        "left": round(min(lefts), 4) if lefts else None,
        "right": round(min(rights), 4) if rights else None,
        "top": round(min(tops), 4) if tops else None,
        "bottom": round(min(bottoms), 4) if bottoms else None,
    }
    vertical, horizontal = _guides(prs.part._element, width, height)
    for master in prs.slide_masters:
        more_v, more_h = _guides(master._element, width, height)
        vertical, horizontal = sorted(set(vertical + more_v)), sorted(set(horizontal + more_h))
    safe_area = None
    if None not in margins.values():
        safe_area = [
            margins["left"],
            margins["top"],
            round(1 - margins["left"] - margins["right"], 4),
            round(1 - margins["top"] - margins["bottom"], 4),
        ]
    edge_counts = Counter(x_edges)
    return {
        "margins": margins,
        "safe_area": safe_area,
        "guides": {"vertical": vertical, "horizontal": horizontal},
        # Края, на которые опирается больше одного плейсхолдера, — линии модульной сетки.
        "columns": sorted(edge for edge, count in edge_counts.items() if count >= 2),
        "rows": sorted(edge for edge, count in Counter(y_edges).items() if count >= 2),
    }


def _box(shape, width: int, height: int) -> list[float] | None:
    if None in (shape.left, shape.top, shape.width, shape.height):
        return None
    return [
        round(shape.left / width, 4),
        round(shape.top / height, 4),
        round(shape.width / width, 4),
        round(shape.height / height, 4),
    ]


def chrome(prs) -> dict[str, Any]:
    """Номер слайда, колонтитул, дата и логотипы: где стоят и используют ли их слайды шаблона."""
    width, height = int(prs.slide_width), int(prs.slide_height)
    placeholders: dict[str, list[float]] = {}
    for master in prs.slide_masters:
        for layout in [master, *master.slide_layouts]:
            for ph in layout.placeholders:
                kind = CHROME_TYPES.get(ph.placeholder_format.type)
                if kind and kind not in placeholders and (box := _box(ph, width, height)):
                    placeholders[kind] = box

    slides = list(prs.slides)
    used: Counter[str] = Counter()
    footer_texts: Counter[str] = Counter()
    for slide in slides:
        kinds = set()
        for shape in slide.placeholders:
            kind = CHROME_TYPES.get(shape.placeholder_format.type)
            if kind:
                kinds.add(kind)
                if kind == "footer" and shape.has_text_frame and shape.text_frame.text.strip():
                    footer_texts[shape.text_frame.text.strip()] += 1
        used.update(kinds)
    share = {kind: round(used[kind] / len(slides), 2) if slides else 0.0 for kind in CHROME_TYPES.values()}

    # Логотипы: картинки мастеров и макетов (на слайдах они есть «бесплатно») и одинаковые
    # картинки в одном и том же месте на многих слайдах-примерах.
    logos: list[dict[str, Any]] = []
    seen: set[str] = set()
    for mi, master in enumerate(prs.slide_masters):
        for source, owner in [("master", master), *(("layout", layout) for layout in master.slide_layouts)]:
            for shape in owner.shapes:
                if shape_type(shape) != MSO_SHAPE_TYPE.PICTURE or not (box := _box(shape, width, height)):
                    continue
                if box[2] * box[3] > 0.05:
                    continue  # крупная картинка — фон или иллюстрация, а не логотип
                try:
                    digest = hashlib.sha1(shape.image.blob).hexdigest()[:12]
                except Exception:  # noqa: BLE001
                    continue
                marker = f"{digest}:{box[0]:.2f}:{box[1]:.2f}"
                if marker in seen:
                    continue
                seen.add(marker)
                logos.append({"source": source, "bbox": box, "image": digest, "layout": getattr(owner, "name", None)})
    positions: Counter[tuple] = Counter()
    boxes: dict[tuple, list[float]] = {}
    for slide in slides:
        for shape in slide.shapes:
            if shape_type(shape) != MSO_SHAPE_TYPE.PICTURE or not (box := _box(shape, width, height)):
                continue
            if box[2] * box[3] > 0.05:
                continue
            try:
                digest = hashlib.sha1(shape.image.blob).hexdigest()[:12]
            except Exception:  # noqa: BLE001
                continue
            key = (digest, round(box[0], 2), round(box[1], 2))
            positions[key] += 1
            boxes[key] = box
    for key, count in positions.items():
        if slides and count >= max(3, len(slides) * 0.3):
            logos.append({"source": "slides", "bbox": boxes[key], "image": key[0], "share": round(count / len(slides), 2)})
    return {
        "placeholders": placeholders,
        "usage": share,
        # Номер слайда переносим на новые слайды, если его ставит хотя бы треть слайдов шаблона.
        "slide_number": share.get("slide_number", 0) >= 0.3,
        "footer_text": footer_texts.most_common(1)[0][0] if footer_texts else None,
        "logos": logos[:12],
    }


def table_style(prs, theme: dict[str, Any]) -> dict[str, Any] | None:
    """Стиль таблиц шаблона: id стиля, флаги шапки и полос, заливки и кегль из примеров."""
    styles: Counter[str] = Counter()
    flags: Counter[tuple] = Counter()
    header_fill = band_fill = header_text = None
    font_sizes: Counter[float] = Counter()
    colors = theme.get("colors") or {}
    clr_map = _clr_map(prs.slide_masters[0]) if prs.slide_masters else {}
    for slide in prs.slides:
        for tbl in slide._element.iter(f"{{{NS_A}}}tbl"):
            props = tbl.find(f"{{{NS_A}}}tblPr")
            if props is not None:
                style_id = props.find(f"{{{NS_A}}}tableStyleId")
                if style_id is not None and style_id.text:
                    styles[style_id.text.strip()] += 1
                flags[(props.get("firstRow") in {"1", "true"}, props.get("bandRow") in {"1", "true"})] += 1
            rows = tbl.findall(f"{{{NS_A}}}tr")
            if rows and header_fill is None:
                cell = rows[0].find(f"{{{NS_A}}}tc/{{{NS_A}}}tcPr/{{{NS_A}}}solidFill")
                header_fill = _color(cell, colors, clr_map) if cell is not None else None
                run = rows[0].find(f".//{{{NS_A}}}rPr/{{{NS_A}}}solidFill")
                header_text = _color(run, colors, clr_map) if run is not None else None
            if len(rows) > 1 and band_fill is None:
                cell = rows[1].find(f"{{{NS_A}}}tc/{{{NS_A}}}tcPr/{{{NS_A}}}solidFill")
                band_fill = _color(cell, colors, clr_map) if cell is not None else None
            for props_el in tbl.iter(f"{{{NS_A}}}rPr"):
                if (size := props_el.get("sz", "")).isdigit():
                    font_sizes[int(size) / 100] += 1
    if not styles and not flags:
        return None
    first_row, band_row = flags.most_common(1)[0][0] if flags else (True, True)
    return {
        "style_id": styles.most_common(1)[0][0] if styles else None,
        "first_row": first_row,
        "band_row": band_row,
        "header_fill": header_fill,
        "header_text": header_text,
        "band_fill": band_fill,
        "font_size_pt": font_sizes.most_common(1)[0][0] if font_sizes else None,
    }


def chart_style(prs, theme: dict[str, Any]) -> dict[str, Any] | None:
    """Цвета рядов и оформление диаграмм шаблона — для нативных диаграмм на макетах."""
    colors = theme.get("colors") or {}
    clr_map = _clr_map(prs.slide_masters[0]) if prs.slide_masters else {}
    series_colors: list[str] = []
    kinds: Counter[str] = Counter()
    labels: Counter[bool] = Counter()
    font_sizes: Counter[float] = Counter()
    gap_widths: Counter[int] = Counter()
    charts = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False) or not shape.has_chart:
                continue
            charts += 1
            space = shape.chart._chartSpace
            for plot in space.iter():
                if isinstance(plot.tag, str) and plot.tag.startswith(f"{{{NS_C}}}") and plot.tag.endswith("Chart"):
                    kinds[etree.QName(plot).localname] += 1
            for ser in space.iter(f"{{{NS_C}}}ser"):
                fill = ser.find(f"{{{NS_C}}}spPr/{{{NS_A}}}solidFill")
                line = ser.find(f"{{{NS_C}}}spPr/{{{NS_A}}}ln/{{{NS_A}}}solidFill")
                color = _color(fill, colors, clr_map) or _color(line, colors, clr_map)
                if color and color not in series_colors:
                    series_colors.append(color)
                for point in ser.iter(f"{{{NS_C}}}dPt"):
                    point_color = _color(point.find(f"{{{NS_C}}}spPr/{{{NS_A}}}solidFill"), colors, clr_map)
                    if point_color and point_color not in series_colors:
                        series_colors.append(point_color)
            for show in space.iter(f"{{{NS_C}}}showVal"):
                labels[show.get("val") in {"1", "true", None}] += 1
            for props in space.iter(f"{{{NS_A}}}defRPr"):
                if (size := props.get("sz", "")).isdigit():
                    font_sizes[int(size) / 100] += 1
            for gap in space.iter(f"{{{NS_C}}}gapWidth"):
                if (value := gap.get("val", "")).isdigit():
                    gap_widths[int(value)] += 1
    if not charts:
        return None
    return {
        "charts": charts,
        "series_colors": series_colors[:8],
        "kinds": dict(kinds),
        "data_labels": labels.most_common(1)[0][0] if labels else None,
        "font_size_pt": font_sizes.most_common(1)[0][0] if font_sizes else None,
        "gap_width": gap_widths.most_common(1)[0][0] if gap_widths else None,
    }


def fonts_check(names: list[str]) -> list[dict[str, Any]] | None:
    """Какие гарнитуры шаблона есть для рендера (fontconfig). Недоступные LibreOffice подменит —
    превью и PDF «поплывут», поэтому это предупреждение для пользователя и подсказка, что доложить."""
    binary = shutil.which("fc-match")
    if binary is None:
        return None
    result = []
    for name in dict.fromkeys(n for n in names if n):
        try:
            matched = subprocess.run(
                [binary, "-f", "%{family}", name], capture_output=True, text=True, timeout=5, check=False
            ).stdout.strip()
        except (OSError, subprocess.TimeoutExpired):
            continue
        families = [family.strip().lower() for family in matched.split(",")]
        available = name.strip().lower() in families
        result.append({"name": name, "available": available, "substitute": None if available else (matched or None)})
    return result


def extract_icons(prs, directory: Path, patterns: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Библиотека пиктограмм: мелкие почти квадратные картинки из слайдов шаблона (с витрин иконок —
    в первую очередь). Файлы кладутся в directory под именем-хешем, повторы схлопываются."""
    width, height = int(prs.slide_width), int(prs.slide_height)
    libraries = {pattern["slide_index"] for pattern in patterns if pattern.get("category") == "library"}
    icons: dict[str, dict[str, Any]] = {}

    def walk(shapes, index: int):
        for shape in shapes:
            kind = shape_type(shape)
            if kind == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes, index)
            elif kind == MSO_SHAPE_TYPE.PICTURE:
                yield shape, index

    ordered = sorted(enumerate(prs.slides), key=lambda pair: pair[0] not in libraries)
    directory.mkdir(parents=True, exist_ok=True)
    for index, slide in ordered:
        for shape, slide_index in walk(slide.shapes, index):
            if len(icons) >= MAX_ICONS:
                break
            if not shape.width or not shape.height:
                continue
            ratio = shape.width / shape.height
            area = (shape.width / width) * (shape.height / height)
            if not (0.6 <= ratio <= 1.65) or area > 0.02:
                continue
            try:
                image = shape.image
                blob, ext = image.blob, image.ext
            except Exception:  # noqa: BLE001 — связанные (внешние) картинки пропускаем
                continue
            digest = hashlib.sha1(blob).hexdigest()[:16]
            if digest in icons:
                icons[digest]["uses"] += 1
                continue
            name = f"{digest}.{ext.lower()}"
            (directory / name).write_bytes(blob)
            icons[digest] = {
                "id": digest,
                "file": name,
                "slide_index": slide_index,
                "size_px": list(image.size) if getattr(image, "size", None) else None,
                "from_library": slide_index in libraries,
                "uses": 1,
            }
    return list(icons.values())
